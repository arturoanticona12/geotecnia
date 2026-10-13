"""Figuras del memo como gráficos nativos de Excel (revisión 12, pedido del usuario 2026-10-13: «que replique tal cual las figuras
del informe»). Lo llama a28_hoja_calculo.py. Escribe dos hojas:
  Datos_figuras: los datos de cada figura. Curvas, banda y ensayos con fórmulas ligadas a Parametros, Ensayos y Base_mundial;
                 φ′1 por set (a34), contraste CU (a16), Bg (a16) y perfil V como valores, con su fuente en la fila 1.
  Figuras: las seis figuras del memo con la misma numeración, series, colores, marcadores, tipos de línea, escalas y leyendas que
           a11, a16 y a34 (matplotlib). Excel no tiene relleno entre curvas en un gráfico de dispersión: la banda y la zona gris
           se dibujan con segmentos verticales anchos y el rayado con segmentos diagonales."""
import math
import numpy as np, pandas as pd
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Font

# paleta de a11 (dataviz)
INK2, MUTED, GRID, SURF = "52514E", "8A8984", "E6E5E1", "FCFCFB"
VIO, VIO_B, ZONA, RAYA, GRIS_M = "4A3AA7", "E3E0F3", "EFEEE9", "D2D1CB", "C9C8C2"
LIT = {"Mármol": ("2A78D6", "circle"), "Hornfels": ("EB6834", "square"), "Intrusivo": ("1BAF7A", "triangle"), "Mezcla": ("EDA100", "diamond")}
LIT_TXT = {"Mármol": "Mármol", "Hornfels": "Hornfels", "Intrusivo": "Intrusivo",
           "Mezcla": "Mezclas (Golder Associates Perú, 2021; Anddes Geolab, 2024)"}
# Leps (1970) e Indraratna et al. (1993) en la forma de la Ec. 9 de Leps 2.0 (m = 0.9); estilos de a11
FONDO = [("Leps (1970), envolvente superior", 46.0, "dashDot"), ("Leps (1970), envolvente media", 41.0, "dash"),
         ("Leps (1970), envolvente inferior", 36.0, "sysDot"), ("Cota inferior de Indraratna et al. (1993)", 33.0, "solid")]
NEG = Font(bold=True)


def _txt(sz):
    cp = CharacterProperties(sz=int(sz * 100))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _ln(color, w_pt, dash="solid"):
    return LineProperties(solidFill=color, w=int(w_pt * 12700), prstDash=dash)


def agregar(wb, *, HB, SNX, SNW, EB, TOPE, PE, PEM, A16BG, A16C, S34, N_AJ, NW, rbg):
    wd = wb.create_sheet("Datos_figuras")
    wf = wb.create_sheet("Figuras", 1)
    col = [1]

    def bloque(titulo, cab, filas, ancho=11):
        c0 = col[0]
        wd.cell(1, c0, titulo).font = NEG
        for j, h in enumerate(cab):
            wd.cell(2, c0 + j, h).font = NEG
            wd.column_dimensions[wd.cell(2, c0 + j).column_letter].width = ancho
        for i, f in enumerate(filas):
            for j, v in enumerate(f):
                if v is not None and not (isinstance(v, float) and math.isnan(v)):
                    wd.cell(3 + i, c0 + j, v)
        col[0] += len(cab) + 1
        return c0, len(filas)

    L = lambda c, r: f"{wd.cell(r, c).column_letter}{r}"
    # fórmulas de la curva (nombres definidos de Parametros)
    phi_c = lambda x: f"DEGREES(ATAN(TAN(RADIANS(pe0-omega*e_base))*({x}/sref)^(m_-1)))"
    phi_e = lambda x, e: f"DEGREES(ATAN(TAN(RADIANS(pe0-omega*({e})))*({x}/sref)^(m_-1)))"
    lo_f = lambda x: "MIN(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$E${h}-omega*e_base))*({x}/sref)^(Parametros!$B${h}-1)))" for h in HB) + ")"
    hi_f = lambda x: "MAX(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$F${h}-omega*e_base))*({x}/sref)^(Parametros!$B${h}-1)))" for h in HB) + ")"
    l70 = lambda x, p1: f"DEGREES(ATAN(TAN(RADIANS({p1}))*({x}/sref)^(0.9-1)))"

    # ---------------------------------------------------------------- bloques de datos
    s_tope = 1.013 * (math.tan(math.radians(float(wb["Parametros"]["B2"].value) - 20 * EB)) / math.tan(math.radians(TOPE))) ** (1 / (1 - 0.85))
    xg = np.unique(np.round(np.concatenate([np.geomspace(0.01, 15, 110), [s_tope, SNX, SNW, 0.1, 1, 8]]), 5))
    filas = []
    for i, x in enumerate(xg):
        r = 3 + i; X = f"A{r}"   # se reubica abajo con el desplazamiento de columna
        filas.append([x])
    cG, nG = bloque("Curvas (fórmulas; a11, Figs. 2, 3 y 5)", ["σ′n (MPa)", "φ′s sin tope (°)", "φ′s curva con tope (°)", "φ′s sin tope, tramo sobre el tope (°)",
                    "Banda inf. (°)", "Banda sup. (°)"] + [nm for nm, _, _ in FONDO] + ["τ curva con tope (MPa)"] + [f"τ {nm}" for nm, _, _ in FONDO], filas)
    for i in range(nG):
        r = 3 + i; x = L(cG, r); ph = L(cG + 1, r)
        wd.cell(r, cG + 1, "=" + phi_c(x)); wd.cell(r, cG + 2, f"=MIN(tope,{ph})"); wd.cell(r, cG + 3, f"=IF({ph}>tope,{ph},NA())")
        wd.cell(r, cG + 4, "=" + lo_f(x)); wd.cell(r, cG + 5, "=" + hi_f(x))
        for j, (_, p1, _) in enumerate(FONDO):
            wd.cell(r, cG + 6 + j, "=" + l70(x, p1)); wd.cell(r, cG + 11 + j, f"={x}*TAN(RADIANS({L(cG + 6 + j, r)}))")
        wd.cell(r, cG + 10, f"={x}*TAN(RADIANS({L(cG + 2, r)}))")

    def banda_seg(titulo, xs, tau=False):
        """Segmentos verticales de la banda: (x, inf), (x, sup), vacío."""
        f = []
        for x in xs:
            f += [[x, None], [x, None], [None, None]]
        c0, n = bloque(titulo, ["σ′n (MPa)", "τ (MPa)" if tau else "φ′s (°)"], f)
        for i, x in enumerate(xs):
            for k, g in ((0, lo_f), (1, hi_f)):
                r = 3 + 3 * i + k; xc = L(c0, r)
                wd.cell(r, c0 + 1, f"={xc}*TAN(RADIANS({g(xc)}))" if tau else "=" + g(xc))
        return c0, n

    cB1, nB1 = banda_seg("Banda, segmentos (Figs. 2 y 3)", np.geomspace(0.1, 15, 130))
    cB2, nB2 = banda_seg("Banda τ, segmentos (Fig. 5)", np.linspace(0.02, 15, 170), tau=True)

    def zonas(titulo, y0, y1, log=True, xmax=15.0):
        """Zona gris (segmentos verticales entre SNX y SNW) y rayado (diagonales sobre SNW) en datos de la figura."""
        f = []
        for x in (np.geomspace(SNX, SNW, 26) if log else np.linspace(SNX, SNW, 26)):
            f += [[x, y0], [x, y1], [None, None]]
        # rayado ///: rectas a 45° en coordenadas de pantalla (eje x en décadas o MPa, eje y normalizado)
        u = (lambda x: math.log10(x)) if log else (lambda x: x)
        ui = (lambda v: 10 ** v) if log else (lambda v: v)
        a, b = u(SNW), u(xmax); paso = (b - a) / 7.0; k = (b - a) / 2.6   # pendiente de las rayas
        g = []
        v0 = a - k
        while v0 < b:
            pts = []
            for t in (0.0, 1.0):
                pts.append((v0 + k * t, t))
            (ua, ta), (ub, tb) = pts
            if ua < a:
                ta = ta + (a - ua) / k; ua = a
            if ub > b:
                tb = tb - (ub - b) / k; ub = b
            if ub > ua:
                g += [[None, None, ui(ua), y0 + ta * (y1 - y0)], [None, None, ui(ub), y0 + tb * (y1 - y0)], [None, None, None, None]]
            v0 += paso
        from itertools import zip_longest
        filas = [[*(fa or [None, None])[:2], *(gg or [None, None, None, None])[2:]] for fa, gg in zip_longest(f, g)]
        return bloque(titulo, ["Zona gris x", "Zona gris y", "Rayado x", "Rayado y"], filas)

    cZ1, nZ1 = zonas("Zonas Fig. 2 (valores)", 22, 56)
    cZ2, nZ2 = zonas("Zonas Fig. 5 (valores)", 0, 10.5, log=False)
    cZ3, nZ3 = zonas("Zonas Fig. 6b (valores)", 22, 58)

    # ensayos del ajuste (fórmulas ligadas a la hoja Ensayos, filas 2 a 1 + N_AJ)
    SER = [(l, tx) for l in LIT for tx in (True, False)]
    cE, nE = bloque("Ensayos del ajuste a e base (fórmulas; Ensayos!M y P)", ["σ′n (MPa)", "Todos φ′s (°)"] +
                    [f"{l} {'TX' if tx else 'DS'} φ′s" for l, tx in SER] + [f"{l} {'TX' if tx else 'DS'} τ" for l, tx in SER], [[None]] * N_AJ)
    for i in range(N_AJ):
        r = 3 + i; re_ = 2 + i
        wd.cell(r, cE, f"=Ensayos!M{re_}"); wd.cell(r, cE + 1, f"=Ensayos!P{re_}")
        for j, (l, tx) in enumerate(SER):
            cond = f'AND(Ensayos!$D{re_}="{l}",{"" if tx else "NOT("}LEFT(Ensayos!$E{re_},7)="Triaxia"{"" if tx else ")"})'
            wd.cell(r, cE + 2 + j, f"=IF({cond},Ensayos!P{re_},NA())")
            wd.cell(r, cE + 2 + len(SER) + j, f"=IF({cond},Ensayos!M{re_}*TAN(RADIANS(Ensayos!P{re_})),NA())")

    # base mundial llevada a e base (fórmulas ligadas a Base_mundial)
    cW, nW = bloque("Base mundial a e base (fórmulas; Base_mundial)", ["σ′n (MPa)", "φ′s a e base (°)", "τ a e base (MPa)"], [[None]] * NW)
    for i in range(NW):
        r = 3 + i; rb = 2 + i
        wd.cell(r, cW, f"=Base_mundial!C{rb}"); wd.cell(r, cW + 1, f"=Base_mundial!E{rb}-omega*(e_base-Base_mundial!B{rb})")
        wd.cell(r, cW + 2, f"={L(cW, r)}*TAN(RADIANS({L(cW + 1, r)}))")

    # Fig. 1: φ′1 por set (a34, valores) y rectas
    sets = [(l, ds) for l in LIT for ds in (False, True)]
    f = []
    for _, r in S34.iterrows():
        f.append([r.e0] + [r.phi1 if (r.litologia == l and bool(r.corte_directo) == ds) else None for l, ds in sets])
    cS, nS = bloque("Fig. 1: φ′1 por set (a34_phi1_sets.csv, m = 0.85)", ["e0 media"] + [f"{l} {'DS' if ds else 'TX'}" for l, ds in sets], f)
    cR, nR = bloque("Fig. 1: rectas (fórmulas)", ["e", "Curva propuesta"] + [nm for nm, _, _ in FONDO], [[0.0], [0.8]])
    for i in range(2):
        r = 3 + i
        wd.cell(r, cR + 1, f"=pe0-omega*{L(cR, r)}")
        for j, (_, p1, _) in enumerate(FONDO):
            wd.cell(r, cR + 2 + j, p1)

    # Fig. 6: perfiles de e
    zz = PE["z"]
    cP, nP = bloque("Fig. 6a: e con la profundidad (a13)", ["Profundidad (m)", "e perfil A", "e perfil V"],
                    [[z, a, v] for z, a, v in zip(zz, PE["e_A"], PE["e_V"])])
    cPb, _ = bloque("Fig. 6a: e base (fórmula)", ["Profundidad (m)", "e base"], [[0.0], [700.0]])
    wd.cell(3, cPb + 1, "=e_base"); wd.cell(4, cPb + 1, "=e_base")
    xs6 = np.geomspace(0.1, 15, 90)
    cQ, nQ = bloque("Fig. 6b: φ′s sin tope, como la Fig. 6 del memo (fórmulas; e perfil V como valor, a13)",
                    ["σ′n (MPa)", "e perfil A", "e perfil V", "φ′s e base (°)", "φ′s perfil A (°)", "φ′s perfil V (°)"],
                    [[x, None, (None if np.isnan(float(PEM.eV_of_sv(x))) else round(float(PEM.eV_of_sv(x)), 5))] for x in xs6])
    for i in range(nQ):
        r = 3 + i; x = L(cQ, r)
        wd.cell(r, cQ + 1, f"=IF(MAX({x},sv_piso)<=1,e_base-cc_bajo*LOG10(MAX({x},sv_piso)/sv_piso),e_base-cc_bajo*LOG10(1/sv_piso)-cc_alto*LOG10({x}))")
        wd.cell(r, cQ + 3, "=" + phi_e(x, 'e_base'))
        wd.cell(r, cQ + 4, "=" + phi_e(x, L(cQ + 1, r)))
        wd.cell(r, cQ + 5, f'=IF(ISNUMBER({L(cQ + 2, r)}),{phi_e(x, L(cQ + 2, r))},NA())')

    # Fig. 3: contraste CU (a16, valores)
    GRU = [("CU Anddes (2024), 6 in, máxima oblicuidad", "E87BA4", "square", "Triaxiales CU, Anddes (2024), 6 in"),
           ("CU Golder (2010), fase A, 2.8 y 4 in, máximo desviador", "008300", "triangle", "Triaxiales CU, Golder (2010), fase A")]
    f = []
    for _, r in A16C.iterrows():
        f.append([r.sn] + [r.phis_e065 if r.grupo == g else None for g, *_ in GRU])
    cC, nC = bloque("Fig. 3: triaxiales CU de contraste a e base (a16_contraste.csv)", ["σ′n (MPa)"] + [g[3] for g in GRU], f)

    # Fig. 4: Bg (a16, valores) y tendencia (fórmula)
    BGS = [(l, esc) for l in ("Hornfels", "Intrusivo", "Mármol") for esc in sorted(A16BG[A16BG.litologia == l].escala.unique())]
    f = []
    for _, r in A16BG.iterrows():
        f.append([r.s3, r.Bg] + [r.Bg if (r.litologia == l and r.escala == esc) else None for l, esc in BGS])
    cG4, nG4 = bloque("Fig. 4: Bg de Marsal (a16_bg.csv)", ["σ′3 (MPa)", "Bg (%)"] + [f"{l} {esc}" for l, esc in BGS], f)
    cT, nT = bloque("Fig. 4: tendencia lineal (fórmula)", ["σ′3 (MPa)", "Bg (%)"], [[0.0], [2.6]])
    xr = f"{L(cG4, 3)}:{L(cG4, 2 + nG4)}"; yr = f"{L(cG4 + 1, 3)}:{L(cG4 + 1, 2 + nG4)}"
    for i in range(2):
        wd.cell(3 + i, cT + 1, f"=INTERCEPT({yr},{xr})+SLOPE({yr},{xr})*{L(cT, 3 + i)}")

    # ---------------------------------------------------------------- gráficos
    def serie(ch, xc, yc, n, titulo, linea=None, marcador=None, msize=5, lleno=True, mcol=None, r0=3, leyenda=True):
        s = Series(Reference(wd, min_col=yc, min_row=r0, max_row=r0 + n - 1), Reference(wd, min_col=xc, min_row=r0, max_row=r0 + n - 1), title=titulo)
        s.smooth = False
        s.graphicalProperties = GraphicalProperties(ln=_ln(*linea) if linea else LineProperties(noFill=True))
        if marcador:
            s.marker.symbol = marcador; s.marker.size = msize
            s.marker.graphicalProperties = (GraphicalProperties(solidFill=mcol, ln=_ln(SURF if lleno and mcol != INK2 else mcol, 0.6)) if lleno
                                            else GraphicalProperties(noFill=True, ln=_ln(mcol, 1.2)))
        else:
            s.marker.symbol = "none"
        ch.series.append(s)
        if not leyenda:
            from openpyxl.chart.legend import Legend
            if ch.legend is None:
                ch.legend = Legend()
            ch.legend.legendEntry.append(LegendEntry(idx=len(ch.series) - 1, delete=True))
        return s

    def grafico(xt, yt, x0, x1, y0, y1, log=False, ymaj=None, w=17.0, h=15.0, invy=False):
        ch = ScatterChart(); ch.style = 13; ch.title = None; ch.scatterStyle = "lineMarker"
        ch.x_axis.title = xt; ch.y_axis.title = yt
        ch.x_axis.scaling.min, ch.x_axis.scaling.max = x0, x1
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = y0, y1
        if log:
            ch.x_axis.scaling.logBase = 10
        if invy:
            ch.y_axis.scaling.orientation = "maxMin"; ch.x_axis.crosses = "max"
        else:
            ch.x_axis.crosses = "min"
        ch.y_axis.crosses = "min"
        if ymaj:
            ch.y_axis.majorUnit = ymaj
        for ax in (ch.x_axis, ch.y_axis):
            ax.delete = False
            ax.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=_ln(GRID, 0.6)))
            ax.txPr = _txt(8); ax.number_format = "General"
            ax.spPr = GraphicalProperties(ln=_ln(MUTED, 0.75))
        ch.legend.position = "b"; ch.legend.txPr = _txt(7)
        ch.width, ch.height = w, h
        return ch

    def zonas_series(ch, cZ, nZ):
        serie(ch, cZ, cZ + 1, nZ, f"{SNX:.2f}–{SNW:.2f} MPa: sin ensayos de Antamina", linea=(ZONA, 4.5))
        serie(ch, cZ + 2, cZ + 3, nZ, f"Sobre {SNW:.2f} MPa: extrapolación", linea=(RAYA, 0.75))

    def fondo(ch, tau=False):
        for j, (nm, _, dash) in enumerate(FONDO):
            serie(ch, cG, cG + (11 if tau else 6) + j, nG, nm, linea=(MUTED, 1.1, dash))

    def ensayos(ch, tau=False):
        for j, (l, tx) in enumerate(SER):
            c_, mk = LIT[l]
            serie(ch, cE, cE + 2 + (len(SER) if tau else 0) + j, nE, f"{LIT_TXT[l]}, {'triaxial drenado' if tx else 'corte directo'}",
                  marcador=mk, msize=6, lleno=tx, mcol=c_)

    def mundo(ch, tau=False):
        serie(ch, cW, cW + (2 if tau else 1), nW, f"Base mundial de Leps 2.0 (Cylwik et al., 2025), llevada a e = {EB:.2f}",
              marcador="circle", msize=2, lleno=True, mcol=GRIS_M)

    XL = "Esfuerzo normal efectivo en el plano de falla, σ′n (MPa)"; YL = f"Ángulo de fricción secante, φ′s (°), a e = {EB:.2f}"
    figs = []
    # Fig. 1
    ch = grafico("Relación de vacíos inicial del ensayo, e0 (media del set)", "φ′1 del set (°), ángulo secante a σ′n = 10pa", 0.2, 0.55, 32, 50, ymaj=2, h=13)
    serie(ch, cR, cR + 1, 2, f"Curva propuesta: φ′1 = {float(wb['Parametros']['B2'].value):.1f}° − 20°·e", linea=(VIO, 2.2))
    for j, (nm, p1, dash) in enumerate(FONDO):
        serie(ch, cR, cR + 2 + j, 2, f"{nm} (φ′1 = {p1:.0f}°)", linea=(MUTED, 1.0, dash))
    for j, (l, ds) in enumerate(sets):
        if S34[(S34.litologia == l) & (S34.corte_directo == ds)].empty:
            continue
        c_, mk = LIT[l]
        serie(ch, cS, cS + 1 + j, nS, f"{'Mezclas' if l == 'Mezcla' else l}, {'corte directo' if ds else 'triaxial'}", marcador=mk, msize=7, lleno=not ds, mcol=c_)
    figs.append(("fig:phi1e", ch))
    # Fig. 2
    ch = grafico(XL, YL, 0.1, 15, 22, 56, log=True, ymaj=5)
    zonas_series(ch, cZ1, nZ1)
    serie(ch, cB1, cB1 + 1, nB1, "Incertidumbre de la curva", linea=(VIO_B, 4.0))
    mundo(ch); fondo(ch)
    serie(ch, cG, cG + 2, nG, f"Curva propuesta (m = 0.85), con tope de {TOPE:.0f}°", linea=(VIO, 2.4))
    serie(ch, cG, cG + 3, nG, "Curva sin tope", linea=(VIO, 1.2, "sysDot"))
    ensayos(ch)
    figs.append(("fig:phi", ch))
    # Fig. 3
    ch = grafico(XL, YL, 0.1, 8, 15, 60, log=True, ymaj=5, h=14)
    serie(ch, cB1, cB1 + 1, nB1, "Incertidumbre de la curva", linea=(VIO_B, 4.0))
    serie(ch, cG, cG + 1, nG, "Curva propuesta (m = 0.85)", linea=(VIO, 2.2))   # la Fig. 3 del memo (a16) no lleva tope
    fondo(ch)
    serie(ch, cE, cE + 1, nE, f"Ensayos del ajuste ({N_AJ}: triaxial CD y corte directo)", marcador="circle", msize=5, lleno=True, mcol=INK2)
    for j, (g, c_, mk, lab) in enumerate(GRU):
        serie(ch, cC, cC + 1 + j, nC, f"{lab} ({int((A16C.grupo == g).sum())})", marcador=mk, msize=6, lleno=False, mcol=c_)
    figs.append(("fig:contraste", ch))
    # Fig. 4
    ch = grafico("Esfuerzo efectivo de confinamiento, σ′3 (MPa)", "Índice de rotura de Marsal, Bg (%)", 0, 2.7, 0, 45, ymaj=5, h=11)
    COLB = {"Mármol": ("2A78D6", "circle"), "Hornfels": ("EB6834", "square"), "Intrusivo": ("1BAF7A", "triangle")}
    for j, (l, esc) in enumerate(BGS):
        c_, mk = COLB[l]
        serie(ch, cG4, cG4 + 2 + j, nG4, f"{l}, probeta de {esc}", marcador=mk, msize=7, lleno=esc != "6 in", mcol=c_)
    serie(ch, cT, cT + 1, 2, f"Tendencia lineal (r = {rbg:.2f})", linea=(INK2, 1.0, "dash"))
    ch.x_axis.majorUnit = 0.5
    figs.append(("fig:bg", ch))
    # Fig. 5
    ch = grafico(XL, "Resistencia al corte, τ (MPa)", 0, 15, 0, 10.5, ymaj=2, h=14)
    ch.x_axis.majorUnit = 2
    zonas_series(ch, cZ2, nZ2)
    serie(ch, cB2, cB2 + 1, nB2, "Incertidumbre de la curva", linea=(VIO_B, 3.0))
    mundo(ch, tau=True); fondo(ch, tau=True)
    serie(ch, cG, cG + 10, nG, f"Curva propuesta, e = {EB:.2f}, con tope de {TOPE:.0f}°", linea=(VIO, 2.4))
    ensayos(ch, tau=True)
    figs.append(("fig:tau", ch))
    # Fig. 6a y 6b
    cha = grafico("Relación de vacíos, e", "Profundidad bajo la superficie (m)", 0.08, 0.70, 0, 700, ymaj=100, w=8.5, h=12, invy=True)
    cha.x_axis.majorUnit = 0.1
    serie(cha, cPb + 1, cPb, 2, f"e = {EB:.2f} constante (base)", linea=(VIO, 2.0))
    serie(cha, cP + 1, cP, nP, "Perfil A, compresibilidad de Antamina", linea=("E87BA4", 2.0))
    serie(cha, cP + 2, cP, nP, "Perfil V, Valenzuela et al. (2008)", linea=("008300", 2.0))
    chb = grafico("σ′n (MPa)", "φ′s (°), curva propuesta", 0.1, 15, 22, 58, log=True, ymaj=5, w=12.5, h=12)
    zonas_series(chb, cZ3, nZ3)
    fondo(chb)
    serie(chb, cQ, cQ + 3, nQ, f"e = {EB:.2f} constante (base)", linea=(VIO, 2.2))
    serie(chb, cQ, cQ + 4, nQ, "Perfil A, compresibilidad de Antamina", linea=("E87BA4", 2.0))
    serie(chb, cQ, cQ + 5, nQ, "Perfil V, Valenzuela et al. (2008)", linea=("008300", 2.0))
    figs.append(("fig:e", (cha, chb)))
    return wf, figs


def colocar(wf, figs, rotulos):
    """Coloca las figuras en la hoja Figuras, una bajo la otra, con su número y título del memo."""
    fila = 1
    wf.column_dimensions["A"].width = 3
    wf.cell(fila, 2, "Figuras del memorándum, recalculadas con las fórmulas de este libro (datos en Datos_figuras). "
                     "La banda y la zona gris se dibujan con segmentos verticales; el rayado, con segmentos diagonales.").font = Font(italic=True)
    fila += 2
    for lab, ch in figs:
        num, tit = rotulos.get(lab, ("", ""))
        wf.cell(fila, 2, f"Figura {num}. {tit}").font = NEG
        fila += 1
        if isinstance(ch, tuple):
            wf.add_chart(ch[0], f"B{fila}"); wf.add_chart(ch[1], f"J{fila}"); h = ch[0].height
        else:
            wf.add_chart(ch, f"B{fila}"); h = ch.height
        fila += int(h / 0.5) + 3
