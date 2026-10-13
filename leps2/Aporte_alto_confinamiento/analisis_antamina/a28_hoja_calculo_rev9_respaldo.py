"""Hoja de cálculo del memo (comentario WA 2026-10-12: «genera una hoja de cálculo que incluya todos los datos, cálculos y curvas»).
Escribe eor_antamina/curva_leps2_antamina_calculos.xlsx con fórmulas vivas de Excel ligadas a la hoja Parametros (celdas con
nombre): ensayos (120, con su uso), curva propuesta y alternativas de e, banda, cotas, hipótesis de m, base mundial (741),
tangentes para FLAC, relación de vacíos, perfiles e(z), edómetros y tamaño de partícula, con gráficos. Escribe también
a28_valores_control.json con los valores que deben dar las fórmulas (para verificarlas al recalcular en Excel)."""
import json, math, sys
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
sys.stdout.reconfigure(encoding="utf-8")
AQ = Path(__file__).parent; sys.path.insert(0, str(AQ.parent / "analisis")); sys.path.insert(0, str(AQ))
from lepsp_r6 import datos  # noqa: E402
import a13_perfil_e as PEM  # noqa: E402
J = lambda n: json.load(open(AQ / n, encoding="utf-8"))
C_DS = J("a9_leps2_antamina.json")["datos"]["c_ds"]; A12 = J("a12_leps2_ponderada.json"); C2 = A12["curvas"]; OP2 = C2["Curva propuesta (mediana mundial)"]; PE = J("a13_perfil_e.json")
R19 = J("a19_m_covariables.json"); EDO = J("a17_edometros.json"); A16 = J("a16_resultados.json")
PE0, M, OM, EB, PA = OP2["pe0"], OP2["m"], 20.0, 0.45, 0.1013
SNX = 4.60; SNW = R19["base_sn_max"]
T27 = pd.read_csv(AQ / "a27_ensayos_todos.csv")
ORD = {"Ajuste de la curva propuesta": 0, "Contraste": 2, "No usado": 3}
T27 = T27.assign(_o=T27.uso.map(ORD)).sort_values(["_o"], kind="stable").drop(columns="_o").reset_index(drop=True)
OUT = AQ.parent / "eor_antamina" / "curva_leps2_antamina_calculos.xlsx"
AZUL = Font(color="1F4E9E"); NEG = Font(bold=True); CAB = PatternFill("solid", fgColor="E8E6F2"); ENT = PatternFill("solid", fgColor="EAF1FB")
wb = Workbook(); control = {}


def hoja(nombre, cab, filas, anchos=None):
    ws = wb.create_sheet(nombre)
    ws.append(cab)
    for c in ws[1]:
        c.font = NEG; c.fill = CAB; c.alignment = Alignment(wrap_text=True, vertical="top")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos or [14] * len(cab), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


# ------------------------------------------------------------------ Leame
ws = wb.active; ws.title = "Leame"
for t in ["Curva de resistencia Leps 2.0 del desmonte ROM de Antamina: datos, cálculos y curvas (memo de Stantec, 12 de octubre de 2026)",
          "",
          "Las celdas en azul de la hoja Parametros son las entradas; todo lo demás se calcula con fórmulas a partir de ellas (nombres definidos: pe0, omega, m_, e_base, sref).",
          "Leps 2.0 (Cylwik et al. 2025, Ec. 26): τ = tan(φ′e=0 − ω·e)·(10pa)^(1−m)·σ′n^m; φ′s = atan(τ/σ′n).",
          "Unidades: esfuerzos en MPa; ángulos en grados.",
          "",
          "Hojas:",
          "  Parametros: parámetros de la curva propuesta, hipótesis de curvatura y control con la composición del DDE.",
          "  Ensayos: los 120 ensayos de resistencia con su uso (ajuste, contraste, no usado), esfuerzos de falla, conversión de Leps 2.0 (fórmula) y corrección a e base (Ec. 27).",
          "  Curva: curva propuesta a e base, 0.55, 0.65 y perfiles A y V; banda de incertidumbre; cotas de Indraratna et al. (1993) y Hawley y Cunning (2017); recta de 37°; sensibilidad del DDE; gráficos.",
          "  phi1_vs_e: φ′1 de cada ensayo y por set contra e, con las rectas de la curva propuesta y de las clases de Leps 2.0 (como la figura de Cylwik et al. 2025).",
          "  Hipotesis_m: φ′s según la hipótesis de curvatura (P10, mediana y P90 de la base mundial).",
          "  Base_mundial: los ensayos de la base de Leps 2.0 (sin los censurados, como en el memo) frente a la curva evaluada a su propia e (Δφ′s) y conteos por rango de esfuerzo.",
          "  FLAC_tangente: Mohr-Coulomb tangente a la curva según σ′3 (para verificar la implementación en FLAC2D).",
          "  Relacion_vacios: e de colocación calculada con los pesos unitarios de los estudios de Antamina.",
          "  Perfiles_e: e y σ′v con la profundidad (perfiles A y V).",
          "  Edometros: edómetros de desmonte de Antamina y e resultante con la ley del perfil A.",
          "  Tamano_particula: sensibilidad al tamaño (Frossard et al. 2012).",
          "",
          "Fuentes: base de ensayos de Antamina verificada contra los PDF (datos_antamina/base, verificacion_ocr); base mundial de Cylwik et al. (2025);",
          "Informe final de estudio del DDE (Stantec 2026, IFF-115-G-0001 Rev. B) para la composición de control y la sensibilidad de 33°/30°."]:
    ws.append([t])
ws.column_dimensions["A"].width = 150

# ------------------------------------------------------------------ Parametros
ws = hoja("Parametros", ["Parámetro", "Valor", "Unidad", "Nota"], [], [34, 12, 10, 90])
PAR = [("pe0", "φ′e=0, curva propuesta", PE0, "°", "Ajuste con todos los ensayos con igual peso (a12)"),
       ("omega", "ω", OM, "°", "Leps 2.0, Ec. 25"),
       ("m_", "m (curvatura)", M, "-", "Mediana de la base mundial (170 sets; a18)"),
       ("e_base", "e base de colocación", EB, "-", "Desmonte ROM (sección 7 del memo)"),
       ("pa", "pa", PA, "MPa", "Presión atmosférica"),
       ("sref", "10·pa", None, "MPa", "Esfuerzo de referencia")]
for k, (nm, et, v, u, n) in enumerate(PAR, 2):
    ws.cell(k, 1, et); ws.cell(k, 3, u); ws.cell(k, 4, n)
    c = ws.cell(k, 2, v if v is not None else "=10*pa"); c.font = AZUL if v is not None else Font(); c.fill = ENT if v is not None else PatternFill()
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"Parametros!$B${k}")
ws.cell(8, 1, "φ′1 a e base"); ws.cell(8, 2, "=pe0-omega*e_base"); ws.cell(8, 3, "°")
control["phi1"] = PE0 - OM * EB
ws.cell(10, 1, "Hipótesis de curvatura").font = NEG
ws.cell(11, 1, "Hipótesis"); ws.cell(11, 2, "m"); ws.cell(11, 3, "φ′e=0 (°)"); ws.cell(11, 4, "Nota"); ws.cell(11, 5, "φ′e=0 P5 (°)"); ws.cell(11, 6, "φ′e=0 P95 (°)")
HIP = [("Mundial P90", "P90 de la base mundial"), ("Curva propuesta (mediana mundial)", "Mediana (curva propuesta)"), ("Mundial P10", "P10 de la base mundial")]
for k, (key, et) in enumerate(HIP, 12):
    ws.cell(k, 1, et); ws.cell(k, 2, C2[key]["m"]); ws.cell(k, 3, C2[key]["pe0"]); ws.cell(k, 4, "Nivel reajustado a los ensayos con esa m (a12; la mediana se reproduce con Solver en la hoja Ajuste)")
    ws.cell(k, 5, C2[key]["pe0_b"][0]); ws.cell(k, 6, C2[key]["pe0_b"][1])
for k_, (nm_, et_, v_, u_, n_) in enumerate([("cc_bajo", "Cc perfil A, 0.2–0.8 MPa", PE["Cc1"], "-", "Edómetro de mármol suelto (a13)"),
                                                ("cc_alto", "Cc perfil A, 1.2–4.8 MPa", PE["Cc2"], "-", "Edómetro de mármol suelto (a13)"),
                                                ("sv_piso", "σ′v mínima de la ley del perfil A", 0.05, "MPa", "a13"),
                                                ("c_ds", "Peso del residuo del corte directo", C_DS, "-", "Iguala la sensibilidad al ángulo de ln q y ln τ (a9)")], 18):
    ws.cell(k_, 1, et_); c_ = ws.cell(k_, 2, v_); c_.font = AZUL; c_.fill = ENT; ws.cell(k_, 3, u_); ws.cell(k_, 4, n_)
    wb.defined_names[nm_] = DefinedName(nm_, attr_text=f"Parametros!$B${k_}")
ws.cell(16, 1, "Control: φ′e=0 con la composición del DDE (44/56)"); ws.cell(16, 2, OP2["pe0_dde"]); ws.cell(16, 4, "IFF-115-G-0001 Rev. B, Tabla 3-1: intrusivo y exoskarn 44 %, hornfels y mármol 56 % (repartido por igual)")

# ------------------------------------------------------------------ Ensayos
cab = ["Fuente", "Set", "Ensayo (id)", "Litología", "Tipo", "Uso", "Motivo", "D (m)", "Dmax (mm)", "e0", "σ′3 (MPa)", "σ′1 (MPa)",
       "σ′n (MPa)", "τ (MPa)", "φ′s (°)", "φ′s a e base (°)", "φ′s predicha por la curva (mismo σ′n y e0 del ensayo) (°)", "Δφ′s ensayo − curva (°)", "εa (%)", "Criterio de falla", "Página",
       "φ′1 del ensayo a su e0 (°)", "φ′1 del ensayo a e base (°)"]
ws = hoja("Ensayos", cab, [], [9, 22, 30, 11, 18, 26, 34, 7, 8, 7, 9, 9, 9, 9, 8, 9, 14, 9, 7, 34, 7, 11, 11])
for i, r in T27.iterrows():
    k = i + 2; tx = not str(r.ensayo).startswith("Corte")
    fila = [r.fuente_id, r.set_id, r.ensayo_id, r.litologia, r.ensayo, r.uso, r.motivo if isinstance(r.motivo, str) else "", r.D_m,
            pd.to_numeric(r.Dmax_mm, errors="coerce"), r.e0, r.s3 if tx else None, r.s1 if tx else None]
    for j, v in enumerate(fila, 1):
        ws.cell(k, j, None if (isinstance(v, float) and math.isnan(v)) else v)
    if tx:   # conversión de Leps 2.0 (Ecs. 13–16) en forma cerrada: D = σ1 − σ3, Sm = σ1 + σ3
        sp = f"((SQRT(4*m_*(m_-1)*(L{k}-K{k})^2+(L{k}+K{k})^2)-(L{k}+K{k}))/(2*(m_-1)*(L{k}-K{k})))"
        ws.cell(k, 13, f"=(L{k}+K{k})/2-(L{k}-K{k})/2*{sp}"); ws.cell(k, 14, f"=(L{k}-K{k})/2*SQRT(1-{sp}^2)")
    else:
        ws.cell(k, 13, r.sn); ws.cell(k, 14, r.tau)
    ws.cell(k, 15, f"=DEGREES(ATAN(N{k}/M{k}))")
    ws.cell(k, 16, f"=O{k}-omega*(e_base-J{k})")
    ws.cell(k, 17, f"=DEGREES(ATAN(TAN(RADIANS(pe0-omega*J{k}))*(M{k}/sref)^(m_-1)))")
    ws.cell(k, 18, f"=O{k}-Q{k}")
    ws.cell(k, 19, None if pd.isna(r.eps_a) else r.eps_a); ws.cell(k, 20, r.criterio); ws.cell(k, 21, None if pd.isna(r.pagina) else str(r.pagina))
    # φ′1: el punto del ensayo llevado a σ′n = 10pa a lo largo de la curvatura m (tan φ′1 = tan φ′s·(σ′n/10pa)^(1−m)); y luego a e base (Ec. 27)
    ws.cell(k, 22, f"=DEGREES(ATAN(TAN(RADIANS(O{k}))*(M{k}/sref)^(1-m_)))"); ws.cell(k, 23, f"=V{k}-omega*(e_base-J{k})")
control["ensayos_phi1"] = np.degrees(np.arctan(np.tan(np.radians(T27.phis)) * (T27.sn / 1.013) ** (1 - M))).tolist()
control["ensayos_sn"] = T27.sn.round(6).tolist(); control["ensayos_phis"] = T27.phis.round(5).tolist()
N_AJ = int((T27.uso == "Ajuste de la curva propuesta").sum())

# ------------------------------------------------------------------ Curva
SN = [0.01, 0.02, 0.03, 0.05, 0.07, 0.1, 0.15, 0.2, 0.3, 0.5, 0.7, 1, 1.5, 2, 3, 4, 4.6, 5, 6, 7, 7.86, 8, 9, 10, 11, 12, 13, 14, 15]
lo = lambda s: min(np.degrees(np.arctan(np.tan(np.radians(v["pe0_b"][0] - OM * EB)) * (s / 1.013) ** (v["m"] - 1))) for v in C2.values())
hi = lambda s: max(np.degrees(np.arctan(np.tan(np.radians(v["pe0_b"][1] - OM * EB)) * (s / 1.013) ** (v["m"] - 1))) for v in C2.values())
cab = ["σ′n (MPa)", "Zona", "τ curva (MPa)", "φ′s curva (°)", "φ′s banda inf. (°)", "φ′s banda sup. (°)", "τ banda inf. (MPa)", "τ banda sup. (MPa)",
       "τ e = 0.55", "τ e = 0.65", "e perfil A", "τ perfil A", "e perfil V", "τ perfil V", "φ′s Indraratna (°)", "φ′s Hawley y Cunning (°)",
       "τ Indraratna", "τ Hawley y Cunning", "τ 37° (MPa)", "φ′ sensibilidad DDE (°)"]
ws = hoja("Curva", cab, [], [9, 30] + [11] * 18)
for k, s in enumerate(SN, 2):
    zona = "con ensayos de Antamina" if s <= SNX else ("solo base mundial" if s <= SNW else "extrapolación")
    eV = float(PEM.eV_of_sv(s)); eA = float(PEM.e_A(s))
    vals = [s, zona, f"=TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)*A{k}^m_", f"=DEGREES(ATAN(C{k}/A{k}))",
            "=MIN(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$E${h}-omega*e_base))*(A{k}/sref)^(Parametros!$B${h}-1)))" for h in (12, 13, 14)) + ")",
            "=MAX(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$F${h}-omega*e_base))*(A{k}/sref)^(Parametros!$B${h}-1)))" for h in (12, 13, 14)) + ")",
            f"=A{k}*TAN(RADIANS(E{k}))", f"=A{k}*TAN(RADIANS(F{k}))", f"=TAN(RADIANS(pe0-omega*0.55))*sref^(1-m_)*A{k}^m_",
            f"=TAN(RADIANS(pe0-omega*0.65))*sref^(1-m_)*A{k}^m_",
            f"=IF(MAX(A{k},sv_piso)<=1,e_base-cc_bajo*LOG10(MAX(A{k},sv_piso)/sv_piso),e_base-cc_bajo*LOG10(1/sv_piso)-cc_alto*LOG10(A{k}))", f"=TAN(RADIANS(pe0-omega*K{k}))*sref^(1-m_)*A{k}^m_",
            None if math.isnan(eV) else round(eV, 4), None if math.isnan(eV) else f"=TAN(RADIANS(pe0-omega*M{k}))*sref^(1-m_)*A{k}^m_",
            f"=DEGREES(ATAN(TAN(RADIANS(33-omega*(e_base-0.6)))*(A{k}/sref)^(0.9-1)))", f"=DEGREES(ATAN(1.446*(A{k}*1000)^0.898/(A{k}*1000)))",
            f"=A{k}*TAN(RADIANS(O{k}))", f"=A{k}*TAN(RADIANS(P{k}))", f"=A{k}*TAN(RADIANS(37))", 33 if 5 <= s <= 10 else (30 if s > 10 else None)]
    for j, v in enumerate(vals, 1):
        ws.cell(k, j, v)
control["banda"] = [[float(lo(s)), float(hi(s))] for s in SN]; control["eA"] = [float(PEM.e_A(s)) for s in SN]
control["curva_phis"] = [float(np.degrees(np.arctan(np.tan(np.radians(PE0 - OM * EB)) * (s / 1.013) ** (M - 1)))) for s in SN]
nC = len(SN) + 1
ch = ScatterChart(); ch.title = "Ángulo de fricción secante a e base"; ch.x_axis.title = "σ′n (MPa)"; ch.y_axis.title = "φ′s (°)"
ch.x_axis.scaling.logBase = 10; ch.x_axis.scaling.min = 0.1; ch.x_axis.scaling.max = 15; ch.y_axis.scaling.min = 22; ch.y_axis.scaling.max = 56
X = Reference(ws, min_col=1, min_row=7, max_row=nC)
for col, nm in ((4, "Curva propuesta"), (5, "Banda inferior"), (6, "Banda superior"), (15, "Cota de Indraratna et al. (1993)"), (16, "Cota de Hawley y Cunning (2017)")):
    se = Series(Reference(ws, min_col=col, min_row=7, max_row=nC), X, title=nm); se.marker.symbol = "none"; ch.series.append(se)
we = wb["Ensayos"]
se = Series(Reference(we, min_col=16, min_row=2, max_row=1 + N_AJ), Reference(we, min_col=13, min_row=2, max_row=1 + N_AJ), title="Ensayos del ajuste, a e base")
se.marker.symbol = "circle"; se.graphicalProperties.line.noFill = True; ch.series.append(se)
ch.height, ch.width = 11, 20; ws.add_chart(ch, "V2")
ch2 = ScatterChart(); ch2.title = "τ–σ′n a e base"; ch2.x_axis.title = "σ′n (MPa)"; ch2.y_axis.title = "τ (MPa)"
X2 = Reference(ws, min_col=1, min_row=2, max_row=nC)
for col, nm in ((3, "Curva propuesta"), (7, "Banda inferior"), (8, "Banda superior"), (17, "Indraratna et al. (1993)"), (18, "Hawley y Cunning (2017)")):
    se = Series(Reference(ws, min_col=col, min_row=2, max_row=nC), X2, title=nm); se.marker.symbol = "none"; ch2.series.append(se)
ch2.height, ch2.width = 11, 20; ws.add_chart(ch2, "V26")

# ------------------------------------------------------------------ phi1_vs_e (φ′1 contra e, como Cylwik et al. 2025)
ws = hoja("phi1_vs_e", ["Set (ensayos del ajuste)", "Litología", "Ensayos", "e0 media", "φ′1 medio a su e0 (°)", "", "e", "Curva propuesta (°)",
                        "Leps 2.0, clase baja (47°)", "Leps 2.0, clase moderada (50°)", "Leps 2.0, clase alta (53°)"], [], [26, 12, 8, 9, 12, 3, 6, 12, 12, 12, 12])
SETS = T27[T27.uso == "Ajuste de la curva propuesta"].drop_duplicates(["fuente_id", "set_id"])[["fuente_id", "set_id", "litologia"]]
for k, (_, r) in enumerate(SETS.iterrows(), 2):
    crit = f'Ensayos!$A$2:$A$500,"{r.fuente_id}",Ensayos!$B$2:$B$500,"{r.set_id}",Ensayos!$F$2:$F$500,"Ajuste de la curva propuesta"'
    ws.cell(k, 1, f"{r.fuente_id} | {r.set_id}"); ws.cell(k, 2, r.litologia); ws.cell(k, 3, f"=COUNTIFS({crit})")
    ws.cell(k, 4, f"=AVERAGEIFS(Ensayos!$J$2:$J$500,{crit})"); ws.cell(k, 5, f"=AVERAGEIFS(Ensayos!$V$2:$V$500,{crit})")
nS = len(SETS) + 1
for k, e_ in enumerate([round(0.05 * i, 2) for i in range(17)], 2):
    ws.cell(k, 7, e_); ws.cell(k, 8, f"=pe0-omega*G{k}")
    for j, p_ in enumerate((47, 50, 53), 9):
        ws.cell(k, j, f"={p_}-omega*G{k}")
ws.cell(nS + 2, 1, "φ′1 de cada ensayo: Ensayos!V (a su e0). Cada punto de la figura de Cylwik et al. (2025) es un set; aquí, el promedio de los ensayos del set.")
ch3 = ScatterChart(); ch3.title = "φ′1 contra relación de vacíos"; ch3.x_axis.title = "e (relación de vacíos)"; ch3.y_axis.title = "φ′1 (°)"
ch3.x_axis.scaling.min = 0; ch3.x_axis.scaling.max = 0.8; ch3.y_axis.scaling.min = 20; ch3.y_axis.scaling.max = 55
we = wb["Ensayos"]
se = Series(Reference(we, min_col=22, min_row=2, max_row=1 + N_AJ), Reference(we, min_col=10, min_row=2, max_row=1 + N_AJ), title="Ensayos de Antamina")
se.marker.symbol = "circle"; se.graphicalProperties.line.noFill = True; ch3.series.append(se)
se = Series(Reference(ws, min_col=5, min_row=2, max_row=nS), Reference(ws, min_col=4, min_row=2, max_row=nS), title="Promedio por set")
se.marker.symbol = "diamond"; se.graphicalProperties.line.noFill = True; ch3.series.append(se)
for col, nm in ((8, "Curva propuesta"), (9, "Clase baja"), (10, "Clase moderada"), (11, "Clase alta")):
    se = Series(Reference(ws, min_col=col, min_row=2, max_row=18), Reference(ws, min_col=7, min_row=2, max_row=18), title=nm); se.marker.symbol = "none"; ch3.series.append(se)
ch3.height, ch3.width = 11, 18; ws.add_chart(ch3, "M2")

# ------------------------------------------------------------------ Hipotesis_m
SNT = [1, 2, 3, 4.6, 8, 15]
ws = hoja("Hipotesis_m", ["Hipótesis", "m", "φ′e=0 (°)"] + [f"φ′s a {s} MPa (°)" for s in SNT], [], [28, 8, 10] + [12] * len(SNT))
for k, (key, et) in enumerate(HIP, 2):
    ws.cell(k, 1, et); ws.cell(k, 2, f"=Parametros!B{k + 10}"); ws.cell(k, 3, f"=Parametros!C{k + 10}")
    for j, s in enumerate(SNT, 4):
        ws.cell(k, j, f"=DEGREES(ATAN(TAN(RADIANS(C{k}-omega*e_base))*({s}/sref)^(B{k}-1)))")

# ------------------------------------------------------------------ Base_mundial
W = datos(); W = W[~W.cens] if "cens" in W else W
ws = hoja("Base_mundial", ["Set", "e", "σ′n (MPa)", "τ (MPa)", "φ′s ensayo (°)", "φ′s curva a su e (°)", "Δφ′s (°)"], [], [40, 8, 10, 10, 12, 14, 10])
setcol = next((c for c in ("set_id", "set", "Set", "set_key") if c in W.columns), None)
for i, (_, r) in enumerate(W.iterrows()):
    k = i + 2
    ws.cell(k, 1, str(r[setcol]) if setcol else ""); ws.cell(k, 2, float(r.e)); ws.cell(k, 3, float(r.sn)); ws.cell(k, 4, float(r.tau))
    ws.cell(k, 5, f"=DEGREES(ATAN(D{k}/C{k}))"); ws.cell(k, 6, f"=DEGREES(ATAN(TAN(RADIANS(pe0-omega*B{k}))*(C{k}/sref)^(m_-1)))"); ws.cell(k, 7, f"=E{k}-F{k}")
nW = len(W) + 1
ws["I1"] = "Rango de σ′n"; ws["J1"] = "Ensayos"; ws["K1"] = "Bajo la curva"; ws["L1"] = "Δφ′s mediana (°)"
for c in ws["I1:L1"][0]:
    c.font = NEG; c.fill = CAB
for k, (a, b) in enumerate(((2, 4.6), (4.6, 7.9), (3, 7.9)), 2):
    ws.cell(k, 9, f"{a}–{b} MPa")
    ws.cell(k, 10, f'=COUNTIFS(C2:C{nW},">{a}",C2:C{nW},"<={b}")'); ws.cell(k, 11, f'=COUNTIFS(C2:C{nW},">{a}",C2:C{nW},"<={b}",G2:G{nW},"<0")')
    ws[f"L{k}"] = ArrayFormula(f"L{k}", f'=MEDIAN(IF((C2:C{nW}>{a})*(C2:C{nW}<={b}),G2:G{nW}))')
A25 = J("a25_alta_tension.json")["ensayos"]; control["mundial_46_79"] = A25["Base mundial (Cylwik et al. 2025) | 4.6–7.9 MPa"]

def g_expr(x, a, s3):
    """σ3(x) − σ′3 para la curva τ = a·x^m (texto de fórmula de Excel)."""
    tau = f"({a}*({x})^m_)"; t = f"(m_*{a}*({x})^(m_-1))"
    return f"(({x})+{tau}*{t}-{tau}*SQRT(1+{t}^2)-({s3}))"


def newton(ws, fila, col0, a, s3, n=10):
    """Escribe n columnas de Newton desde col0 (x0 = 2·σ′3); devuelve la letra de la última."""
    prev = f"2*{s3}"
    for i in range(n):
        c = col0 + i; x = prev
        h = f"(({x})*1E-6)"
        ws.cell(fila, c, f"={x}-{g_expr(x, a, s3)}*{h}/({g_expr(x + '+' + h, a, s3)}-{g_expr(x, a, s3)})")
        prev = f"{get_column_letter(c)}{fila}"
    return get_column_letter(col0 + n - 1)


# ------------------------------------------------------------------ Ajuste (función objetivo de a12: todos los ensayos con igual peso)
AJ = pd.read_csv(AQ / "ensayos_ajuste.csv", low_memory=False)
cab = ["Fuente", "Set", "Tipo", "e0", "σ′3 (MPa)", "σ′1 (MPa)", "q obs (MPa)", "σ′n obs, DS (MPa)", "τ obs, DS (MPa)", "a = tan(φ′e0 − ω·e0)·(10pa)^(1−m)"] + \
      [f"Newton {i + 1}: σ′n (MPa)" for i in range(10)] + ["τ (MPa)", "tan φ′t", "R (MPa)", "q pred (MPa)", "Residuo", "Residuo²"]
ws = hoja("Ajuste", cab, [], [9, 22, 5, 7, 8, 8, 8, 9, 9, 12] + [9] * 10 + [8, 8, 8, 9, 9, 9])
for i, r in AJ.iterrows():
    k = i + 2; tx = r.tipo == "TX"
    for j, v in enumerate([r.fuente_id, r.set_id, r.tipo, r.e0_adoptado, r.s3 if tx else None, r.s1 if tx else None,
                           f"=F{k}-E{k}" if tx else None, None if tx else r.sn_ds, None if tx else r.tau_ds,
                           f"=TAN(RADIANS(pe0-omega*D{k}))*sref^(1-m_)"], 1):
        ws.cell(k, j, v)
    if tx:
        ul = newton(ws, k, 11, f"J{k}", f"E{k}")
        ws.cell(k, 21, f"=J{k}*{ul}{k}^m_"); ws.cell(k, 22, f"=m_*J{k}*{ul}{k}^(m_-1)"); ws.cell(k, 23, f"=U{k}*SQRT(1+V{k}^2)")
        ws.cell(k, 24, f"=2*W{k}"); ws.cell(k, 25, f"=LN(X{k})-LN(G{k})")
    else:
        ws.cell(k, 21, f"=J{k}*H{k}^m_"); ws.cell(k, 25, f"=c_ds*(LN(U{k})-LN(I{k}))")
    ws.cell(k, 26, f"=Y{k}^2")
nA = len(AJ) + 1
ws.cell(nA + 2, 25, "SSE"); ws.cell(nA + 2, 26, f"=SUM(Z2:Z{nA})").font = NEG
ws.cell(nA + 3, 1, "φ′e0 de la curva propuesta = el valor de Parametros!B2 que minimiza SSE (Datos > Solver: minimizar Z" + str(nA + 2) + " cambiando Parametros!$B$2).")
ws.cell(nA + 4, 1, "Tangencia: σ′3 = σ′n + τ·tan φ′t − τ·√(1 + tan² φ′t) (círculo de Mohr tangente a la curva), resuelta con 10 pasos de Newton; q = 2R, R = τ·√(1 + tan² φ′t).")
sys.path.insert(0, str(AQ.parent / "analisis")); from modelos import sigma1, tau_leps2  # noqa: E402
def _res(pe0):
    e_ = AJ.e0_adoptado.values; t_ = (AJ.tipo == "TX").values; rr = np.empty(len(AJ))
    rr[t_] = np.log(sigma1(tau_leps2, AJ.s3.values[t_], pe0 - OM * e_[t_], M)[0] - AJ.s3.values[t_]) - np.log((AJ.s1 - AJ.s3).values[t_])
    rr[~t_] = C_DS * (np.log(tau_leps2(AJ.sn_ds.values[~t_], pe0 - OM * e_[~t_], M)) - np.log(AJ.tau_ds.values[~t_]))
    return rr
control["ajuste_resid"] = _res(PE0).tolist(); control["sse"] = {str(round(p, 2)): float((_res(p) ** 2).sum()) for p in (PE0 - 0.05, PE0, PE0 + 0.05)}

# ------------------------------------------------------------------ FLAC_tangente
def tangente(s3):
    a_ = math.tan(math.radians(PE0 - OM * EB)) * 1.013 ** (1 - M); lo_, hi_ = math.log(s3), math.log(40 * s3)
    for _ in range(80):
        mid = 0.5 * (lo_ + hi_); sn_ = math.exp(mid); t_ = a_ * sn_ ** M; b_ = 0.5 * math.atan(M * t_ / sn_) + math.pi / 4
        lo_, hi_ = (lo_, mid) if sn_ - t_ / math.tan(b_) - s3 > 0 else (mid, hi_)
    sn_ = math.exp(0.5 * (lo_ + hi_)); return sn_, a_ * sn_ ** M


ws = hoja("FLAC_tangente", ["σ′3 (MPa)", "σ′n tangencia (MPa)", "τ (MPa)", "φ′t (°)", "c′t (kPa)", "σ′1 − σ′3 (MPa)", "a"] + [f"Newton {i + 1}" for i in range(10)],
          [], [10, 16, 10, 10, 10, 14, 10] + [9] * 10)
control["flac_sn"] = []
for k, s3 in enumerate((0.1, 0.5, 1.0, 2.0, 5.0), 2):
    sn_, t_ = tangente(s3); control["flac_sn"].append(sn_)
    ws.cell(k, 1, s3); ws.cell(k, 7, "=TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)")
    ul = newton(ws, k, 8, f"G{k}", f"A{k}"); ws.cell(k, 2, f"={ul}{k}")
    ws.cell(k, 3, "=TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)*B%d^m_" % k)
    ws.cell(k, 4, f"=DEGREES(ATAN(m_*C{k}/B{k}))"); ws.cell(k, 5, f"=(1-m_)*C{k}*1000"); ws.cell(k, 6, f"=2*C{k}/COS(RADIANS(D{k}))")
ws["A8"] = "σ′n de tangencia: σ′3 = σ′n + τ·tan φ′t − τ·√(1 + tan² φ′t), resuelta con 10 pasos de Newton (columnas H a Q); tan φ′t = m·τ/σ′n, c′t = (1 − m)·τ."

# ------------------------------------------------------------------ Relacion_vacios
EV = [("ROM", "SDA, INF-115-R-0005 (Stantec 2025), valor adoptado", 20.0, 0.029), ("ROM", "DDE, INF-115-R-0004, Tabla 9-3", 19.6, 0.029),
      ("ROM", "Golder 2010, citado en el DDE (p. 59)", 20.75, 0.029), ("Chancado", "Golder 2010, citado en el DDE", 21.66, 0.028),
      ("Chancado", "SDA, valor adoptado", 22.0, 0.028), ("Enrocado 3B", "DDE y SDA", 22.0, 0.029)]
ws = hoja("Relacion_vacios", ["Material", "Fuente", "γ (kN/m³)", "w", "Gs", "γw", "e"], [], [12, 50, 10, 8, 7, 7, 8])
for k, (mt, fu, g, w) in enumerate(EV, 2):
    for j, v in enumerate([mt, fu, g, w, 2.75, 9.81, f"=E{k}*F{k}*(1+D{k})/C{k}-1"], 1):
        ws.cell(k, j, v)

# ------------------------------------------------------------------ Perfiles_e, Edometros, Tamano_particula
ws = hoja("Perfiles_e", ["Profundidad (m)", "σ′v perfil A (MPa)", "e perfil A", "σ′v perfil V (MPa)", "e perfil V"], [], [14, 16, 10, 16, 10])
for z, sa, ea, sv, ev in zip(PE["z"], PE["sv_A"], PE["e_A"], PE["sv_V"], PE["e_V"]):
    ws.append([z, sa, ea, sv, ev])
ws["G1"] = "Valenzuela et al. (2008), curva de diseño (lectura de la figura): altura (m)"; ws["H1"] = "γd (t/m³)"
for i_, (h_, g_) in enumerate(zip(PEM.HV, PEM.GDV), 2):
    ws.cell(i_, 7, float(h_)); ws.cell(i_, 8, float(g_))
ws = hoja("Edometros", ["Ensayo", "Diámetro (mm)", "e0", "Cc1 (0.2–0.8 MPa)", "Cc2 (1.2–4.8 MPa)", "e a 300 m", "e a 700 m"], [], [40, 12, 8, 14, 14, 10, 10])
for r in EDO:
    p_ = {str(k): v for k, v in (r["perfil"] or {}).items()}
    ws.append([r["ensayo"], r["diametro_mm"], r["e0"], r["Cc1"], r["Cc2"], p_.get("300"), p_.get("700")])
ws = hoja("Tamano_particula", ["D lab (mm)", "m_w", "D campo (mm)", "λ = (Dlab/Dcampo)^(3/mw)", "Factor sobre τ = λ^(1−m)", "φ′1 (°)", "Diferencia (°)"],
          [], [10, 8, 12, 22, 22, 10, 12])
k = 2
for dl in (178, 102, 25):
    for mw in (4.0, 3 / 0.45, 10.0):
        for j, v in enumerate([dl, round(mw, 4), 1000, f"=(A{k}/C{k})^(3/B{k})", f"=D{k}^(1-m_)", f"=DEGREES(ATAN(E{k}*TAN(RADIANS(pe0-omega*e_base))))",
                               f"=F{k}-(pe0-omega*e_base)"], 1):
            ws.cell(k, j, v)
        k += 1
wb.save(OUT)
json.dump(control, open(AQ / "a28_valores_control.json", "w", encoding="utf-8"))
print("escrito", OUT, "| ensayos", len(T27), "| base mundial", len(W), "| curva", len(SN), "puntos")
