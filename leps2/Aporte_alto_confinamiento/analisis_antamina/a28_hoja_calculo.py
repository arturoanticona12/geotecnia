"""Hoja de cálculo del memo (comentario WA 2026-10-12: «genera una hoja de cálculo que incluya todos los datos, cálculos y curvas»).
Escribe eor_antamina/curva_leps2_antamina_calculos.xlsx con fórmulas vivas de Excel ligadas a la hoja Parametros (celdas con
nombre): ensayos (120, con su uso), curva propuesta y alternativas de e, banda, cotas, hipótesis de m, base mundial (741),
tangentes para FLAC, relación de vacíos, perfiles e(z), edómetros y tamaño de partícula, con gráficos. Escribe también
a28_valores_control.json con los valores que deben dar las fórmulas (para verificarlas al recalcular en Excel)."""
import json, math, sys
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
sys.stdout.reconfigure(encoding="utf-8")
AQ = Path(__file__).parent; sys.path.insert(0, str(AQ.parent / "analisis")); sys.path.insert(0, str(AQ))
from lepsp_r6 import datos  # noqa: E402
import a13_perfil_e as PEM  # noqa: E402
J = lambda n: json.load(open(AQ / n, encoding="utf-8"))
C_DS = J("a9_leps2_antamina.json")["datos"]["c_ds"]; A12 = J("a12_leps2_ponderada.json"); C2 = A12["curvas"]; OP2 = C2["Curva propuesta (mediana mundial)"]; PE = J("a13_perfil_e.json")
R19 = J("a19_m_covariables.json"); EDO = J("a17_edometros.json"); A16 = J("a16_resultados.json")
# revisión 10: curva simple (a33): m = 0.85, nivel = media de φ′1 + 20·e0 de los sets; m sin sesgo en la base mundial (a31)
A33 = J("a33_curva_propuesta.json"); NIV = A33["nivel"]; A31 = J("a31_m_defendible.json"); S34 = pd.read_csv(AQ / "a34_phi1_sets.csv")
_ml = S34.groupby("litologia").phi_e0.mean()
OP2 = dict(pe0=A33["parametros"]["pe0"], m=A33["parametros"]["m_a"], pe0_dde=float(0.28 * _ml["Mármol"] + 0.28 * _ml["Hornfels"] + 0.44 * _ml["Intrusivo"]),
           pe0_le2=float(0.674 * _ml["Mármol"] + 0.155 * _ml["Hornfels"] + 0.170 * _ml["Intrusivo"]) / 0.999)   # MR2024 ≥2035; caliza y margas (29.6 %) asimiladas al mármol
from tope import TOPE   # tope del ángulo secante (tope.py)
PE0, M, OM, EB, PA = OP2["pe0"], OP2["m"], 20.0, 0.45, 0.1013
SNW = R19["base_sn_max"]
T27 = pd.read_csv(AQ / "a27_ensayos_todos.csv")
SNX = float(T27[(T27.uso == "Ajuste de la curva propuesta") & ~T27.ensayo.str.startswith("Corte")].sn.max())
ORD = {"Ajuste de la curva propuesta": 0, "Contraste": 2, "No usado": 3}
T27 = T27.assign(_o=T27.uso.map(ORD)).sort_values(["_o"], kind="stable").drop(columns="_o").reset_index(drop=True)
OUT = AQ.parent / "eor_antamina" / "curva_leps2_antamina_calculos.xlsx"
AZUL = Font(color="1F4E9E"); NEG = Font(bold=True); CAB = PatternFill("solid", fgColor="E8E6F2"); ENT = PatternFill("solid", fgColor="EAF1FB")
wb = Workbook(); control = {}


def hoja(nombre, cab, filas, anchos=None):
    ws = wb.create_sheet(nombre)
    ws.append(cab)
    for c in ws[1]:
        c.font = NEG; c.fill = CAB; c.alignment = Alignment(wrap_text=True, vertical="top")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos or [14] * len(cab), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


# ------------------------------------------------------------------ Leame
ws = wb.active; ws.title = "Leame"
for t in ["Curva de resistencia Leps 2.0 del desmonte ROM de Antamina: datos, cálculos y curvas (memo de Stantec, 12 de octubre de 2026)",
          "",
          "Las celdas en azul de la hoja Parametros son las entradas; todo lo demás se calcula con fórmulas a partir de ellas (nombres definidos: pe0, omega, m_, e_base, sref, tope).",
          "Leps 2.0 (Cylwik et al. 2025, Ec. 26): τ = tan(φ′e=0 − ω·e)·(10pa)^(1−m)·σ′n^m; φ′s = atan(τ/σ′n). La curva se aplica con un tope de " + f"{TOPE:.0f}° en φ′s (τ ≤ σ′n·tan {TOPE:.0f}°).",
          "Unidades: esfuerzos en MPa; ángulos en grados.",
          "",
          "Hojas:",
          "  Parametros: parámetros de la curva propuesta, hipótesis de curvatura y controles con la composición del DDE y la de LE2.",
          "  Ensayos: los 120 ensayos de resistencia con su uso (ajuste, contraste, no usado), esfuerzos de falla, conversión de Leps 2.0 (fórmula) y corrección a e base (Ec. 27).",
          "  Figuras: las figuras del memo como gráficos de Excel, con la misma numeración (datos y fórmulas en Datos_figuras).",
          "  Curva: curva propuesta a e base, 0.55, 0.65 y perfiles A y V; tope de " + f"{TOPE:.0f}" + "° en todas las curvas de Leps 2.0; banda de incertidumbre (sin tope); cotas de Indraratna et al. (1993) y Hawley y Cunning (2017); recta de 37°; sensibilidad del DDE; gráficos.",
          "  phi1_vs_e: φ′1 de cada ensayo y por set contra e, φ′e=0 por set y su media (el nivel de la curva), con la recta de la curva propuesta, las envolventes de Leps (1970) y la cota de Indraratna et al. (1993).",
          "  Comparacion_envolventes: curva propuesta, Leps (1970) e Indraratna et al. (1993) frente a los ensayos (sección «Por qué esta curva» del memo).",
          "  Hipotesis_m: φ′s según la hipótesis de curvatura (m del intervalo al 90 % y P10–P90 entre materiales de la base mundial, y m = 1 de Antamina sola como sensibilidad).",
          "  Base_mundial: los ensayos de la base de Leps 2.0 (sin los censurados, como en el memo) frente a la curva evaluada a su propia e (Δφ′s) y conteos por rango de esfuerzo.",
          "  FLAC_tangente: Mohr-Coulomb tangente a la curva según σ′3 (para verificar la implementación en FLAC2D).",
          "  Relacion_vacios: e de colocación calculada con los pesos unitarios de los estudios de Antamina.",
          "  Perfiles_e: e y σ′v con la profundidad (perfiles A y V).",
          "  Edometros: edómetros de desmonte de Antamina y e resultante con la ley del perfil A.",
          "  Tamano_particula: sensibilidad al tamaño (Frossard et al. 2012).",
          "",
          "Fuentes: base de ensayos de Antamina verificada contra los PDF (datos_antamina/base, verificacion_ocr); base mundial de Cylwik et al. (2025);",
          "Informe final de estudio del DDE (Stantec 2026, IFF-115-G-0001 Rev. B) para la composición de control y la sensibilidad de 33°/30°."]:
    ws.append([t])
ws.column_dimensions["A"].width = 150

# ------------------------------------------------------------------ Parametros
ws = hoja("Parametros", ["Parámetro", "Valor", "Unidad", "Nota"], [], [34, 12, 10, 90])
PAR = [("pe0", "φ′e=0, curva propuesta", PE0, "°", "Media de φ′1 + 20·e0 de los 13 sets con m = 0.85, todos con igual peso (a33; ver phi1_vs_e)"),
       ("omega", "ω", OM, "°", "Leps 2.0, Ec. 25"),
       ("m_", "m (curvatura)", M, "-", "Sin sesgo en los ensayos mundiales de mayor esfuerzo (a31; IC 90 % 0.834–0.862)"),
       ("e_base", "e base de colocación", EB, "-", "Desmonte ROM (sección 7 del memo)"),
       ("pa", "pa", PA, "MPa", "Presión atmosférica"),
       ("sref", "10·pa", None, "MPa", "Esfuerzo de referencia")]
for k, (nm, et, v, u, n) in enumerate(PAR, 2):
    ws.cell(k, 1, et); ws.cell(k, 3, u); ws.cell(k, 4, n)
    c = ws.cell(k, 2, v if v is not None else "=10*pa"); c.font = AZUL if v is not None else Font(); c.fill = ENT if v is not None else PatternFill()
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"Parametros!$B${k}")
ws.cell(8, 1, "φ′1 a e base"); ws.cell(8, 2, "=pe0-omega*e_base"); ws.cell(8, 3, "°")
ws.cell(9, 1, "Tope de φ′s"); c = ws.cell(9, 2, TOPE); c.font = AZUL; c.fill = ENT; ws.cell(9, 3, "°"); ws.cell(9, 4, "Ángulo secante máximo de la curva (bajo confinamiento; sección 7 del memo)")
wb.defined_names["tope"] = DefinedName("tope", attr_text="Parametros!$B$9")
control["phi1"] = PE0 - OM * EB
ws.cell(10, 1, "Hipótesis de curvatura").font = NEG
ws.cell(11, 1, "Hipótesis"); ws.cell(11, 2, "m"); ws.cell(11, 3, "φ′e=0 (°)"); ws.cell(11, 4, "Nota"); ws.cell(11, 5, "φ′e=0 P5 (°)"); ws.cell(11, 6, "φ′e=0 P95 (°)")
HIP = [("m_p90", "m P90 entre materiales"), ("m0862", "m alto (IC 90 %)"), ("m085", "Curva propuesta"), ("m0834", "m bajo (IC 90 %)"), ("m_p10", "m P10 entre materiales"),
       ("lineal", "Antamina sola (m = 1), sensibilidad")]
HB = (12, 13, 14, 15, 16)
for k, (key, et) in enumerate(HIP, 12):          # filas 12–16: banda (mínimo de los P5 y máximo de los P95); fila 17: sensibilidad, fuera de la banda
    ws.cell(k, 1, et); ws.cell(k, 2, A33["hipotesis"][key]["m_a"]); ws.cell(k, 3, NIV[key]["pe0"]); ws.cell(k, 4, "Nivel reajustado a los ensayos de Antamina con esa m (a33)")
    ws.cell(k, 5, NIV[key]["p5"]); ws.cell(k, 6, NIV[key]["p95"])
for k_, (nm_, et_, v_, u_, n_) in enumerate([("cc_bajo", "Cc perfil A, 0.2–0.8 MPa", PE["Cc1"], "-", "Edómetro de mármol suelto (a13)"),
                                                ("cc_alto", "Cc perfil A, 1.2–4.8 MPa", PE["Cc2"], "-", "Edómetro de mármol suelto (a13)"),
                                                ("sv_piso", "σ′v mínima de la ley del perfil A", 0.05, "MPa", "a13"),
                                                ("c_ds", "Peso del residuo del corte directo", C_DS, "-", "Iguala la sensibilidad al ángulo de ln q y ln τ (a9)")], 19):
    ws.cell(k_, 1, et_); c_ = ws.cell(k_, 2, v_); c_.font = AZUL; c_.fill = ENT; ws.cell(k_, 3, u_); ws.cell(k_, 4, n_)
    wb.defined_names[nm_] = DefinedName(nm_, attr_text=f"Parametros!$B${k_}")
ws.cell(24, 1, "Control: φ′e=0 con la composición de LE2"); ws.cell(24, 2, OP2["pe0_le2"]); ws.cell(24, 3, "°"); ws.cell(24, 4, "MR2024, períodos ≥2035: mármol 37.8 %, hornfels 15.5 %, intrusivo 17.0 %, caliza y margas 29.6 % (sin ensayos; asimiladas al mármol)")
ws.cell(23, 1, "Control: φ′e=0 con la composición del DDE (44/56)"); ws.cell(23, 2, OP2["pe0_dde"]); ws.cell(23, 3, "°"); ws.cell(23, 4, "IFF-115-G-0001 Rev. B, Tabla 3-1: intrusivo y exoskarn 44 %, hornfels y mármol 56 % (repartido por igual)")

# ------------------------------------------------------------------ Ensayos
cab = ["Fuente", "Set", "Ensayo (id)", "Litología", "Tipo", "Uso", "Motivo", "D (m)", "Dmax (mm)", "e0", "σ′3 (MPa)", "σ′1 (MPa)",
       "σ′n (MPa)", "τ (MPa)", "φ′s (°)", "φ′s a e base (°)", "φ′s predicha por la curva, con tope (mismo σ′n y e0 del ensayo) (°)", "Δφ′s ensayo − curva (°)", "εa (%)", "Criterio de falla", "Página",
       "φ′1 del ensayo a su e0 (°)", "φ′1 del ensayo a e base (°)"]
ws = hoja("Ensayos", cab, [], [9, 22, 30, 11, 18, 26, 34, 7, 8, 7, 9, 9, 9, 9, 8, 9, 14, 9, 7, 34, 7, 11, 11])
for i, r in T27.iterrows():
    k = i + 2; tx = not str(r.ensayo).startswith("Corte")
    fila = [r.fuente_id, r.set_id, r.ensayo_id, r.litologia, r.ensayo, r.uso, r.motivo if isinstance(r.motivo, str) else "", r.D_m,
            pd.to_numeric(r.Dmax_mm, errors="coerce"), r.e0, r.s3 if tx else None, r.s1 if tx else None]
    for j, v in enumerate(fila, 1):
        ws.cell(k, j, None if (isinstance(v, float) and math.isnan(v)) else v)
    if tx:   # conversión de Leps 2.0 (Ecs. 13–16) en forma cerrada: D = σ1 − σ3, Sm = σ1 + σ3
        sp = f"((SQRT(4*m_*(m_-1)*(L{k}-K{k})^2+(L{k}+K{k})^2)-(L{k}+K{k}))/(2*(m_-1)*(L{k}-K{k})))"
        ws.cell(k, 13, f"=(L{k}+K{k})/2-(L{k}-K{k})/2*{sp}"); ws.cell(k, 14, f"=(L{k}-K{k})/2*SQRT(1-{sp}^2)")
    else:
        ws.cell(k, 13, r.sn); ws.cell(k, 14, r.tau)
    ws.cell(k, 15, f"=DEGREES(ATAN(N{k}/M{k}))")
    ws.cell(k, 16, f"=O{k}-omega*(e_base-J{k})")
    ws.cell(k, 17, f"=MIN(tope,DEGREES(ATAN(TAN(RADIANS(pe0-omega*J{k}))*(M{k}/sref)^(m_-1))))")   # curva entregada: con el tope
    ws.cell(k, 18, f"=O{k}-Q{k}")
    ws.cell(k, 19, None if pd.isna(r.eps_a) else r.eps_a); ws.cell(k, 20, r.criterio); ws.cell(k, 21, None if pd.isna(r.pagina) else str(r.pagina))
    # φ′1: el punto del ensayo llevado a σ′n = 10pa a lo largo de la curvatura m (tan φ′1 = tan φ′s·(σ′n/10pa)^(1−m)); y luego a e base (Ec. 27)
    ws.cell(k, 22, f"=DEGREES(ATAN(TAN(RADIANS(O{k}))*(M{k}/sref)^(1-m_)))"); ws.cell(k, 23, f"=V{k}-omega*(e_base-J{k})")
control["ensayos_phi1"] = np.degrees(np.arctan(np.tan(np.radians(T27.phis)) * (T27.sn / 1.013) ** (1 - M))).tolist()
control["ensayos_sn"] = T27.sn.round(6).tolist(); control["ensayos_phis"] = T27.phis.round(5).tolist()
N_AJ = int((T27.uso == "Ajuste de la curva propuesta").sum())

# ------------------------------------------------------------------ Curva
SN = [0.01, 0.02, 0.03, 0.05, 0.07, 0.1, 0.15, 0.2, 0.3, 0.5, 0.7, 1, 1.5, 2, 3, 4, 4.6, 5, 6, 7, 7.86, 8, 9, 10, 11, 12, 13, 14, 15]
_hb = [A33["hipotesis"][k_]["m_a"] for k_ in ("m0862", "m085", "m0834")]
lo = lambda s: min(np.degrees(np.arctan(np.tan(np.radians(NIV[k_]["p5"] - OM * EB)) * (s / 1.013) ** (A33["hipotesis"][k_]["m_a"] - 1))) for k_ in ("m_p90", "m0862", "m085", "m0834", "m_p10"))
hi = lambda s: max(np.degrees(np.arctan(np.tan(np.radians(NIV[k_]["p95"] - OM * EB)) * (s / 1.013) ** (A33["hipotesis"][k_]["m_a"] - 1))) for k_ in ("m_p90", "m0862", "m085", "m0834", "m_p10"))
cab = ["σ′n (MPa)", "Zona", "τ curva, con tope (MPa)", "φ′s curva, con tope (°)", "φ′s banda inf. (°)", "φ′s banda sup. (°)", "τ banda inf. (MPa)", "τ banda sup. (MPa)",
       "τ e = 0.55", "τ e = 0.65", "e perfil A", "τ perfil A", "e perfil V", "τ perfil V", "φ′s Indraratna (°)", "φ′s Hawley y Cunning (°)",
       "τ Indraratna", "τ Hawley y Cunning", "τ 37° (MPa)", "φ′ sensibilidad DDE (°)"]
ws = hoja("Curva", cab, [], [9, 30] + [11] * 18)
for k, s in enumerate(SN, 2):
    zona = "con ensayos de Antamina" if s <= SNX else ("solo base mundial" if s <= SNW else "extrapolación")
    eV = float(PEM.eV_of_sv(s)); eA = float(PEM.e_A(s))
    vals = [s, zona, f"=MIN(TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)*A{k}^m_,A{k}*TAN(RADIANS(tope)))", f"=DEGREES(ATAN(C{k}/A{k}))",
            "=MIN(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$E${h}-omega*e_base))*(A{k}/sref)^(Parametros!$B${h}-1)))" for h in HB) + ")",
            "=MAX(" + ",".join(f"DEGREES(ATAN(TAN(RADIANS(Parametros!$F${h}-omega*e_base))*(A{k}/sref)^(Parametros!$B${h}-1)))" for h in HB) + ")",
            f"=A{k}*TAN(RADIANS(E{k}))", f"=A{k}*TAN(RADIANS(F{k}))", f"=MIN(TAN(RADIANS(pe0-omega*0.55))*sref^(1-m_)*A{k}^m_,A{k}*TAN(RADIANS(tope)))",
            f"=MIN(TAN(RADIANS(pe0-omega*0.65))*sref^(1-m_)*A{k}^m_,A{k}*TAN(RADIANS(tope)))",
            f"=IF(MAX(A{k},sv_piso)<=1,e_base-cc_bajo*LOG10(MAX(A{k},sv_piso)/sv_piso),e_base-cc_bajo*LOG10(1/sv_piso)-cc_alto*LOG10(A{k}))", f"=MIN(TAN(RADIANS(pe0-omega*K{k}))*sref^(1-m_)*A{k}^m_,A{k}*TAN(RADIANS(tope)))",
            None if math.isnan(eV) else round(eV, 4), None if math.isnan(eV) else f"=MIN(TAN(RADIANS(pe0-omega*M{k}))*sref^(1-m_)*A{k}^m_,A{k}*TAN(RADIANS(tope)))",
            f"=DEGREES(ATAN(TAN(RADIANS(33))*(A{k}/sref)^(0.9-1)))", f"=DEGREES(ATAN(1.446*(A{k}*1000)^0.898/(A{k}*1000)))",
            f"=A{k}*TAN(RADIANS(O{k}))", f"=A{k}*TAN(RADIANS(P{k}))", f"=A{k}*TAN(RADIANS(37))", 33 if 5 <= s <= 10 else (30 if s > 10 else None)]
    for j, v in enumerate(vals, 1):
        ws.cell(k, j, v)
control["banda"] = [[float(lo(s)), float(hi(s))] for s in SN]; control["eA"] = [float(PEM.e_A(s)) for s in SN]
control["curva_phis"] = [min(TOPE, float(np.degrees(np.arctan(np.tan(np.radians(PE0 - OM * EB)) * (s / 1.013) ** (M - 1))))) for s in SN]
nC = len(SN) + 1

# ------------------------------------------------------------------ phi1_vs_e (φ′1 contra e, como Cylwik et al. 2025)
ws = hoja("phi1_vs_e", ["Set (ensayos del ajuste)", "Litología", "Ensayos", "e0 media", "φ′1 medio a su e0 (°)", "", "e", "Curva propuesta (°)",
                        "Leps (1970), superior (46°)", "Leps (1970), media (41°)", "Leps (1970), inferior (36°)", "Indraratna et al. (1993) (33°)", "",
                        "φ′e=0 del set = φ′1 + ω·e0 (°)"], [], [26, 12, 8, 9, 12, 3, 6, 12, 12, 12, 12, 12, 3, 14])
SETS = T27[T27.uso == "Ajuste de la curva propuesta"].drop_duplicates(["fuente_id", "set_id"])[["fuente_id", "set_id", "litologia"]]
for k, (_, r) in enumerate(SETS.iterrows(), 2):
    crit = f'Ensayos!$A$2:$A$500,"{r.fuente_id}",Ensayos!$B$2:$B$500,"{r.set_id}",Ensayos!$F$2:$F$500,"Ajuste de la curva propuesta"'
    ws.cell(k, 1, f"{r.fuente_id} | {r.set_id}"); ws.cell(k, 2, r.litologia); ws.cell(k, 3, f"=COUNTIFS({crit})")
    ws.cell(k, 4, f"=AVERAGEIFS(Ensayos!$J$2:$J$500,{crit})"); ws.cell(k, 5, f"=AVERAGEIFS(Ensayos!$V$2:$V$500,{crit})")
    ws.cell(k, 14, f"=E{k}+omega*D{k}")
nS = len(SETS) + 1
ws.cell(nS + 1, 13, "Media"); ws.cell(nS + 1, 14, f"=AVERAGE(N2:N{nS})").font = NEG
ws.cell(nS + 3, 13, "Nota"); ws.cell(nS + 3, 14, "Media de los φ′1 de los ensayos de cada set; la curva usa el φ′1 del set ajustado por mínimos cuadrados (a33), "
                                              "por eso esta media difiere algo de Parametros!B2.")
for k, e_ in enumerate([round(0.05 * i, 2) for i in range(17)], 2):
    ws.cell(k, 7, e_); ws.cell(k, 8, f"=pe0-omega*G{k}")
    for j, p_ in enumerate((46, 41, 36, 33), 9):          # Leps (1970) e Indraratna (1993): no dependen de e (Cylwik et al. 2025, p. 2 y Tabla 1)
        ws.cell(k, j, p_)
ws.cell(nS + 2, 1, "φ′1 de cada ensayo: Ensayos!V (a su e0). Cada punto de la figura de Cylwik et al. (2025) es un set; aquí, el promedio de los ensayos del set.")

# ------------------------------------------------------------------ Comparacion_envolventes (capítulo «Por qué esta curva», a35)
A35 = J("a35_comparacion.json")
ws = hoja("Comparacion_envolventes", ["Set", "σ′n (MPa)", "e0", "φ′s medido (°)", "Curva propuesta (°)", "Leps (1970) superior (°)",
                                      "Leps (1970) media (°)", "Leps (1970) inferior (°)", "Indraratna et al. (1993) (°)",
                                      "Res. curva", "Res. Leps sup.", "Res. Leps media", "Res. Leps inf.", "Res. Indraratna",
                                      "φ′e=0 sin el set (°)", "Curva sin su set (°)", "Res. curva sin su set"], [], [26] + [11] * 16)
n1 = 1 + N_AJ
_C0 = 21   # tabla de sets (a34) en las columnas U:V
ws.cell(1, _C0, "Set (a34)").font = NEG; ws.cell(1, _C0 + 1, "φ′e=0 del set (°)").font = NEG
for k, (_, r) in enumerate(S34.iterrows(), 2):
    ws.cell(k, _C0, r.set); ws.cell(k, _C0 + 1, float(r.phi_e0))
nS34 = 1 + len(S34); SR, PR = f"$U$2:$U${nS34}", f"$V$2:$V${nS34}"
p09 = lambda x, p1: f"DEGREES(ATAN(TAN(RADIANS({p1}))*({x}/sref)^(0.9-1)))"
for i in range(N_AJ):
    k = 2 + i
    ws.cell(k, 1, f'=Ensayos!A{k}&"|"&Ensayos!B{k}'); ws.cell(k, 2, f"=Ensayos!M{k}"); ws.cell(k, 3, f"=Ensayos!J{k}"); ws.cell(k, 4, f"=Ensayos!O{k}")
    ws.cell(k, 5, f"=MIN(tope,DEGREES(ATAN(TAN(RADIANS(pe0-omega*C{k}))*(B{k}/sref)^(m_-1))))")
    for j, p1 in enumerate((46, 41, 36, 33), 6):
        ws.cell(k, j, "=" + p09(f"B{k}", p1))
    for j in range(5):
        ws.cell(k, 10 + j, f"=D{k}-{get_column_letter(5 + j)}{k}")
    ws.cell(k, 15, f"=(SUM({PR})-SUMIF({SR},A{k},{PR}))/(COUNT({PR})-1)")
    ws.cell(k, 16, f"=MIN(tope,DEGREES(ATAN(TAN(RADIANS(O{k}-omega*C{k}))*(B{k}/sref)^(m_-1))))"); ws.cell(k, 17, f"=D{k}-P{k}")
# resumen
r0 = n1 + 3
ws.cell(r0, 1, "Residuo = φ′s medido − φ′s de la envolvente, a σ′n y e0 del ensayo (°). Leps (1970) e Indraratna: Ec. 9 de Cylwik et al. (2025), m = 0.9, sin e.").font = NEG
cab = ["Envolvente", "n", "Medio", "Mediana", "Medio absoluto", "RMSE", "Sobreestimados (res. < 0)", "Medio σ′n < 0.5", "Medio 0.5–2", f"Medio 2–{SNX:.2f}"]
for j, h in enumerate(cab, 1):
    c = ws.cell(r0 + 1, j, h); c.font = NEG; c.fill = CAB
RES = [("Curva propuesta", "J"), ("Curva propuesta, validación por sets", "Q"), ("Leps (1970), superior", "K"), ("Leps (1970), media", "L"),
       ("Leps (1970), inferior", "M"), ("Indraratna et al. (1993)", "N")]
B_ = f"$B$2:$B${n1}"
for i, (nm, c) in enumerate(RES):
    k = r0 + 2 + i; R = f"{c}2:{c}{n1}"
    for j, v in enumerate([nm, f"=COUNT({R})", f"=AVERAGE({R})", f"=MEDIAN({R})", f"=SUMPRODUCT(ABS({R}))/COUNT({R})",
                           f"=SQRT(SUMSQ({R})/COUNT({R}))", f'=COUNTIF({R},"<0")', f'=AVERAGEIFS({R},{B_},"<=0.5")',
                           f'=AVERAGEIFS({R},{B_},">0.5",{B_},"<=2")', f'=AVERAGEIFS({R},{B_},">2")'], 1):
        ws.cell(k, j, v)
# diseño a e = 0.45
r1 = r0 + 10
ws.cell(r1, 1, "φ′s de cada envolvente con e = e_base (°)").font = NEG
for j, s_ in enumerate((1, 4.6, 10), 2):
    c = ws.cell(r1 + 1, j, s_); c.font = NEG; c.fill = CAB
ws.cell(r1 + 1, 1, "σ′n (MPa) →").font = NEG
for i, (nm, f_) in enumerate([("Curva propuesta", lambda x: f"MIN(tope,DEGREES(ATAN(TAN(RADIANS(pe0-omega*e_base))*({x}/sref)^(m_-1))))"),
                               ("Leps (1970), superior", lambda x: p09(x, 46)), ("Leps (1970), media", lambda x: p09(x, 41)),
                               ("Leps (1970), inferior", lambda x: p09(x, 36)), ("Indraratna et al. (1993)", lambda x: p09(x, 33))]):
    k = r1 + 2 + i; ws.cell(k, 1, nm)
    for j in range(2, 5):
        ws.cell(k, j, "=" + f_(f"{get_column_letter(j)}${r1 + 1}"))
# triaxiales CU de contraste (a16, valores)
r2 = r1 + 9
_CU = pd.read_csv(AQ / "a16_contraste.csv")
ws.cell(r2, 1, "Triaxiales CU de contraste, no usados en el ajuste (a16_contraste.csv): σ′n, e0 y φ′s medidos como valores").font = NEG
for j, h in enumerate(["Grupo", "σ′n (MPa)", "e0", "φ′s medido (°)", "Res. curva", "Res. Leps sup.", "Res. Leps media", "Res. Leps inf.", "Res. Indraratna"], 1):
    c = ws.cell(r2 + 1, j, h); c.font = NEG; c.fill = CAB
for i, (_, r) in enumerate(_CU.iterrows()):
    k = r2 + 2 + i
    ws.cell(k, 1, r.grupo); ws.cell(k, 2, float(r.sn)); ws.cell(k, 3, float(r.e)); ws.cell(k, 4, float(r.phis))
    ws.cell(k, 5, f"=D{k}-MIN(tope,DEGREES(ATAN(TAN(RADIANS(pe0-omega*C{k}))*(B{k}/sref)^(m_-1))))")
    for j, p1 in enumerate((46, 41, 36, 33), 6):
        ws.cell(k, j, f"=D{k}-" + p09(f"B{k}", p1))
kz = r2 + 2 + len(_CU)
ws.cell(kz, 1, "Medio absoluto").font = NEG
ws.cell(kz + 1, 1, "Sobreestimados (res. < 0)").font = NEG
for j in range(5, 10):
    L_ = get_column_letter(j); R = f"{L_}{r2 + 2}:{L_}{kz - 1}"
    ws.cell(kz, j, f"=SUMPRODUCT(ABS({R}))/COUNT({R})"); ws.cell(kz + 1, j, f'=COUNTIF({R},"<0")')
ws.cell(kz + 3, 1, "Forma sobre 4.6 MPa (a31, prueba de extrapolación con 26 sets mundiales): sesgo con m = 0.85: "
        f"{A35['forma']['0.85']['sesgo']:+.2f}°, con m = 0.9 (Leps 1970 e Indraratna): {A35['forma']['0.9']['sesgo']:+.2f}°.")
control["comparacion"] = {k: v["Todos"] for k, v in A35["ajuste"].items()}; control["comparacion_loso"] = A35["loso"]["Todos"]
control["comparacion_cu"] = A35["cu"]; control["comparacion_diseno"] = A35["diseno"]

# ------------------------------------------------------------------ Hipotesis_m
SNT = [1, 2, 3, 4.6, 8, 15]
ws = hoja("Hipotesis_m", ["Hipótesis", "m", "φ′e=0 (°)"] + [f"φ′s a {s} MPa (°)" for s in SNT], [], [28, 8, 10] + [12] * len(SNT))
for k, (key, et) in enumerate(HIP, 2):
    ws.cell(k, 1, et); ws.cell(k, 2, f"=Parametros!B{k + 10}"); ws.cell(k, 3, f"=Parametros!C{k + 10}")
    for j, s in enumerate(SNT, 4):
        ws.cell(k, j, f"=DEGREES(ATAN(TAN(RADIANS(C{k}-omega*e_base))*({s}/sref)^(B{k}-1)))")

# ------------------------------------------------------------------ Base_mundial
W = datos(); W = W[~W.cens] if "cens" in W else W
ws = hoja("Base_mundial", ["Set", "e", "σ′n (MPa)", "τ (MPa)", "φ′s ensayo (°)", "φ′s curva a su e (°)", "Δφ′s (°)"], [], [40, 8, 10, 10, 12, 14, 10])
setcol = next((c for c in ("set_id", "set", "Set", "set_key") if c in W.columns), None)
for i, (_, r) in enumerate(W.iterrows()):
    k = i + 2
    ws.cell(k, 1, str(r[setcol]) if setcol else ""); ws.cell(k, 2, float(r.e)); ws.cell(k, 3, float(r.sn)); ws.cell(k, 4, float(r.tau))
    ws.cell(k, 5, f"=DEGREES(ATAN(D{k}/C{k}))"); ws.cell(k, 6, f"=DEGREES(ATAN(TAN(RADIANS(pe0-omega*B{k}))*(C{k}/sref)^(m_-1)))"); ws.cell(k, 7, f"=E{k}-F{k}")
nW = len(W) + 1
ws["I1"] = "Rango de σ′n"; ws["J1"] = "Ensayos"; ws["K1"] = "Bajo la curva"; ws["L1"] = "Δφ′s mediana (°)"
for c in ws["I1:L1"][0]:
    c.font = NEG; c.fill = CAB
for k, (a, b) in enumerate(((2, 4.6), (4.6, 7.9), (3, 7.9)), 2):
    ws.cell(k, 9, f"{a}–{b} MPa")
    ws.cell(k, 10, f'=COUNTIFS(C2:C{nW},">{a}",C2:C{nW},"<={b}")'); ws.cell(k, 11, f'=COUNTIFS(C2:C{nW},">{a}",C2:C{nW},"<={b}",G2:G{nW},"<0")')
    ws[f"L{k}"] = ArrayFormula(f"L{k}", f'=MEDIAN(IF((C2:C{nW}>{a})*(C2:C{nW}<={b}),G2:G{nW}))')
_k = (W.sn > 4.6) & (W.sn <= 7.9); _d = np.degrees(np.arctan(W.tau[_k] / W.sn[_k])) - np.degrees(np.arctan(np.tan(np.radians(PE0 - OM * W.e[_k])) * (W.sn[_k] / 1.013) ** (M - 1)))
control["mundial_46_79"] = dict(n=int(_k.sum()), bajo_curva=int((_d < 0).sum()), dphi_mediana=float(np.median(_d)))

def g_expr(x, a, s3):
    """σ3(x) − σ′3 para la curva τ = a·x^m (texto de fórmula de Excel)."""
    tau = f"({a}*({x})^m_)"; t = f"(m_*{a}*({x})^(m_-1))"
    return f"(({x})+{tau}*{t}-{tau}*SQRT(1+{t}^2)-({s3}))"


def newton(ws, fila, col0, a, s3, n=10):
    """Escribe n columnas de Newton desde col0 (x0 = 2·σ′3); devuelve la letra de la última."""
    prev = f"2*{s3}"
    for i in range(n):
        c = col0 + i; x = prev
        h = f"(({x})*1E-6)"
        ws.cell(fila, c, f"={x}-{g_expr(x, a, s3)}*{h}/({g_expr(x + '+' + h, a, s3)}-{g_expr(x, a, s3)})")
        prev = f"{get_column_letter(c)}{fila}"
    return get_column_letter(col0 + n - 1)


# ------------------------------------------------------------------ Ajuste (función objetivo de a12: todos los ensayos con igual peso)
AJ = pd.read_csv(AQ / "ensayos_ajuste.csv", low_memory=False)
cab = ["Fuente", "Set", "Tipo", "e0", "σ′3 (MPa)", "σ′1 (MPa)", "q obs (MPa)", "σ′n obs, DS (MPa)", "τ obs, DS (MPa)", "a = tan(φ′e0 − ω·e0)·(10pa)^(1−m)"] + \
      [f"Newton {i + 1}: σ′n (MPa)" for i in range(10)] + ["τ (MPa)", "tan φ′t", "R (MPa)", "q pred (MPa)", "Residuo", "Residuo²"]
ws = hoja("Ajuste", cab, [], [9, 22, 5, 7, 8, 8, 8, 9, 9, 12] + [9] * 10 + [8, 8, 8, 9, 9, 9])
for i, r in AJ.iterrows():
    k = i + 2; tx = r.tipo == "TX"
    for j, v in enumerate([r.fuente_id, r.set_id, r.tipo, r.e0_adoptado, r.s3 if tx else None, r.s1 if tx else None,
                           f"=F{k}-E{k}" if tx else None, None if tx else r.sn_ds, None if tx else r.tau_ds,
                           f"=TAN(RADIANS(pe0-omega*D{k}))*sref^(1-m_)"], 1):
        ws.cell(k, j, v)
    if tx:
        ul = newton(ws, k, 11, f"J{k}", f"E{k}")
        ws.cell(k, 21, f"=J{k}*{ul}{k}^m_"); ws.cell(k, 22, f"=m_*J{k}*{ul}{k}^(m_-1)"); ws.cell(k, 23, f"=U{k}*SQRT(1+V{k}^2)")
        ws.cell(k, 24, f"=2*W{k}"); ws.cell(k, 25, f"=LN(X{k})-LN(G{k})")
    else:
        ws.cell(k, 21, f"=J{k}*H{k}^m_"); ws.cell(k, 25, f"=c_ds*(LN(U{k})-LN(I{k}))")
    ws.cell(k, 26, f"=Y{k}^2")
nA = len(AJ) + 1
ws.cell(nA + 2, 25, "SSE"); ws.cell(nA + 2, 26, f"=SUM(Z2:Z{nA})").font = NEG
ws.cell(nA + 3, 1, "Control: ajuste global por mínimos cuadrados (Solver: minimizar Z" + str(nA + 2) + " cambiando Parametros!$B$2). La curva propuesta no usa este estimador sino la media por set (hoja phi1_vs_e, a33).")
ws.cell(nA + 4, 1, "Tangencia: σ′3 = σ′n + τ·tan φ′t − τ·√(1 + tan² φ′t) (círculo de Mohr tangente a la curva), resuelta con 10 pasos de Newton; q = 2R, R = τ·√(1 + tan² φ′t).")
sys.path.insert(0, str(AQ.parent / "analisis")); from modelos import sigma1, tau_leps2  # noqa: E402
def _res(pe0):
    e_ = AJ.e0_adoptado.values; t_ = (AJ.tipo == "TX").values; rr = np.empty(len(AJ))
    rr[t_] = np.log(sigma1(tau_leps2, AJ.s3.values[t_], pe0 - OM * e_[t_], M)[0] - AJ.s3.values[t_]) - np.log((AJ.s1 - AJ.s3).values[t_])
    rr[~t_] = C_DS * (np.log(tau_leps2(AJ.sn_ds.values[~t_], pe0 - OM * e_[~t_], M)) - np.log(AJ.tau_ds.values[~t_]))
    return rr
control["ajuste_resid"] = _res(PE0).tolist(); control["sse"] = {str(round(p, 2)): float((_res(p) ** 2).sum()) for p in (PE0 - 0.05, PE0, PE0 + 0.05)}

# ------------------------------------------------------------------ FLAC_tangente
def tangente(s3):
    a_ = math.tan(math.radians(PE0 - OM * EB)) * 1.013 ** (1 - M); lo_, hi_ = math.log(s3), math.log(40 * s3)
    for _ in range(80):
        mid = 0.5 * (lo_ + hi_); sn_ = math.exp(mid); t_ = a_ * sn_ ** M; b_ = 0.5 * math.atan(M * t_ / sn_) + math.pi / 4
        lo_, hi_ = (lo_, mid) if sn_ - t_ / math.tan(b_) - s3 > 0 else (mid, hi_)
    sn_ = math.exp(0.5 * (lo_ + hi_)); return sn_, a_ * sn_ ** M


ws = hoja("FLAC_tangente", ["σ′3 (MPa)", "σ′n tangencia (MPa)", "τ (MPa)", "φ′t (°)", "c′t (kPa)", "σ′1 − σ′3 (MPa)", "a"] + [f"Newton {i + 1}" for i in range(10)],
          [], [10, 16, 10, 10, 10, 14, 10] + [9] * 10)
control["flac_sn"] = []
for k, s3 in enumerate((0.1, 0.5, 1.0, 2.0, 5.0), 2):
    sn_, t_ = tangente(s3); control["flac_sn"].append(sn_)
    ws.cell(k, 1, s3); ws.cell(k, 7, "=TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)")
    ul = newton(ws, k, 8, f"G{k}", f"A{k}"); ws.cell(k, 2, f"={ul}{k}")
    ws.cell(k, 3, "=TAN(RADIANS(pe0-omega*e_base))*sref^(1-m_)*B%d^m_" % k)
    ws.cell(k, 4, f"=DEGREES(ATAN(m_*C{k}/B{k}))"); ws.cell(k, 5, f"=(1-m_)*C{k}*1000"); ws.cell(k, 6, f"=2*C{k}/COS(RADIANS(D{k}))")
ws["A8"] = "σ′n de tangencia: σ′3 = σ′n + τ·tan φ′t − τ·√(1 + tan² φ′t), resuelta con 10 pasos de Newton (columnas H a Q); tan φ′t = m·τ/σ′n, c′t = (1 − m)·τ."

# ------------------------------------------------------------------ Relacion_vacios
EV = [("ROM", "SDA, INF-115-R-0005 (Stantec 2025), valor adoptado", 20.0, 0.029), ("ROM", "DDE, INF-115-R-0004, Tabla 9-3", 19.6, 0.029),
      ("ROM", "Golder 2010, citado en el DDE (p. 59)", 20.75, 0.029), ("Chancado", "Golder 2010, citado en el DDE", 21.66, 0.028),
      ("Chancado", "SDA, valor adoptado", 22.0, 0.028), ("Enrocado 3B", "DDE y SDA", 22.0, 0.029)]
ws = hoja("Relacion_vacios", ["Material", "Fuente", "γ (kN/m³)", "w", "Gs", "γw", "e"], [], [12, 50, 10, 8, 7, 7, 8])
for k, (mt, fu, g, w) in enumerate(EV, 2):
    for j, v in enumerate([mt, fu, g, w, 2.75, 9.81, f"=E{k}*F{k}*(1+D{k})/C{k}-1"], 1):
        ws.cell(k, j, v)

# ------------------------------------------------------------------ Perfiles_e, Edometros, Tamano_particula
ws = hoja("Perfiles_e", ["Profundidad (m)", "σ′v perfil A (MPa)", "e perfil A", "σ′v perfil V (MPa)", "e perfil V"], [], [14, 16, 10, 16, 10])
for z, sa, ea, sv, ev in zip(PE["z"], PE["sv_A"], PE["e_A"], PE["sv_V"], PE["e_V"]):
    ws.append([z, sa, ea, sv, ev])
ws["G1"] = "Valenzuela et al. (2008), curva de diseño (lectura de la figura): altura (m)"; ws["H1"] = "γd (t/m³)"
for i_, (h_, g_) in enumerate(zip(PEM.HV, PEM.GDV), 2):
    ws.cell(i_, 7, float(h_)); ws.cell(i_, 8, float(g_))
ws = hoja("Edometros", ["Ensayo", "Diámetro (mm)", "e0", "Cc1 (0.2–0.8 MPa)", "Cc2 (1.2–4.8 MPa)", "e a 300 m", "e a 700 m"], [], [40, 12, 8, 14, 14, 10, 10])
for r in EDO:
    p_ = {str(k): v for k, v in (r["perfil"] or {}).items()}
    ws.append([r["ensayo"], r["diametro_mm"], r["e0"], r["Cc1"], r["Cc2"], p_.get("300"), p_.get("700")])
ws = hoja("Tamano_particula", ["D lab (mm)", "m_w", "D campo (mm)", "λ = (Dlab/Dcampo)^(3/mw)", "Factor sobre τ = λ^(1−m)", "φ′1 (°)", "Diferencia (°)"],
          [], [10, 8, 12, 22, 22, 10, 12])
k = 2
for dl in (178, 102, 25):
    for mw in (4.0, 3 / 0.45, 10.0):
        for j, v in enumerate([dl, round(mw, 4), 1000, f"=(A{k}/C{k})^(3/B{k})", f"=D{k}^(1-m_)", f"=DEGREES(ATAN(E{k}*TAN(RADIANS(pe0-omega*e_base))))",
                               f"=F{k}-(pe0-omega*e_base)"], 1):
            ws.cell(k, j, v)
        k += 1
# ------------------------------------------------------------------ Figuras del memo (gráficos nativos; a28_figuras_excel.py)
import re  # noqa: E402
import a28_figuras_excel as FX  # noqa: E402
_aux = AQ.parent / "eor_antamina" / "curva_leps2_antamina.aux"
ROT = {}
if _aux.exists():
    for _l in open(_aux, encoding="utf-8", errors="ignore"):
        _m = re.match(r"\\newlabel\{(fig:[^}]+)\}\{\{(\d+)\}", _l)
        if _m:
            ROT[_m.group(1)] = _m.group(2)
TIT = {"fig:phi1e": "φ′1 de cada set de Antamina frente a su e0 media, y la recta de la curva propuesta",
       "fig:phi": "Curva propuesta y su banda de incertidumbre, con los ensayos del ajuste y la base mundial llevados a e base",
       "fig:contraste": "Contraste de la curva propuesta con los triaxiales CU de desmonte ROM no usados en el ajuste",
       "fig:bg": "Índice de rotura de Marsal (1967) en los triaxiales drenados del ajuste",
       "fig:tau": "Curva propuesta y su banda de incertidumbre en τ–σ′n",
       "fig:residuos": "Residuos del ajuste: (a) curva propuesta, (b) nivel reajustado con m = 1, (c) residuo por litología",
       "fig:extrap": "Prueba de extrapolación: (a) ejemplo, (b) base mundial, (c) Antamina, (d) sesgo según m",
       "fig:e": "(a) Relación de vacíos con la profundidad; (b) ángulo de fricción secante de la curva propuesta con cada supuesto"}
_wf, _figs = FX.agregar(wb, HB=HB, SNX=SNX, SNW=SNW, EB=EB, TOPE=TOPE, PE=PE, PEM=PEM, A16BG=pd.read_csv(AQ / "a16_bg.csv"),
                        A16C=pd.read_csv(AQ / "a16_contraste.csv"), S34=S34, N_AJ=N_AJ, NW=len(W), rbg=A16["bg"]["r"],
                        A36=J("a36_residuos.json"), P35=pd.read_csv(AQ / "a35_residuos.csv"), FIT=T27.iloc[:N_AJ],
                        A38=J("a38_extrapolacion.json"), WDF=W)
FX.colocar(_wf, _figs, {k: (ROT.get(k, ""), t) for k, t in TIT.items()})
wb.save(OUT)
json.dump(control, open(AQ / "a28_valores_control.json", "w", encoding="utf-8"))
print("escrito", OUT, "| ensayos", len(T27), "| base mundial", len(W), "| curva", len(SN), "puntos")
