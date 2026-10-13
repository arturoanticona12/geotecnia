"""Manuscrito «Leps 3.0» (artículo completo, español). Contenido aquí; números de Slide2 y de Python leídos de resultados.
Uso: python construir_leps3.py"""
import json, re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

AQUI = Path(__file__).parent; RAIZ = AQUI.parent; FIG = RAIZ / "figuras"; AN = RAIZ / "analisis"
RES = Path(r"H:\Mi unidad\slide2-bench-intercambio\resultados")
S52 = json.load(open(RES / "52-botadero700-leps3" / "fs_por_escenario.json", encoding="utf-8"))["botadero_700m"]["fs"]
EJ = json.load(open(AN / "ejemplo_leps3.json", encoding="utf-8"))
WG = json.load(open(AN / "wedge_global.json", encoding="utf-8"))

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5); s.left_margin = s.right_margin = Cm(2.5)


def runs(p, texto):
    for tok in re.split(r"(\*[^*]+\*|\^[^^]+\^|_\{[^}]+\}|\[\[[^\]]+\]\])", texto):
        if not tok:
            continue
        if tok.startswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("^"):
            r = p.add_run(tok[1:-1]); r.font.superscript = True
        elif tok.startswith("_{"):
            r = p.add_run(tok[2:-1]); r.font.subscript = True
        elif tok.startswith("[["):
            r = p.add_run(tok[2:-2]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p.add_run(tok)
    return p


def par(texto, sangria=True, just=True):
    p = doc.add_paragraph()
    if just:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if sangria:
        p.paragraph_format.first_line_indent = Cm(0.6)
    return runs(p, texto)


def vineta(texto):
    p = doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return runs(p, texto)


def titulo(texto, nivel=1):
    p = doc.add_paragraph(); r = p.add_run(texto); r.bold = True
    r.font.size = Pt(12 if nivel == 1 else 11); r.italic = nivel == 2
    p.paragraph_format.space_before = Pt(12 if nivel == 1 else 8); p.paragraph_format.keep_with_next = True


def ecuacion(texto, n):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(8), alignment=1); pf.tab_stops.add_tab_stop(Cm(16), alignment=2)
    p.add_run("\t"); runs(p, texto); p.add_run(f"\t({n})")


def rotulo(tipo, n, texto):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8)
    p.add_run(f"{tipo} {n}").bold = True
    p2 = doc.add_paragraph(); p2.paragraph_format.keep_with_next = True; runs(p2, f"*{texto}*")


def fuente(texto):
    p = doc.add_paragraph(); r = p.add_run(texto); r.font.size = Pt(9); p.paragraph_format.space_after = Pt(10)


def figura(n, archivo, texto, fte, ancho=15.5):
    rotulo("Figura", n, texto)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / archivo), width=Cm(ancho)); fuente(fte)


def bordes(c, arriba=None, abajo=None):
    tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for lado, v in (("top", arriba), ("bottom", abajo)):
        if v:
            e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v)); b.append(e)
    tcPr.append(b)


def tabla(n, texto, enc, filas, anchos, fte, nota=None):
    rotulo("Tabla", n, texto)
    t = doc.add_table(rows=1 + len(filas), cols=len(enc)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for j, h in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; runs(c.paragraphs[0], h)
        for r in c.paragraphs[0].runs:
            r.bold = True
        bordes(c, 8, 6)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f):
            c = t.rows[i].cells[j]; c.text = ""; runs(c.paragraphs[0], str(v))
            if i == len(filas):
                bordes(c, abajo=8)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        for j, c in enumerate(row.cells):
            c.width = Cm(anchos[j])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
                for r in p.runs:
                    r.font.size = Pt(9)
    if nota:
        p = doc.add_paragraph(); runs(p, nota); p.paragraph_format.space_before = Pt(3)
        for r in p.runs:
            r.font.size = Pt(9)
    fuente(fte)


f2 = lambda x: f"{x:.2f}"
ej = lambda p1, nm: EJ[f"{p1}|{nm}"]["F"]
s2 = lambda esc, met: S52[f"Fundacion horizontal / {esc}"][met]

# ============================================================== portada y resumen
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Leps 3.0: envolventes de resistencia al corte de enrocados y desmonte de mina para grandes confinamientos")
r.bold = True; r.font.size = Pt(14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, "[[Autor(es) y afiliaciones por confirmar]]")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, "Borrador del manuscrito. Formato de artículo de investigación.")

titulo("Resumen")
par("Las envolventes de Leps (1970), actualizadas por Cylwik et al. (2025) como Leps 2.0, son la referencia empírica "
    "para estimar la resistencia al corte de enrocados y desmonte de mina. La nueva generación de botaderos, con alturas "
    "de 500–700 m, impone en su base esfuerzos de 10–15 MPa, mientras que la base de datos de Leps 2.0 llega a "
    "σ′_{n} = 7.9 MPa y solo el 3.6 % de sus ensayos supera 5 MPa. Este trabajo revisa las ecuaciones de Leps 2.0, evalúa su "
    "extrapolación con validación cruzada dejando fuera cada publicación y en esfuerzos principales medidos, y propone "
    "Leps 3.0. Las ecuaciones cerradas de Leps 2.0 son correctas; la aproximación explícita de su Ecuación 23 tiene errores "
    "de hasta 7.4 % en σ′_{1}, frente al 1 % declarado. El exponente de curvatura típico m = 0.9 sobreestima la resistencia "
    "a alta presión: al calibrar con σ′_{n} ≤ 3 MPa y predecir por encima, el sesgo es de +10.4 % y el percentil 90 de la "
    "sobreestimación, de +25.7 %. Tres estimaciones independientes coinciden en m ≈ 0.87. Cuatro formas alternativas con "
    "piso de estado crítico o curvatura dependiente de la densidad no mejoran la predicción fuera de muestra, porque los "
    "datos no permiten identificar el piso. Leps 3.0 conserva la tabla de φ′_{1} de Leps 2.0 y agrega: m = 0.87; una "
    "distribución poblacional de m (media 0.851, desviación entre materiales 0.056) para actualizar con ensayos del "
    "sitio; reducciones de diseño de φ′_{1} según el nivel de esfuerzo; un dominio de aplicación explícito; un criterio de "
    "drenaje para zonas saturadas; y una implementación abierta. En un botadero de 700 m, la envolvente de diseño reduce "
    "el factor de seguridad en 12 % respecto de Leps 2.0, y una base saturada con respuesta no drenada lo lleva "
    "por debajo de 1.0.", sangria=False)
par("*Palabras clave:* enrocado; desmonte de mina; resistencia al corte; envolvente no lineal; botaderos de gran altura; "
    "rotura de partículas; actualización bayesiana; presión de poros.", sangria=False, just=False)

# ============================================================== 1
titulo("1. Introducción")
par("Leps (1970) reunió 109 ensayos triaxiales de gran escala y propuso tres envolventes no lineales para enrocados. "
    "Durante cinco décadas, esas curvas se usaron en presas y botaderos cuando no había ensayos del sitio. Cylwik et al. "
    "(2025) las actualizaron con 938 ensayos, propusieron un criterio de potencia con un solo parámetro de ajuste, el "
    "ángulo secante φ′_{1} a σ′_{n} = 10 p_{a}, y un modelo empírico que estima φ′_{1} a partir de la relación de vacíos "
    "y de la calidad de la roca. Llamaron al conjunto Leps 2.0.")
par("La minería a cielo abierto plantea ahora botaderos de 500–700 m de altura (Valenzuela et al., 2008). Su base trabaja "
    "con esfuerzos efectivos de 10–15 MPa, en un material que se tritura, pierde dilatancia, reduce su permeabilidad y "
    "puede responder de forma contractiva (Linero et al., 2007; Ovalle et al., 2020). La base de datos de Leps 2.0 "
    "cubre bien el rango de 0.1–3 MPa, pero tiene pocos ensayos por encima. Extrapolar una envolvente de potencia con "
    "curvatura constante hasta la base de esos botaderos es una decisión de diseño cuyo error no se había cuantificado.")
par("Este trabajo tiene tres objetivos: (a) revisar la consistencia de las ecuaciones de Leps 2.0; (b) medir, fuera de "
    "muestra, el error de su extrapolación a alta presión y probar formas alternativas con base física; y (c) proponer "
    "Leps 3.0, un criterio que conserve la parte de Leps 2.0 que los datos respaldan, corrija lo que no, cuantifique la "
    "incertidumbre según el nivel de esfuerzo y pueda recalibrarse cuando existan ensayos a mayor presión.")

# ============================================================== 2
titulo("2. Revisión de las ecuaciones de Leps 2.0")
par("El criterio de Leps 2.0 (Cylwik et al., 2025, Ecuación 8) es:")
ecuacion("τ = tan φ′_{1} (10 p_{a})^1−m^ σ′_{n}^m^", 1)
par("con p_{a} = 0.1013 MPa, m = 0.9 como valor típico y φ′_{1} = φ′_{e=0} − ω e, donde φ′_{e=0} vale 53°, 50° o 47° "
    "para rocas de calidad buena, moderada o baja y ω = 20°. Se verificaron de forma simbólica las Ecuaciones 2, 3, 6 y 15 "
    "de ese trabajo y, en 2 000 combinaciones aleatorias de m (0.6–0.99), φ′_{1} (30–50°) y σ′_{n} (0.01–16 MPa), las "
    "Ecuaciones 10–12, 16–18 y 21. Las siete coinciden con la solución numérica del círculo de Mohr tangente con errores menores que "
    "10^−13^.")
par("La aproximación explícita de la Ecuación 23 de Cylwik et al. (2025) se declara con error menor que 1 % en σ′_{1} para "
    "0.7 < m < 1 y 25° < φ′_{1} < 50°. El error máximo encontrado en ese rango es de 7.4 % (m = 0.7, φ′_{1} = 50°, "
    "σ′_{3} < 60 kPa). Con m ≤ 0.85, el error llega a 1–2.3 % en el rango de 0.06–2 MPa; con m ≥ 0.9 no supera "
    "1.3 %. Para valores de m menores que 0.9 conviene resolver la Ecuación 21 por iteración, como hace la "
    "implementación que acompaña este trabajo.")

# ============================================================== 3
titulo("3. Datos y métodos")
titulo("3.1 Base de datos", 2)
par("Se usó la base de datos publicada con Leps 2.0 (Arrieta, 2025): 938 ensayos de 40 publicaciones, con muestras de "
    "51–1 130 mm de diámetro y D_{max} de 4.8–200 mm. Los sets se identificaron por publicación y código, porque dos códigos "
    "se repiten entre publicaciones. Para los análisis en esfuerzos principales se usaron los 918 ensayos triaxiales con "
    "σ′_{3} > 0 (282 sets, 39 publicaciones). La calidad de roca se asignó según la Tabla 5 de Cylwik et al. (2025). Las "
    "mezclas sin categoría directa se asignaron así: caliza con basalto a calidad moderada, y grava de río de gneis, granito "
    "y esquisto a calidad baja.")
figura(1, "fig1_cobertura_esfuerzos.png", "Distribución de σ′n en la base de datos y rangos de esfuerzo de un botadero de 700 m",
       "Fuente: elaboración propia con datos de Arrieta (2025). Las bandas indican el rango de σ′n en las superficies "
       "críticas calculadas en la sección 6.")

titulo("3.2 Evaluación en esfuerzos principales", 2)
par("Los pares (σ′_{n}, τ) de la base dependen del exponente con que se convirtieron. Para comparar envolventes de "
    "distinta forma sin ese sesgo, cada modelo se evaluó prediciendo el desviador de falla q = σ′_{1} − σ′_{3} a partir del "
    "σ′_{3} medido, con el círculo de Mohr tangente a la envolvente (Ecuaciones 10–12 de Cylwik et al., 2025, "
    "generalizadas a envolventes de otra forma y resueltas por bisección). El error relativo es ε = q_{pred}/q − 1; ε > 0 "
    "indica sobreestimación. Se informan el sesgo medio, el error absoluto medio (EAM) y el percentil 90 de ε, que mide "
    "la cola de sobreestimación relevante para el diseño.")

titulo("3.3 Modos de uso y validación cruzada", 2)
par("Se evaluaron tres modos que reproducen el uso práctico del criterio:")
vineta("Modo A, predictivo: sin ensayos del sitio, φ′_{1} se estima con la calidad de roca y la relación de vacíos.")
vineta("Modo B, extrapolación desde ensayos convencionales: φ′_{1} se calibra en cada set con los ensayos de "
       "σ′_{n} ≤ 1.5 MPa y se predicen los de σ′_{n} > 2 MPa (46–47 sets, 13 publicaciones).")
vineta("Modo C, extrapolación desde ensayos de presión intermedia: calibración con σ′_{n} ≤ 3 MPa y predicción por "
       "encima de 3 MPa (52 sets, 8 publicaciones).")
par("En los tres modos, los parámetros globales de cada modelo se eligieron sin la publicación evaluada y se aplicaron "
    "a ella (validación cruzada dejando fuera una publicación). Así, el error informado corresponde a un laboratorio, "
    "un material y un procedimiento que el modelo no vio.")

titulo("3.4 Formas de envolvente evaluadas", 2)
par("Además de la ley de potencia de la Ecuación 1, se probaron cuatro formas con base física, ancladas en "
    "φ′_{1} a 10 p_{a} para conservar la tabla de Leps 2.0:")
vineta("Bolton con piso: φ′_{s} = φ′_{cs} + (φ′_{1} − φ′_{cs}) ln(σ_{cr}/σ′_{n})/ln(σ_{cr}/10 p_{a}), con φ′_{s} ≥ φ′_{cs}, "
       "que traslada a la forma secante la relación entre dilatancia y esfuerzo de Bolton (1986) y la resistencia "
       "equivalente de partícula de Barton y Kjærnsli (1981).")
vineta("Decaimiento hacia el estado crítico: tan φ′_{s} = tan φ′_{cs} + (tan φ′_{1} − tan φ′_{cs})(σ′_{n}/10 p_{a})^−n^.")
vineta("Piso suave: φ′_{s} = φ′_{cs} + (φ′_{1} − φ′_{cs}) ln(1 + σ_{cr}/σ′_{n})/ln(1 + σ_{cr}/10 p_{a}).")
vineta("Curvatura dependiente de la resistencia: m = m_{0} − k(φ′_{1} − 41°), que reproduce la mayor pérdida de "
       "dilatancia de los materiales densos predicha por Bolton (1986).")
par("Los parámetros globales (φ′_{cs}, σ_{cr}, n, m_{0}, k) se buscaron en grillas y se eligieron por el mínimo error "
    "cuadrático en ln q, con la misma validación cruzada.")

titulo("3.5 Distribución poblacional de m y actualización con ensayos del sitio", 2)
par("En los sets con al menos tres ensayos con σ′_{3} ≥ 60 kPa y una razón σ′_{3,max}/σ′_{3,min} ≥ 4, se ajustaron φ′_{1} y m "
    "con su error estándar (134 sets, 27 publicaciones). La media poblacional μ y la varianza entre materiales τ^2^ se "
    "estimaron por el método de momentos con ponderación de efectos aleatorios. Para un material nuevo con ensayos que "
    "permiten ajustar m_{s} con error estándar s_{s}, el valor actualizado es:")
ecuacion("m_{post} = (m_{s}/s_{s}^2^ + μ/τ^2^) / (1/s_{s}^2^ + 1/τ^2^)", 2)
par("que corresponde a la posterior de un modelo normal con previa N(μ, τ^2^). En la validación, μ y τ se estimaron sin "
    "la publicación evaluada.")

titulo("3.6 Análisis de estabilidad y de presión de poros", 2)
par("Se analizó un botadero genérico de 700 m con talud global de 2H:1V, γ = 21 kN/m^3^ y fundación de roca competente "
    "(c′ = 1 MPa, φ′ = 45°), horizontal o con pendiente de 10°. Se implementó en Python el método de Bishop (1955) con "
    "envolvente no lineal, linealizada en la tangente al σ′_{n} de cada dovela. Reproduce la solución de talud infinito "
    "(1.566 frente a 1.562) y los factores de seguridad de Slide2 9.043 (Rocscience, s. f.) con diferencias de 0.008 o "
    "menos. En Slide2 se usaron además la búsqueda no circular por recocido simulado y los métodos de Spencer (1967) y "
    "GLE. Los escenarios con agua y la presión de poros por consolidación se describen en la sección 6.")

# ============================================================== 4
titulo("4. Resultados")
titulo("4.1 Sesgo de Leps 2.0 a alta presión", 2)
par("Dentro de cada set, el φ′_{1} implícito con m = 0.9 cae con el esfuerzo (Figura 2a): respecto de la referencia de "
    "0.3–1.5 MPa, la mediana es −0.8° entre 1 y 2 MPa y de −2.2° a −3.1° entre 3 y 8 MPa. Un modelo de efectos mixtos con "
    "intercepto aleatorio por set (Seabold y Perktold, 2010) da una caída de 3.2° por década de σ′_{n} bajo 1 MPa y de "
    "2.2° por década sobre 1 MPa: la envolvente es más curva que m = 0.9 en el rango ensayado, y ese exceso disminuye a "
    "alta presión. La Tabla 1 muestra el efecto en los tres modos. Con m = 0.9, el sesgo a alta presión es de +4.5 % en el "
    "modo A, +7.3 % en el modo B y +10.4 % en el modo C, y el percentil 90 de la sobreestimación llega a 16–26 %.")
figura(2, "fig2_extrapolacion_m.png",
       "(a) Caída de φ′1 dentro de cada set respecto de la referencia de 0.3–1.5 MPa. (b) Error de τ al calibrar con "
       "σ′n ≤ 1.5 MPa y extrapolar a σ′n ≥ 2 MPa, según el exponente m",
       "Fuente: elaboración propia con datos de Arrieta (2025). En (b), el error se evaluó en τ dentro de cada set; la "
       "banda gris es el intervalo de 90 % del exponente sin sesgo.")
tabla(1, "Error de predicción de q a alta presión con Leps 2.0 y Leps 3.0 (validación dejando fuera cada publicación)",
      ["Modo y rango evaluado", "Modelo", "Sesgo", "EAM", "P90 de sobreestimación"],
      [["A, σ′_{3} de 2–5 MPa (90 ensayos)", "Leps 2.0 (m = 0.9)", "+4.5 %", "9.0 %", "+16.3 %"],
       ["", "Leps 3.0 (m = 0.87)", "−2.6 %", "8.6 %", "+8.8 %"],
       ["A, σ′_{3} de 1–2 MPa (90 ensayos)", "Leps 2.0 (m = 0.9)", "+3.4 %", "10.2 %", "+21.4 %"],
       ["", "Leps 3.0 (m = 0.87)", "−0.3 %", "10.1 %", "+17.6 %"],
       ["B, σ′_{n} > 2 MPa (64 ensayos)", "Leps 2.0 (m = 0.9)", "+7.3 %", "10.2 %", "+17.9 %"],
       ["", "Leps 3.0 (m = 0.87)", "+0.1 %", "6.9 %", "+10.7 %"],
       ["C, σ′_{n} > 3 MPa (90 ensayos)", "Leps 2.0 (m = 0.9)", "+10.4 %", "12.1 %", "+25.7 %"],
       ["", "Leps 3.0 (m = 0.87)", "+2.8 %", "8.1 %", "+17.2 %"],
       ["", "Leps 3.0 con actualización (Ec. 2)", "−1.0 %", "8.1 %", "+13.1 %"]],
      [4.6, 4.6, 1.6, 1.4, 2.8],
      "Fuente: elaboración propia con datos de Arrieta (2025).",
      "Nota: errores en q = σ′1 − σ′3. Modo A con la tabla de φ′1 de Cylwik et al. (2025); modos B y C con φ′1 calibrado "
      "en cada set.")

titulo("4.2 Exponente de extrapolación", 2)
par("Tres estimaciones independientes coinciden. Dentro de cada set, el exponente que anula el sesgo medio al "
    "extrapolar τ desde σ′_{n} ≤ 1.5 MPa es 0.866 (intervalo de 90 % por remuestreo de sets: 0.856–0.875). En esfuerzos "
    "principales, la ley de potencia elegida dejando fuera cada publicación tiene m = 0.87–0.88 en los 13 pliegues. Al "
    "reajustar el modelo empírico de Leps 2.0 (tabla de φ′_{1} y exponente) con la guía 60-600 de Cylwik et al. (2025), sin ensayos con "
    "σ′_{3} < 60 kPa, el exponente resulta 0.867 y la tabla de φ′_{e=0} queda en 54.0°, 49.4° y 46.1°, con ω = 18.1°, "
    "valores cercanos a los publicados. Con m = 0.87 y la tabla publicada, el sesgo en el modo A queda entre −2.6 % y "
    "+0.5 % en el rango de 0.06–5 MPa (Tabla 1), frente a −7.5 % y +4.5 % con m = 0.9. En el modo A, el valor 0.87 se "
    "fijó con la base entera y no dentro de cada pliegue; la validación por publicación de los modos B y C, donde sí se "
    "eligió dentro del pliegue, da el mismo valor.")
par("La distribución poblacional de m tiene media 0.851 y desviación entre materiales de 0.056, y no cambia con el "
    "esfuerzo máximo alcanzado por los sets (0.848, 0.851 y 0.853 para σ′_{3,max} de 0–1, 1–2.5 y 2.5–5 MPa). La "
    "diferencia entre 0.851 y 0.87 es coherente con la disminución de la curvatura con el esfuerzo: el m de un set "
    "promedia su rango de ensayo, mientras que el exponente de extrapolación describe el tramo de mayor esfuerzo.")

titulo("4.3 Formas alternativas: lo que los datos no permiten identificar", 2)
par("La Tabla 2 resume la comparación en el modo B. Las tres formas con piso de estado crítico no mejoran a la ley de "
    "potencia con m = 0.87, y en las tres el piso elegido queda en el límite inferior de la grilla (φ′_{cs} = 26°), es decir, "
    "fuera de la zona donde influye. Con datos hasta 7.9 MPa, la curva sigue bajando sin señal de un piso, aunque la "
    "curvatura disminuye. Este resultado no contradice la existencia de un ángulo de estado crítico a muy alta presión "
    "(Coop et al., 2004; Yamamuro y Lade, 1996); indica que la base actual no permite calibrarlo.")
tabla(2, "Formas de envolvente en el modo B (calibración σ′n ≤ 1.5 MPa, predicción σ′n ≥ 2 MPa)",
      ["Forma", "Parámetros globales", "Sesgo, σ′_{n} ≥ 2", "EAM, σ′_{n} ≥ 2", "Sesgo, σ′_{n} ≥ 3", "EAM, σ′_{n} ≥ 3"],
      [["Potencia (Leps 2.0 y 3.0)", "m = 0.87", "+1.0 %", "7.3 %", "−3.8 %", "8.4 %"],
       ["Decaimiento hacia φ′_{cs}", "φ′_{cs} = 26°, n = 0.3", "+1.0 %", "7.0 %", "−2.1 %", "8.2 %"],
       ["Bolton con piso", "φ′_{cs} = 26°, σ_{cr} = 80 MPa", "+2.8 %", "8.0 %", "−0.2 %", "9.3 %"],
       ["Piso suave", "φ′_{cs} = 26°, σ_{cr} = 60 MPa", "+1.0 %", "7.9 %", "−2.0 %", "9.6 %"],
       ["m dependiente de φ′_{1}", "m_{0} = 0.88, k = 0.008", "+2.1 %", "8.6 %", "−2.2 %", "10.4 %"]],
      [3.6, 4.0, 1.9, 1.9, 1.9, 1.9],
      "Fuente: elaboración propia con datos de Arrieta (2025).",
      "Nota: 47 sets, 67 ensayos de evaluación, 13 publicaciones; errores en q. Los parámetros indicados son los elegidos "
      "con la base entera; en la validación se eligieron sin la publicación evaluada.")
par("La interacción entre densidad y esfuerzo sí aparece en los datos. Con un moderador independiente de los ensayos "
    "(φ′_{1} estimado por la calidad de roca y la relación de vacíos), cada grado adicional de φ′_{1} agrega −0.47° de "
    "caída por década de esfuerzo (intervalo de 90 %: −0.76° a −0.19°). Sin embargo, la forma con m dependiente de φ′_{1} "
    "aumenta el error fuera de muestra en los modos A y B. La interacción existe, pero con 47 sets no mejora la "
    "predicción, y Leps 3.0 no la incorpora.")

titulo("4.4 Rotura de partículas y comportamiento volumétrico", 2)
par("El índice de rotura de Marsal (B_{g}), disponible en 627 ensayos, tiene una mediana de 5.7 % para σ′_{n} de "
    "0.1–0.3 MPa, 24.4 % para 1–2 MPa y 36.0 % para 2–4 MPa (Figura 3). La deformación volumétrica neta en la falla, "
    "disponible en 347 ensayos, es de compresión en los 347 y aumenta de −1.9 % en promedio para σ′_{3} ≤ 0.1 MPa a "
    "−7.1 % para σ′_{3} de 3–5 MPa (valores negativos indican compresión, según su tendencia con σ′_{3} y lo informado por "
    "Valenzuela et al., 2008). A los esfuerzos de la mitad inferior de un botadero alto, el material es contractivo.")
figura(3, "fig4_rotura_Bg.png", "Índice de rotura de Marsal en función de σ′n, por clase de resistencia de la roca",
       "Fuente: elaboración propia con datos de Arrieta (2025). Se omiten 2 ensayos con Bg > 120 %.")

titulo("4.5 Incertidumbre según el nivel de esfuerzo", 2)
par("El error del modo A disminuye con el esfuerzo, como observaron Cylwik et al. (2025) en grados de φ′_{1}. Para "
    "obtener una envolvente de diseño se calculó, en cada rango de σ′_{3}, el percentil 10 y 20 de q_{obs}/q_{pred} con "
    "Leps 3.0 y se expresó como una reducción equivalente de φ′_{1} (Tabla 3). La reducción va de 4.8° a bajo esfuerzo a "
    "1.6° entre 2 y 5 MPa en el percentil 10.")
tabla(3, "Reducción de diseño de φ′1 en Leps 3.0 (modo predictivo)",
      ["Rango de σ′_{3} (MPa)", "Ensayos", "Δφ′_{1}, percentil 10", "Δφ′_{1}, percentil 20"],
      [["0.06–0.3", "297", "−4.8°", "−3.1°"], ["0.3–1", "333", "−3.9°", "−2.0°"], ["1–2", "90", "−3.0°", "−1.4°"],
       ["2–5", "90", "−1.6°", "−0.7°"]],
      [4.0, 2.2, 3.6, 3.6],
      "Fuente: elaboración propia con datos de Arrieta (2025).",
      "Nota: reducción de φ′1 que produce el cuantil indicado de q_obs/q_pred con m = 0.87, evaluada para φ′1 = 41° en el "
      "σ′3 central de cada rango. Entre rangos se interpola en log σ′3; fuera de 0.06–5 MPa se mantiene el valor extremo.")

# ============================================================== 5
titulo("5. Leps 3.0")
titulo("5.1 Formulación", 2)
par("Leps 3.0 conserva la forma de la Ecuación 1 y la tabla de φ′_{1} de Cylwik et al. (2025), y cambia el exponente "
    "central a m = 0.87:")
ecuacion("τ = tan φ′_{1} (10 p_{a})^0.13^ σ′_{n}^0.87^ ,   φ′_{1} = φ′_{e=0} − 20° e", 3)
par("La envolvente de diseño reduce φ′_{1} según el σ′_{3} asociado a cada σ′_{n} (Tabla 3):")
ecuacion("τ_{d} = tan[φ′_{1} + Δφ′_{1}(σ′_{3})] (10 p_{a})^0.13^ σ′_{n}^0.87^", 4)
par("La Figura 4 compara ambas con Leps 2.0. Con φ′_{1} = 41°, el ángulo secante a 10 MPa es 34.7° con Leps 2.0 y 32.8° "
    "con Leps 3.0. Bajo 1 MPa, la envolvente central de Leps 3.0 queda sobre la de Leps 2.0; los datos lo respaldan (el "
    "sesgo con σ′_{3} de 0.06–0.3 MPa pasa de −7.5 % a −0.8 %), pero la cola de sobreestimación a bajo esfuerzo aumenta "
    "(percentil 90 de +23 % a +32 %). Por eso la envolvente de diseño es parte del criterio y no un complemento opcional.")
figura(4, "fig6_leps3_envolventes.png", "Envolventes de Leps 2.0 y Leps 3.0 para φ′1 = 41°",
       "Fuente: elaboración propia. La banda sombreada es la diferencia entre la envolvente central de Leps 3.0 y su "
       "envolvente de diseño P10.")

titulo("5.2 Procedimiento", 2)
par("El procedimiento de Leps 3.0 tiene seis pasos:")
vineta("Paso 1, aplicabilidad. El material debe cumplir D_{50} > 2 mm y finos < 25 % (Cylwik et al., 2025). El criterio "
       "está respaldado por datos hasta σ′_{n} ≈ 8 MPa (σ′_{3} ≈ 4.5 MPa).")
vineta("Paso 2, φ′_{1}. Sin ensayos, se toma de la tabla de calidad de roca y relación de vacíos. Con ensayos del sitio, "
       "se ajusta en ln q con m fijo.")
vineta("Paso 3, m. Sin ensayos o con ensayos solo bajo 1.5 MPa, m = 0.87. Si los ensayos del sitio llegan a "
       "σ′_{3} ≥ 1.5 MPa y cubren una razón σ′_{3,max}/σ′_{3,min} ≥ 4, se ajusta m_{s} y se actualiza con la Ecuación 2 "
       "(μ = 0.851, τ = 0.056).")
vineta("Paso 4, envolvente de diseño. Se aplica Δφ′_{1}(σ′_{3}) de la Tabla 3, con el percentil que corresponda a la "
       "consecuencia de falla de la estructura.")
vineta("Paso 5, drenaje. En zonas saturadas donde σ′_{3} supera 0.2–0.5 MPa, sobre el cual el desmonte suelto deja de "
       "dilatar (Linero et al., 2007; Valenzuela et al., 2008), se verifica la estabilidad con resistencia no drenada (S_{u}/σ′_{v}), medida en ensayos CIU a alta presión sobre la "
       "granulometría triturada. Mientras no existan ensayos, S_{u}/σ′_{3c} = 0.28–0.29 (Valenzuela et al., 2008) sirve "
       "como valor de tamizaje.")
vineta("Paso 6, ensayos y dominio. La guía 60-600 se extiende con al menos un ensayo cuyo σ′_{3} se acerque al de las "
       "superficies críticas. Para σ′_{n} > 8 MPa la envolvente se extrapola con el mismo m, sin piso de estado crítico, "
       "salvo que ensayos del sitio sobre 5 MPa justifiquen calibrar una de las formas de la sección 3.4.")

titulo("5.3 Actualización del criterio", 2)
par("Leps 3.0 se define por sus parámetros poblacionales (μ, τ, el exponente central y la Tabla 3), no por valores fijos. "
    "Cuando se publiquen ensayos con σ′_{3} sobre 5 MPa, estos parámetros se recalculan con el mismo procedimiento, y las "
    "formas con piso de la sección 3.4 pueden volver a evaluarse con la validación cruzada de la sección 3.3. La "
    "implementación de referencia (módulo en Python con pruebas) calcula la envolvente, σ′_{1}(σ′_{3}), la envolvente de "
    "diseño, la actualización de m y los parámetros para el modelo de potencia o de función de Slide2.")

# ============================================================== 6
titulo("6. Aplicación a un botadero de 700 m")
titulo("6.1 Estabilidad drenada", 2)
par(f"Sobre fundación horizontal, el círculo crítico alcanza 160–180 m de profundidad y σ′_{{n}} de hasta 2.8–3.1 MPa. Con "
    f"φ′_{{1}} = 41°, el factor de seguridad es {f2(ej(41, 'Leps 2.0 (m = 0.9)'))} con Leps 2.0, "
    f"{f2(ej(41, 'Leps 3.0 central (m = 0.87)'))} con la envolvente central de Leps 3.0 y "
    f"{f2(ej(41, 'Leps 3.0 diseño P10'))} con su envolvente de diseño P10 (Tabla 4). El cambio de m reduce el FS en 2 %; "
    "la envolvente de diseño, en 12 %. En la ladera de 10°, la búsqueda no circular de Slide2 encontró una superficie casi "
    "circular, con σ′_{n} de hasta 6.4 MPa en el mecanismo basal, que solo pasa a ser crítico con agua en la base.")
filas = []
for p1 in (41, 38):
    filas.append([f"{p1}°", f2(ej(p1, "Leps 2.0 (m = 0.9)")), f2(ej(p1, "Leps 3.0 central (m = 0.87)")),
                  f2(ej(p1, "Leps 3.0 diseño P20")), f2(ej(p1, "Leps 3.0 diseño P10")),
                  f"{f2(s2(f'Leps 2.0 phi1 {p1}', 'bishop simplified'))} / {f2(s2(f'm 0.87 phi1 {p1}', 'bishop simplified'))} / "
                  f"{f2(s2(f'L3 diseno phi1 {p1}', 'bishop simplified'))}"])
tabla(4, "Factor de seguridad del botadero genérico de 700 m sobre fundación horizontal, método de Bishop",
      ["φ′_{1}", "Leps 2.0", "Leps 3.0 central", "Leps 3.0 diseño P20", "Leps 3.0 diseño P10",
       "Slide2: Leps 2.0 / central / P10"],
      filas, [1.2, 1.8, 2.3, 2.5, 2.5, 4.4],
      "Fuente: elaboración propia. Python con el método de Bishop en las columnas 2–5 y Slide2 9.043 (Rocscience, s. f.) en la última.")

titulo("6.2 Presión de poros", 2)
par(f"En el mecanismo basal sobre la ladera de 10°, una capa basal saturada de 30 m con razón de exceso r_{{u}} = 0.2 "
    f"reduce el FS de {f2(WG[chr(76)+'eps 2.0 φ'+chr(39)+'1=41° m=0.87|hw=0|ru=0']['F'])} a "
    f"{f2(WG[chr(76)+'eps 2.0 φ'+chr(39)+'1=41° m=0.87|hw=30|ru=0.2']['F'])} con la envolvente central de Leps 3.0, y "
    f"un contacto basal con respuesta no drenada (S_{{u}} = 0.28 σ′_{{v0}}) lo deja en "
    f"{f2(WG['Base no drenada Su/σ'+chr(39)+'v=0.28 (resto φ'+chr(39)+'1=41° m=0.87)|hw=30|ru=0']['F'])} (método "
    "simplificado de Janbu sin factor de corrección). Ese valor es probablemente optimista, porque la razón medida en "
    "consolidación isótropa equivale a una menor referida a σ′_{v} en condición K_{0}. La presión de poros por "
    "consolidación durante el crecimiento del botadero, calculada con la ecuación de Terzaghi con carga creciente, no "
    "supera 0.4 % de σ_{v} en una capa basal saturada de 30 m con k ≥ 10^−7^ m/s, aun con 100 m/año; para llegar a "
    "Δu/σ_{v} = 0.1 se requiere k ≈ 10^−9^ m/s (Figura 5). El riesgo de presión de poros viene de la saturación de la base "
    "y de la respuesta contractiva del material, no de la consolidación, como en los deslizamientos por licuación de "
    "botaderos de carbón de 100–400 m (Dawson et al., 1998; Hungr et al., 2002).")
figura(5, "fig5_presion_poros_consolidacion.png",
       "Exceso de presión de poros por consolidación en la base de una capa saturada de 30 m al terminar un botadero de 700 m",
       "Fuente: elaboración propia. M = 50 MPa; con M = 150 MPa los valores se dividen por tres.")

# ============================================================== 7
titulo("7. Discusión")
par("Leps 3.0 no cambia la forma funcional de Leps 2.0 y no lo hace por decisión metodológica: la comparación fuera de "
    "muestra no favorece las formas con piso de estado crítico ni la curvatura variable con los datos disponibles. Los "
    "cambios que sí respaldan los datos son tres: el exponente central, la cuantificación de la incertidumbre según el "
    "esfuerzo y el uso de los ensayos del sitio para ajustar m. En el ejemplo, la envolvente de diseño reduce el FS en "
    "10 % respecto de la central y el exponente en 2 %: la incertidumbre pesa más, y Leps 2.0 no la entregaba de forma "
    "operativa.")
par("Las limitaciones principales son cinco:")
vineta("Los ensayos con σ′_{3} ≥ 3 MPa provienen de cuatro publicaciones de 1963–1983, y la base no contiene "
       "esfuerzos de la mitad inferior de un botadero de 700 m.")
vineta("Los ensayos usan muestras escalpadas o de granulometría paralela con D_{max} ≤ 200 mm; su correspondencia con "
       "el material de campo depende de la razón entre el diámetro de la muestra y D_{max} (Girumugisha et al., 2024).")
vineta("La Tabla 3 proviene del modo A. Con ensayos del sitio las bandas son más estrechas: en el modo B, el percentil "
       "10 de q_{obs}/q_{pred} es 0.89 para σ′_{n} de 2–3 MPa y 0.95 para 3–8 MPa.")
vineta("La base no tiene ensayos no drenados, de modo que el paso 5 depende de ensayos del sitio.")
vineta("El efecto de la resistencia de la partícula (UCS) no es concluyente con esta muestra.")
par("La agenda que define la vigencia de Leps 3.0 es concreta: ensayos triaxiales drenados de gran escala con σ′_{3} de "
    "5–10 MPa en desmonte de distintas litologías, ensayos no drenados a alta presión sobre granulometrías trituradas, "
    "permeabilidad después de la rotura y monitoreo piezométrico en la base de botaderos altos. Cada uno de esos datos "
    "entra en el criterio por la vía de la sección 5.3.")

# ============================================================== 8
titulo("8. Conclusiones")
vineta("Las ecuaciones cerradas de Leps 2.0 son correctas. La aproximación explícita de su Ecuación 23 tiene errores de "
       "hasta 7.4 % en σ′_{1} dentro del rango declarado; conviene usar la solución iterativa.")
vineta("Con m = 0.9, Leps 2.0 sobreestima la resistencia a alta presión entre 4.5 % y 10.4 % en promedio, según el modo "
       "de uso, con una cola de sobreestimación de 16–26 % en el percentil 90.")
vineta("Tres estimaciones independientes dan m ≈ 0.87. Con ese valor, el sesgo baja a −2.6 % a +2.8 % y la cola de "
       "sobreestimación a alta presión se reduce a la mitad en el modo predictivo.")
vineta("Con datos hasta 7.9 MPa no se puede calibrar un piso de estado crítico, y la curvatura dependiente de la densidad, "
       "aunque presente, no mejora la predicción.")
vineta("Leps 3.0 combina m = 0.87, una distribución poblacional de m para actualizar con ensayos del sitio, reducciones "
       "de diseño de φ′_{1} según el esfuerzo, un dominio de aplicación explícito, un criterio de drenaje y una "
       "implementación abierta que permite recalibrarlo con datos nuevos.")
vineta("En botaderos de 700 m, la condición de drenaje de la base pesa más que la forma de la envolvente drenada.")

titulo("Disponibilidad de datos y programas")
par("La base de datos es pública (Arrieta, 2025). La implementación de referencia de Leps 3.0, los programas de análisis "
    "y el modelo de Slide2 están disponibles a solicitud de los autores [[repositorio público por definir]].", sangria=False)

# ============================================================== referencias
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
titulo("Referencias")
REFS = [
    "Arrieta, M. (2025). *Rockfill test database* [Base de datos]. Zenodo. https://doi.org/10.5281/zenodo.14665504",
    "Barton, N. y Kjærnsli, B. (1981). Shear strength of rockfill. *Journal of the Geotechnical Engineering Division*, *107*(7), 873–891. https://doi.org/10.1061/AJGEB6.0001167",
    "Bishop, A. W. (1955). The use of the slip circle in the stability analysis of slopes. *Géotechnique*, *5*(1), 7–17. https://doi.org/10.1680/geot.1955.5.1.7",
    "Bolton, M. D. (1986). The strength and dilatancy of sands. *Géotechnique*, *36*(1), 65–78. https://doi.org/10.1680/geot.1986.36.1.65",
    "Coop, M. R., Sørensen, K. K., Bodas Freitas, T. y Georgoutsos, G. (2004). Particle breakage during shearing of a carbonate sand. *Géotechnique*, *54*(3), 157–163. https://doi.org/10.1680/geot.2004.54.3.157",
    "Cylwik, S. D., Arrieta, M., Cunning, J. C. y Hawley, P. M. (2025). Update to the Leps strength envelopes for rockfill based on 900 large-scale triaxial tests. *International Journal of Geomechanics*, *25*(12), 04025285. https://doi.org/10.1061/IJGNAI.GMENG-12050",
    "Dawson, R. F., Morgenstern, N. R. y Stokes, A. W. (1998). Liquefaction flowslides in Rocky Mountain coal mine waste dumps. *Canadian Geotechnical Journal*, *35*(2), 328–343. https://doi.org/10.1139/t98-009",
    "Girumugisha, G., Ovalle, C. y Ouellet, S. (2024). Grading scalping and sample size effects on critical shear strength of mine waste rock through laboratory and in-situ testing. *International Journal of Rock Mechanics and Mining Sciences*, *183*, 105915. https://doi.org/10.1016/j.ijrmms.2024.105915",
    "Hungr, O., Dawson, R. F., Kent, A., Campbell, D. y Morgenstern, N. R. (2002). Rapid flow slides of coal-mine waste in British Columbia, Canada. En S. G. Evans y J. V. DeGraff (Eds.), *Catastrophic landslides: Effects, occurrence, and mechanisms* (pp. 191–208). Geological Society of America. https://doi.org/10.1130/REG15-p191",
    "Leps, T. M. (1970). Review of shearing strength of rockfill. *Journal of the Soil Mechanics and Foundations Division*, *96*(4), 1159–1170. https://doi.org/10.1061/JSFEAQ.0001433",
    "Linero, S., Palma, C. y Apablaza, R. (2007). Geotechnical characterisation of waste material in very high dumps with large scale triaxial testing. En *Proceedings of the 2007 International Symposium on Rock Slope Stability in Open Pit Mining and Civil Engineering* (pp. 59–75). Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/708_2",
    "Ovalle, C., Linero, S., Dano, C., Bard, E., Hicher, P.-Y. y Osses, R. (2020). Data compilation from large drained compression triaxial tests on coarse crushable rockfill materials. *Journal of Geotechnical and Geoenvironmental Engineering*, *146*(9), 06020013. https://doi.org/10.1061/(ASCE)GT.1943-5606.0002314",
    "Rocscience. (s. f.). *Slide2* (versión 9.043) [Programa de cómputo]. Rocscience Inc.",
    "Seabold, S. y Perktold, J. (2010). Statsmodels: Econometric and statistical modeling with Python. En *Proceedings of the 9th Python in Science Conference* (pp. 92–96). https://doi.org/10.25080/Majora-92bf1922-011",
    "Spencer, E. (1967). A method of analysis of the stability of embankments assuming parallel inter-slice forces. *Géotechnique*, *17*(1), 11–26. https://doi.org/10.1680/geot.1967.17.1.11",
    "Valenzuela, L., Bard, E., Campaña, J. y Anabalón, M. E. (2008). High waste rock dumps: Challenges and developments. En *Proceedings of the First International Seminar on the Management of Rock Dumps, Stockpiles and Heap Leach Pads* (pp. 65–78). Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/802_6",
    "Yamamuro, J. A. y Lade, P. V. (1996). Drained sand behavior in axisymmetric tests at high pressures. *Journal of Geotechnical Engineering*, *122*(2), 109–119. https://doi.org/10.1061/(ASCE)0733-9410(1996)122:2(109)",
]
for ref in REFS:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0)
    runs(p, ref)
out = AQUI / "Manuscrito_Leps3.docx"
doc.save(str(out)); print("escrito", out)
