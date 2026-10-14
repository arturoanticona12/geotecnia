"""Leps-P manuscript, round 2 (ASCE Int. J. Geomech., Technical Paper). All numbers read from result files.
ASCE format: double spacing, continuous line numbers, unnumbered headings, equations as Word equation objects (OMML),
figure caption list at the end followed by tables, ASCE author–date references."""
import json, math, re, sys
from pathlib import Path
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from lxml import etree
import latex2mathml.converter as l2m

AQ = Path(__file__).parent; R = AQ.parent; AN = R / "analisis"; AA = R / "analisis_antamina"
RES = Path(r"H:\Mi unidad\slide2-bench-intercambio\resultados")
J = lambda p: json.load(open(p, encoding="utf-8"))
PAR = J(R / "leps3" / "lepsp_parametros.json"); FX = PAR["formas"]["x2"]; MC2 = J(AN / "lepsp_mcmc2.json")
MET = J(AN / "lepsp_cv_metricas.json"); M2 = J(AN / "lepsp_cv2_metricas.json"); ESC = J(AN / "lepsp_escala2.json")
SIT = J(AA / "lepsp_sitio2.json"); EST = J(AN / "lepsp_estabilidad2.json")
S62 = {}   # lote 62: casos no drenados (sin malla de presiones); lote 63: barrido drenado con la malla solo en la capa basal
for lote, pref in (("62-lepsp-presion-poros", "Su "), ("63-lepsp-presion-poros-capa", "R")):
    try:
        for mod, v in J(RES / lote / "fs_por_escenario.json").items():
            S62.update({k_: v_ for k_, v_ in v.get("fs", {}).items() if k_.startswith(pref)})
    except FileNotFoundError:
        print(f"AVISO: faltan resultados de Slide2 (lote {lote})")
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))

# ------------------------------------------------------------------ numbers
b = dict(zip(FX["nombres"], FX["beta"])); SE = dict(zip(FX["nombres"], np.sqrt(np.diag(FX["cov_beta"]))))
H = np.array([[math.exp(p[0]), math.exp(p[1]), math.tanh(p[2]), math.exp(p[3]), math.exp(p[4]), math.tanh(p[5]), math.exp(p[6]), p[7]]
              for p in (np.array(s["psi"]) for s in FX["muestras"])])
HP = [(np.median(H[:, i]), np.percentile(H[:, i], 5), np.percentile(H[:, i], 95)) for i in range(8)]
deg = lambda th: math.degrees(math.atan(math.exp(th)))
phi_e0 = {c: deg(b[c]) for c in ("good", "moderate", "low")}
omega = -b["e"] * math.sin(math.radians(41)) * math.cos(math.radians(41)) * 180 / math.pi
fl_th = HP[0][0] ** 2 / (HP[0][0] ** 2 + HP[3][0] ** 2); fl_m = HP[1][0] ** 2 / (HP[1][0] ** 2 + HP[4][0] ** 2)
fb = (HP[0][0] ** 2 + HP[3][0] ** 2) / (HP[0][0] ** 2 + HP[3][0] ** 2 + HP[6][0] ** 2)
N = PAR["n"]
A_all = MET["pub|M2|todo"]; A_hi = MET["pub|M2|σ′n>2.75"]; ALR = MET["Material Type|M0|todo"]; AU = MET["pub|M5|todo"]
BN = M2["B|nivel"]; BJ = M2["B|conjunta"]; BI86 = M2["B|ingenuo m=0.86"]; BI90 = M2["B|ingenuo m=0.9"]; RVI = M2["regla_vs_ingenuo086"]
JJ = M2["J|conjunta"]; JN = M2["J|nivel"]; AMIX = M2["A|clase desconocida"]
import pandas as pd  # noqa: E402
_b = pd.read_csv(AN / "lepsp_cv2_B.csv"); _b = _b[_b.variante == "nivel"]          # previa del modo B: muA, sdA
_z = ((_b.y - _b.muA) / _b.sdA).to_numpy(); _Phi = np.vectorize(lambda v: 0.5 * (1 + math.erf(v / math.sqrt(2))))(_z)
_crps = _b.sdA.to_numpy() * (_z * (2 * _Phi - 1) + 2 * np.exp(-_z ** 2 / 2) / math.sqrt(2 * math.pi) - 1 / math.sqrt(math.pi))
PRB = dict(n=len(_z), cob80=float(np.mean(np.abs(_z) < 1.2816)), bajo10=float(np.mean(_z < -1.2816)), crps=float(_crps.mean()))
print("previa modo B:", PRB, "| crps nivel + mejora:", BN["crps"] + BN["mejora_vs_previa"][0])
thr = [M2[f"B|regla r{r} p{p}"]["crps"] for r in (2, 3, 4) for p in (0.01, 0.05, 0.1)]
EG = ESC["todos"]; EP = ESC["similares Cu≤1.5 Δe≤0.05"]; MR = ESC["metaregresion"]
SNs = np.array(SIT["sn"]); isn = lambda s: int(np.argmin(np.abs(SNs - s)))
SC = SIT["escenarios"]
pc = lambda v: f"{100 * v:.0f}%"; f2 = lambda v: f"{v:.2f}"; f3 = lambda v: f"{v:.3f}"; f4 = lambda v: f"{v:.4f}"
neg = lambda s: s.replace("-", "−")
sg = lambda v, d=4: neg(f"{v:+.{d}f}")
ci = lambda a, d=4: f"{sg(a[1], d)} to {sg(a[2], d)}"
DE = lambda k, h, nm: EST["det"][f"{k}|{h}|{nm}"]
MCx = lambda k, nm: EST["mc"][f"{k}|700|{nm}"]
KMOD = "model"; KCOM = "comonotonic"
kweak = [k for k in EST["mc"] if k.startswith("e=0.65|700|weakest (")][0].split("|")[-1]
kweakt = [k for k in EST["mc"] if k.startswith("e=0.65|700|weakest tested")][0].split("|")[-1]
used = [s for s in SIT["sets"] if s["usado"]]; notused = [s for s in SIT["sets"] if not s["usado"]]
SITB = J(AA / "lepsp_sitio2_marmol_buena.json")                       # sensibilidad: mármol en clase «buena»
pmb = [s["p_conflicto"] for s in SITB["sets"] if s["lit"] == "Mármol"][0]
_mb, _m0 = SITB["escenarios"]["e=0.65|general"]["botadero"]["model"], SIT["escenarios"]["e=0.65|general"]["botadero"]["model"]
dmb50 = _mb["P50"][isn(3)] - _m0["P50"][isn(3)]; dmb05 = _mb["P05"][isn(3)] - _m0["P05"][isn(3)]


def s62(R_, k, esc="P05", met="gle/morgenstern-price"):
    key = f"R{R_} k{k:.0e} / {esc}"
    return S62.get(key, {}).get(met, float("nan"))


def s62u(r, R_, k, fam):
    key = f"Su r{r} R{R_} k{k:.0e} {fam} / P50"
    return S62.get(key, {}).get("gle/morgenstern-price", float("nan"))


# ------------------------------------------------------------------ document setup
doc = Document()
cp = doc.core_properties; cp.author = ""; cp.title = "Leps-P"; cp.comments = ""; cp.last_modified_by = ""; cp.subject = ""; cp.keywords = ""
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.line_spacing = 2.0; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
for m_ in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m_, Inches(1))
ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); sec._sectPr.append(ln)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    r_ = fp.add_run()
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r_._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r_._r.append(e)


def omml(latex):
    return XSL(etree.fromstring(l2m.convert(latex).encode())).getroot()


SYM = [("b^(p)", r"\mathbf{b}^{(p)}"), ("b^(s)", r"\mathbf{b}^{(s)}"), ("C_β", r"\mathbf{C}_\beta"), ("β̂", r"\hat{\boldsymbol{\beta}}"),
       ("β_c", r"\beta_c"), ("β_e", r"\beta_e"), ("β_m", r"\beta_m"), ("β_D", r"\beta_D"), ("G_p", r"\mathbf{G}_p"), ("G_s", r"\mathbf{G}_s"),
       ("σε²", r"\sigma_\varepsilon^2"), ("σε", r"\sigma_\varepsilon"), ("φ′e=0", r"\phi'_{e=0}"),
       ("φ′1", r"\phi'_1"), ("φ′s", r"\phi'_s"), ("σ′n", r"\sigma'_n"), ("σ′3", r"\sigma'_3"), ("σ′v", r"\sigma'_v"), ("Δu/σv", r"\Delta u/\sigma_v"),
       ("σv", r"\sigma_v"), ("σ0", r"\sigma_0"), ("p_a", r"p_a"), ("*D*max", r"D_{max}"), ("*D*f", r"D_f"), ("*D*l", r"D_l"),
       ("*x*i", r"x_i"), ("*y*i", r"y_i"), ("*x̄*", r"\bar{x}"), ("x̃", r"\tilde{\mathbf{x}}"), ("**z**", r"\mathbf{z}"),
       ("Su = ", r"S_u="), ("Su/σ′v", r"S_u/\sigma'_v"), ("K0", r"K_0"), ("c2", r"c_2"), ("a0, a1", r"a_0, a_1"), ("D², ", r"D^2, ")]


def simbolos(text):
    parts = re.split(r"(\$[^$]+\$|https?://\S+)", text); out = []
    for s in parts:
        if s.startswith("$") or s.startswith("http"):
            out.append(s); continue
        for a_, l_ in SYM:
            s = s.replace(a_, f"${l_}$")
        out.append(s)
    return "".join(out)


def runs(p, text, _sym=True):
    if _sym:
        text = simbolos(text)
    for tok in re.split(r"(\$[^$]+\$|\*[^*]+\*|\[\[[^\]]+\]\])", text):
        if not tok:
            continue
        if tok.startswith("$"):
            p._p.append(parse_xml(etree.tostring(omml(tok[1:-1]))))
        elif tok.startswith("*") and "$" in tok:
            runs(p, tok[1:-1], False)
        elif tok.startswith("*"):
            p.add_run(tok[1:-1]).italic = True
        elif tok.startswith("[["):
            rr = p.add_run(tok[2:-2]); rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p.add_run(tok)
    return p


def P(text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    return runs(p, text)


def Hd(text, level=1):
    p = doc.add_paragraph(); rr = p.add_run(text); rr.bold = level == 1; rr.italic = level == 2; p.paragraph_format.keep_with_next = True


NEQ = [0]


def EQ(latex):
    NEQ[0] += 1; p = doc.add_paragraph(); pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(3.25), alignment=1); pf.tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    p.add_run("\t"); p._p.append(parse_xml(etree.tostring(omml(latex)))); p.add_run(f"\t({NEQ[0]})")


def br():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================== TITLE, ABSTRACT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Leps-P: A Probabilistic and Updatable Extension of the Leps 2.0 Shear Strength Envelopes for Rockfill and Mine Waste Rock").bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, "[[Author names, ASCE grade, affiliations, and corresponding author e-mail to be completed.]]")
Hd("Abstract")
P("The Leps 2.0 envelopes, calibrated on a public database of more than 900 large-scale triaxial tests, are the standard empirical "
  "estimate of the drained shear strength of rockfill and mine waste rock. Reliability-based design of waste dumps up to 700 m "
  "high also needs the uncertainty of that estimate, its dependence on stress, and a rule to combine it with site tests. "
  "Leps-P is a three-level linear mixed model of the Leps power envelope in which the reference secant friction angle and the "
  "curvature vary between publications and between materials, the within-material dispersion depends on stress, and variance "
  f"parameters are sampled by Markov chain Monte Carlo. It was calibrated on {N['ensayos']} tests from {N['pub']} publications and "
  f"evaluated by leaving out one publication at a time. Without site tests, the 80% predictive interval covered {pc(A_all['m']['cob80'])} "
  "of held-out tests, with an accuracy equal to that of the Leps 2.0 envelope with constant dispersion. With a few site tests "
  "at low stress, updating only the level of the envelope reduced the continuous ranked probability score at higher stress from "
  f"{f3(BN['crps'] + BN['mejora_vs_previa'][0])} to {f3(BN['crps'])}; updating the curvature as well was overconfident, and a "
  f"least-squares fit with the curvature fixed at 0.86 performed similarly. Between-publication differences accounted for {pc(fl_th)} "
  "of the between-material variance of the reference angle. Within-material comparisons showed a heterogeneous specimen-size "
  "effect, negative only for nearly parallel gradations. For waste rock from a polymetallic skarn deposit, the placement void "
  "ratio, untested lithologies, and the pore pressure generated in a saturated basal layer by rapid loading controlled the "
  "factor of safety of a 700 m dump more than the extrapolation of the envelope beyond the data. The model, code, and parameters "
  "can be re-estimated as tests at higher stress become available.", indent=False)
Hd("Practical Applications")
P("Designers of high waste dumps and rockfill dams often have few or no large-scale shear tests of their own and rely on "
  "published charts of rockfill strength. This work turns the most widely used chart into a tool that gives a best estimate of "
  "strength together with its range, shows how that range grows at the high stresses found under tall dumps, and shows how "
  "much it narrows when a few laboratory tests of the site material are added. The tool also warns the user when site tests "
  "disagree with published experience. An example for a very high dump shows that the safety margin depends strongly on how "
  "loosely the waste rock is placed, on rock types that have not been tested, and on keeping the base of the dump drained "
  "during construction. Fast advance of the dump over a saturated base of low permeability can cause failure even when the "
  "strength of the waste rock is well known. Testing every major rock type at large scale, measuring the density of the placed "
  "material, and controlling the rate of construction over saturated ground are the most effective actions. Free software "
  "accompanies the paper.", indent=False)
P("*Author keywords:* Rockfill; Mine waste rock; Shear strength; Nonlinear failure envelope; Hierarchical model; "
  "Reliability-based design; High waste dumps; Pore pressure.", indent=False)

# ================================================================== INTRODUCTION
Hd("Introduction")
P("Leps (1970) compiled 109 large-scale triaxial tests and proposed three nonlinear envelopes for rockfill. For five decades, these "
  "curves have been used for dams and waste dumps whenever site-specific tests were unavailable. Cylwik et al. (2025) updated them "
  "with a public database of 938 tests (Arrieta 2025), adopted the power criterion of De Mello (1977) and Charles and Watts (1980) "
  "with a single fitting parameter, the secant friction angle at σ′n = 10p_a, and related that angle to rock quality and void "
  "ratio. That update is referred to as Leps 2.0.")
P("Two developments in mining practice motivate a further extension. First, waste dumps now reach 500–700 m (Valenzuela et al. "
  "2008), so effective normal stresses at their base can exceed the 7.9 MPa reached in the database, in a material that crushes, "
  "loses dilatancy and permeability, and may respond contractively (Linero et al. 2007; Ovalle et al. 2020); loose, saturated "
  "waste on weak foundations has produced rapid flow slides (Hungr et al. 2002). Second, design is "
  "moving toward reliability-based criteria (Read and Stacey 2009; Hawley and Cunning 2017), which require a strength "
  "distribution and a rule to combine published experience with site tests.")
P("Hierarchical models that borrow strength from generic databases have been developed for soil properties (Ching and Phoon "
  "2019; Ching et al. 2021), and a hierarchical Bayesian analysis of the Barton and Kjærnsli (1981) parameters has been applied "
  "to rockfill groups of a single project (Contreras et al. 2022; Contreras and Linero-Molina 2025). The novelty of this work lies "
  "not in the hierarchical framework itself but in (1) its calibration on the Leps 2.0 database with random reference angle and "
  "curvature; (2) the separation of publication and material effects; (3) out-of-sample evaluation by publication, including a "
  "comparison with the simple fits used in practice; (4) a validated site-updating rule with a prior–data conflict check; and "
  "(5) explicit treatment of specimen size, placement density, and pore pressure in the design of high dumps. Leps-P is an "
  "independent extension built on the public Leps 2.0 database; it is not affiliated with the authors of Leps 2.0.")

# ================================================================== MODEL
Hd("Database and Model Formulation")
Hd("Calibration Database", 2)
P(f"The public database (Arrieta 2025) was filtered to tests with σ′3 ≥ 60 kPa, reported initial void ratio, and at least two tests "
  f"per set, giving {N['ensayos']} tests on {N['sets']} sets from {N['pub']} publications (Fig. 1). A set is a material with a given "
  "gradation, density, and rock type tested in one program; failure is the maximum deviator stress. Rock quality classes follow "
  "Table 5 of Cylwik et al. (2025). Two features of the database limit inference: tests with σ′n > 3 MPa (90 tests) come from eight "
  "publications, seven of them published between 1963 and 1983 (e.g., Marachi et al. 1972; Marsal 1973), and sets with φ′1 < 31° were excluded by Cylwik et al. (2025), so the lower tail of the population is truncated "
  "and predictive lower percentiles for weak materials may be optimistic.")
Hd("Hierarchical Envelope", 2)
P("The Leps 2.0 criterion is the power law")
EQ(r"\tau=\tan\phi'_1\,\sigma_0^{1-m}\,\sigma_n'^{\,m}")
P("where σ0 = 10p_a, p_a = 0.1013 MPa is atmospheric pressure, φ′1 is the secant friction angle at σ′n = σ0, and *m* is the "
  "curvature exponent; the secant friction angle at any stress is φ′s = arctan(τ/σ′n). With *x* = ln(σ′n/σ0) and *y* = ln(τ/σ0), "
  "Eq. (1) is a straight line with intercept θ = ln tan φ′1 and slope *m*. Leps-P treats each set *j* of publication *k* as a "
  "realization of a population:", indent=False)
EQ(r"y_{ijk}=\theta_{jk}+m_{jk}x_{ijk}+c_2x_{ijk}^2+\varepsilon_{ijk}")
EQ(r"\theta_{jk}=\beta_{c(jk)}+\beta_e e_{jk}+b_{0k}^{(p)}+b_{0jk}^{(s)},\quad m_{jk}=\beta_m+b_{1k}^{(p)}+b_{1jk}^{(s)}")
EQ(r"\mathbf{b}^{(p)}\sim N(\mathbf{0},\mathbf{G}_p),\quad\mathbf{b}^{(s)}\sim N(\mathbf{0},\mathbf{G}_s),\quad\ln\sigma_\varepsilon(x)=a_0+a_1x")
P("where β_c is the intercept of rock quality class *c*, β_e the effect of initial void ratio *e*, β_m the mean curvature, c2 a "
  "population term that lets the local slope vary with stress, b^(p) and b^(s) the publication and set random effects on (θ, *m*) "
  "with covariance matrices G_p and G_s, and ε the dispersion of tests on one material (Laird and Ware 1982). The publication "
  "effect combines laboratory, procedure, era, and material selection, which cannot be separated with these data; it plays the "
  "role of measurement and transformation uncertainty in the classification of Phoon and Kulhawy (1999). Through β_c and "
  "β_e, the median φ′1 reproduces the relation φ′1 = φ′e=0 − ωe of Leps 2.0, where ω is the decrease in φ′1 per unit void ratio.", indent=False)
Hd("Predictive Distribution", 2)
P("For a new material of class *c* and void ratio *e* tested by a new laboratory, *y* is normal with")
EQ(r"E[y]=\beta_c+\beta_e e+\beta_m x+c_2x^2,\quad\mathrm{Var}[y]=\mathbf{z}^T(\mathbf{G}_p+\mathbf{G}_s)\mathbf{z}+\tilde{\mathbf{x}}^T\mathbf{C}_\beta\tilde{\mathbf{x}}+\sigma_\varepsilon^2(x)")
P("with **z** = (1, *x*)ᵀ, x̃ the covariate vector, and C_β the covariance of the fixed effects. The first two terms describe the "
  "uncertainty of the mean envelope of the material, which does not average out along a failure surface within that material "
  "(Christian et al. 1994); the last describes test-to-test scatter. Variance parameters are sampled from their posterior and "
  "fixed effects from N(β̂, C_β) for each sample. Samples whose local slope *m* + 2c2*x* exceeds one within the calibrated stress "
  "range are discarded, because a convex envelope has no support in granular materials; the fraction discarded is reported.", indent=False)
Hd("Updating with Site Tests", 2)
P("Site triaxial tests are converted to (σ′n, τ) with Eqs. (13), (14), and (16) of Cylwik et al. (2025). Consistency with the prior "
  "is checked with the prior predictive Mahalanobis distance")
EQ(r"D^2=\mathbf{r}^T\left(\mathbf{Z}(\mathbf{G}_p+\mathbf{G}_s)\mathbf{Z}^T+\tilde{\mathbf{X}}\mathbf{C}_\beta\tilde{\mathbf{X}}^T+\mathbf{R}\right)^{-1}\mathbf{r}\sim\chi^2_n")
P(r"where $\mathbf{r}$ is the vector of residuals of the *n* site tests with respect to the prior mean, $\mathbf{Z}$ has rows (1, *x*i), "
  r"and $\mathbf{R}=\mathrm{diag}\,\sigma_\varepsilon^2(x_i)$. A *p*-value below 0.05 flags the tests for review. Tests that pass are "
  "used to update the level of the envelope at their centroid *x̄* (the mean of *x*i): the random-effect combination θ + *m**x̄* is "
  "updated with the exact Gaussian posterior and the slope keeps its prior conditional distribution, so the envelope pivots about "
  "the tested range. This rule was chosen after the out-of-sample comparison reported below. When several materials of one site "
  "were tested by the same laboratory, their publication-level effect is shared and all sets are updated jointly.", indent=False)
Hd("Specimen Size and Placement Density", 2)
P("Weibull scaling (Frossard et al. 2012; Ovalle et al. 2014) shifts the envelope of a material with maximum particle size *D*f "
  "with respect to tests at size *D*l by")
EQ(r"\Delta\theta=\beta_D\ln\left(D_f/D_l\right),\quad\beta_D=-\kappa\,(1-m)")
P("where κ is the ratio of a size exponent to the Weibull modulus. β_D was estimated within materials of the database tested at "
  "several sizes. The void ratio of placed waste rock, which differs from that of compacted specimens, enters through β_e.", indent=False)

# ================================================================== ESTIMATION & EVALUATION
Hd("Estimation and Out-of-Sample Evaluation")
P("Fixed and random effects were integrated analytically and the restricted likelihood was evaluated with publication blocks. "
  "Variance parameters were sampled with an adaptive Metropolis algorithm under weakly informative priors (half-normal with scale "
  "0.25 on standard deviations, uniform on correlations, and normal on a0 and a1), with four chains started at dispersed values "
  f"(split potential scale reduction ≤ {f3(max(MC2['x2']['rhat']))}; effective sample size ≥ {min(MC2['x2']['ess']):.0f}; Gelman et al. 2013). "
  "A quadratic form (Eq. 2) and a power form (c2 = 0) were fitted.")
P("Each publication was left out in turn and its tests predicted. In Mode A (no site tests), the predictive distribution of "
  "Eq. (5) was scored. In Mode B, each held-out set with at least two tests at σ′n ≤ 1.5 MPa and one at σ′n ≥ 2 MPa was updated with "
  "the former and the latter were predicted; competing procedures were joint updating of (θ, *m*), level-only updating, and "
  "least-squares fits of φ′1 to the site tests with *m* fixed at 0.90 or 0.86 and a predictive standard deviation equal to their "
  "out-of-sample root-mean-square error. Scores were the coverage of the 80% interval, the fraction of tests below the 10th "
  "percentile (nominally 10% for a calibrated 10th-percentile envelope), and the continuous ranked probability score (CRPS) of "
  "ln τ (Gneiting and Raftery 2007). Confidence intervals were obtained by bootstrapping publications. The reference for Mode A was "
  "the published Leps 2.0 envelope (*m* = 0.9) with a constant dispersion; because its constants were fitted to the whole database, "
  "this comparison favors the reference.")

# ================================================================== RESULTS
Hd("Results")
Hd("Parameters", 2)
P(f"Table 1 lists the parameters. The mean curvature is β_m = {f3(b['x'])} and c2 = {f4(b['x2'])} ± {f4(SE['x2'])}, so the local "
  "slope rises from about 0.84 at 0.2 MPa to about 0.89 at 15 MPa, qualitatively consistent with Charles and Watts (1980) and "
  f"Indraratna et al. (1993). The void-ratio effect is about {omega:.0f}° per unit of *e* at φ′1 ≈ 41°, compared with 20° in Leps 2.0. "
  f"Publication effects account for {pc(fl_th)} of the between-set variance of θ and {pc(fl_m)} of that of *m*. At σ0, the "
  f"between-set and between-publication terms make up {pc(fb)} of the variance excluding fixed-effect uncertainty.")
Hd("Mode A: No Site Tests", 2)
P(f"Leaving out each publication, the 80% interval covered {pc(A_all['m']['cob80'])} of the tests (95% confidence interval "
  f"{pc(A_all['ci']['cob80'][0])}–{pc(A_all['ci']['cob80'][1])}) and {pc(A_all['m']['bajo10'])} fell below the 10th percentile (Fig. 3, "
  f"Table 2). The CRPS of Leps-P and of the reference did not differ significantly (difference {sg(A_all['dcrps'][0])}, 95% CI "
  f"{sg(A_all['dcrps'][1])} to {sg(A_all['dcrps'][2])}): without site tests, Leps-P does not predict the strength of a new material "
  "better than Leps 2.0 with a constant dispersion, but it provides a stress-dependent description of uncertainty. Above 2.75 MPa "
  f"the intervals were conservative ({pc(A_hi['m']['cob80'])} coverage, {pc(A_hi['m']['bajo10'])} below the 10th percentile, 95% CI "
  f"{pc(A_hi['ci']['bajo10'][0])}–{pc(A_hi['ci']['bajo10'][1])}), mainly because the predicted median was low for some publications, "
  "and the 10% quantile loss favored the reference. Leaving out each material type gave similar scores "
  f"({pc(ALR['m']['cob80'])} and {pc(ALR['m']['bajo10'])}); this does not rule out leakage through the class labels, which Cylwik et al. "
  f"(2025) assigned from measured strength. When the class was treated as unknown (a mixture of the three classes), the CRPS rose "
  f"from {f4(A_all['m']['crps'])} to {f4(AMIX['crps'])}, which bounds the value of the class information.")
Hd("Mode B: Value of Site Tests", 2)
P(f"In Mode B ({BN['n']} held-out tests from 13 publications), level-only updating reduced the CRPS from "
  f"{f4(BN['crps'] + BN['mejora_vs_previa'][0])} to {f4(BN['crps'])} (difference {sg(BN['mejora_vs_previa'][0])}, 95% CI "
  f"{sg(BN['mejora_vs_previa'][1])} to {sg(BN['mejora_vs_previa'][2])}) with conservative intervals ({pc(BN['cob80'])} coverage, "
  f"{pc(BN['bajo10'])} below the 10th percentile). Joint updating of θ and *m* was overconfident (standard deviation of standardized "
  f"residuals {f2(BJ['sdz'])}) and did not improve significantly. In a specific test of joint updating on {JJ['sets']} sets whose "
  "lower half of tests spanned a stress ratio of at least three, level-only updating still scored better "
  f"(CRPS {f4(JN['crps'])} versus {f4(JJ['crps'])}). A least-squares fit of φ′1 with *m* = 0.86 achieved CRPS {f4(BI86['crps'])}, not "
  f"significantly worse than Leps-P (difference {sg(RVI[0])}, 95% CI {sg(RVI[1])} to {sg(RVI[2])}), whereas the same fit with *m* = 0.90 "
  f"was biased ({pc(BI90['bajo10'])} of tests below its 10th percentile). Varying the thresholds of a mixed rule (stress ratio 2–4, "
  f"*p* 0.01–0.10) changed the CRPS only between {f4(min(thr))} and {f4(max(thr))} (Fig. 4). The gain from site tests therefore comes "
  "from adjusting the level of the envelope with the correct population curvature; the curvature of a single material cannot be "
  "reliably estimated from a few tests.")
Hd("Form of the Envelope", 2)
P("The quadratic and power forms are indistinguishable within the data but diverge beyond them (Fig. 2): for moderate rock "
  "quality and *e* = 0.45, the median φ′s at 15 MPa is about 1° lower with the power form, which is the conservative choice for "
  "extrapolation. A critical-state floor was not imposed because the literature is divided on whether the critical-state friction "
  "angle of crushable rockfill changes with breakage (Coop et al. 2004; Bandini and Coop 2011; Xiao et al. 2014a, b; Ovalle et al. "
  "2014; Ventini et al. 2021) and the data cannot identify it; it remains a scenario to be tested with data above 8 MPa.")
Hd("Specimen-Size Effect", 2)
P(f"Nineteen materials were tested at more than one maximum particle size (Fig. 5). Across all of them, the random-effects mean of "
  f"β_D was {sg(EG['b'], 3)} (Hartung–Knapp 95% CI {sg(EG['ci_hk'][0], 3)} to {sg(EG['ci_hk'][1], 3)}; Hartung and Knapp 2001; DerSimonian and Laird 1986) "
  f"with a between-material standard deviation of {f3(EG['tau'])}. For {EP['n']} materials with nearly parallel gradation and similar "
  f"density, β_D was {sg(EP['b'], 3)} (95% CI {sg(EP['ci_hk'][0], 3)} to {sg(EP['ci_hk'][1], 3)}), equivalent to κ = {f2(EP['kappa'])} "
  f"({f2(EP['kappa_ci_hk'][0])}–{f2(EP['kappa_ci_hk'][1])}), consistent in sign and compatible in magnitude with Weibull scaling. This "
  "result depends on the similarity criterion: stricter or looser criteria gave intervals that include zero. In a meta-regression, "
  f"an increase of the uniformity coefficient with size tended to raise β_D (coefficient {sg(MR['coef']['dlnCu'], 3)} ± {f3(MR['se'][1])}), "
  "while particle angularity had no clear effect, contrary to the explanation of Varadarajan et al. (2003). Size therefore enters "
  "Leps-P as an uncertainty: by default with the distribution from all groups, and, when field and laboratory gradations are "
  "parallel, with the parallel-grading distribution.")

# ================================================================== APPLICATION
Hd("Application: Waste Rock from a Polymetallic Skarn Deposit")
Hd("Site Data", 2)
P("The waste rock of a polymetallic skarn deposit in the Andes comprises metamorphic carbonate rock (marble), hornfels, intrusive "
  "rock, and sedimentary carbonate rocks (limestone and marl). Drained triaxial tests on large-scale specimens of scalped gradation "
  f"by one laboratory were available for the first three lithologies ({len(SIT['sets'])} sets). Illustrative mass fractions of 35%, 25%, "
  "20%, 10%, and 10% were used for the dump. Rock quality was taken as low for the carbonate and intrusive rocks and moderate for "
  "hornfels; the carbonate class is uncertain because Cylwik et al. (2025) list quarried calcareous rock as good and limestone as low.")
Hd("Envelopes by Lithology and Prior–Data Conflict", 2)
P(f"The conflict check flagged {len(notused)} of the {len(SIT['sets'])} sets (Fig. 6a): both hornfels sets (*p* < 0.001), whose secant "
  "angle increased with stress while the specimens were still contracting at the maximum deviator, and one intrusive set "
  f"(*p* = {min(s['p_conflicto'] for s in notused if s['lit'] == 'Intrusivo'):.3f}). Increasing secant angles over this range are "
  "unusual; densification during consolidation, which is collinear with confining stress in these sets, is a plausible cause that "
  "the data cannot isolate. The flagged sets were not used to update the envelope; for hornfels the prior was kept, which is "
  "conservative here because the tests are stronger than the prior. The remaining sets were updated jointly with a common "
  f"laboratory effect. The slope restriction retained {pc(SIT['retencion_restriccion']['conjunta'])} of the joint samples. "
  f"When the marble was assigned to the good class instead, its set was flagged (*p* = {pmb:.3f}) because the tests are weaker than "
  "that prior, so the conflict check also discriminates between candidate classes. Keeping the good-class prior for the marble "
  f"would have raised the median and 5th percentile of the dump-average φ′s of placed waste rock at 3 MPa by {dmb50:.1f}° and "
  f"{dmb05:.1f}°; the low class, which agrees with the tests, was retained.")
Hd("Placement Density and Dump-Average Envelope", 2)
lab3 = SC["e=lab|ninguno"]["botadero"]["model"]; f3_ = SC["e=0.65|general"]["botadero"]["model"]
P(f"Laboratory specimens had void ratios of {min(s['e'] for s in SIT['sets']):.2f}–{max(s['e'] for s in SIT['sets']):.2f}, whereas "
  "end-dumped waste rock is looser; a placement void ratio of 0.65 was adopted as the base case, with degradation of the marl "
  "represented by an additional increase of 0.15 in void ratio. A failure surface through a mixed dump mobilizes the mass-weighted "
  "strength of the lithologies it crosses; the lithology envelopes were sampled jointly, sharing hyperparameters, fixed effects, "
  "laboratory effect, and size shift. At σ′n = 3 MPa, the median and 5th percentile of the dump-average φ′s are "
  f"{lab3['P50'][isn(3)]:.1f}° and {lab3['P05'][isn(3)]:.1f}° at laboratory density and specimen size, and {f3_['P50'][isn(3)]:.1f}° and "
  f"{f3_['P05'][isn(3)]:.1f}° for placed waste rock at field scale (Table 3; Fig. 6b). Placement density is the largest single effect; "
  "full correlation between lithologies changes the 5th percentile by less than 1°, whereas a surface confined to the weakest "
  "lithology lowers it by several degrees.")
Hd("Stability of Dumps on a Competent Foundation", 2)
d7 = DE("e=0.65", 700, KMOD); d7l = DE("e=lab", 700, KMOD)
P("Dumps 150–700 m high with an overall slope of 2H:1V and a competent, drained level foundation were analyzed with the method of Bishop "
  "(1955) and a nonlinear envelope (Table 4). For 700 m, the critical circle reached σ′n ≈ "
  f"{d7['P50']['snmax']:.1f} MPa, within the calibration range. With the dump-average envelope of placed waste rock, the factor of "
  f"safety was {f2(d7['P50']['F'])} (median envelope) and {f2(d7['P05']['F'])} (5th-percentile envelope), compared with "
  f"{f2(d7l['P50']['F'])} and {f2(d7l['P05']['F'])} at laboratory density. A Monte Carlo analysis with 600 envelope samples per "
  f"correlation assumption gave {MCx('e=0.65', KMOD)['fallas']} failures with the model correlation (median factor of safety "
  f"{f2(MCx('e=0.65', KMOD)['FS_P50'])}) and {MCx('e=0.65', kweak)['fallas']} failures when the surface was confined to the untested, "
  f"degraded marl (probability of failure ≈ {pc(MCx('e=0.65', kweak)['fallas'] / MCx('e=0.65', kweak)['n'])}); at laboratory density "
  f"there were no failures (probability of failure < {MCx('e=lab', KMOD)['pf_sup95'] * 100:.1f}% at 95% confidence). The factor of safety "
  "computed with the 5th-percentile envelope was close to the 5th percentile of the factor of safety "
  f"({f2(d7['P05']['F'])} versus {f2(MCx('e=0.65', KMOD)['FS_P05'])}). The purpose-written Bishop code agreed with the Bishop method in "
  "Slide2 9.043 (Rocscience [[year]]) within 0.7% in an earlier set of cases and was systematically slightly higher.")
Hd("Pore Pressure in a Saturated Basal Layer", 2)
ks = (1e-6, 1e-7, 1e-8, 1e-9); Rs = (25, 100, 300, 600)
SUP = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")
kx = lambda k: "10" + str(int(round(math.log10(k)))).translate(SUP)


def umbral(lim, esc):
    """Mayor k (más permeable) con FS < lim, por velocidad de carga."""
    out = []
    for R_ in Rs:
        kk = [k for k in ks if s62(R_, k, esc) < lim]
        out.append(f"*k* ≤ {kx(max(kk))} m/s at {R_} m/year" if kk else f"none at {R_} m/year")
    return ", ".join(out[:-1]) + ", and " + out[-1]


fsP05 = [s62(R_, k) for R_ in Rs for k in ks[:2]]
P("The 700 m dump was placed on a 10° foundation slope over a 30 m saturated basal layer that does not daylight at the toe. "
  "Excess pore pressure in each column was computed with one-dimensional consolidation under a load increasing at rate *R* "
  "(Olson 1977), for constrained modulus 50 MPa, drained top, and impervious base. It was imposed in Slide2 as a pore-pressure "
  "grid assigned to the basal layer only, with a nonlinear search optimized by simulated annealing and the Morgenstern and Price "
  f"(1965) method (Table 5, Fig. 7). With *k* ≥ 10⁻⁷ m/s the excess pore pressure dissipated during loading: the factor of safety was "
  f"{f2(min(s62(R_, k, 'P50') for R_ in Rs for k in ks[:2]))}–{f2(max(s62(R_, k, 'P50') for R_ in Rs for k in ks[:2]))} with the median "
  f"envelope and {f2(min(fsP05))}–{f2(max(fsP05))} with the 5th-percentile envelope, the latter close to the level-foundation case. "
  f"With *k* = 10⁻⁹ m/s it fell to {f2(s62(25, 1e-9))} at 25 m/year and {f2(s62(600, 1e-9))} at 600 m/year, with the critical surface "
  "sliding along the basal layer under most of the dump. The factor of safety with the median envelope fell below 1.3 for "
  f"{umbral(1.3, 'P50')}, and that with the 5th-percentile envelope fell below 1.0 for {umbral(1.0, 'P05')}. "
  "An undrained basal layer with Su = 0.17σ′v (the isotropic ratio of 0.28 reported by "
  "Valenzuela et al. (2008) referred to σ′v under K0 ≈ 0.4) and with the liquefied ratio Su = 0.09σ′v of Olson and Stark (2002), "
  "extrapolated far beyond its stress range, gave factors of safety of "
  f"{f2(s62u(0.17, 100, 1e-8, 'libre'))} and {f2(s62u(0.09, 100, 1e-8, 'libre'))} (σ′v consistent with *R* = 100 m/year and *k* = 10⁻⁸ m/s). "
  "These undrained cases are sensitivity analyses: the strength of crushed, contractive waste rock under undrained loading at "
  "several megapascals has not been measured.")

# ================================================================== DISCUSSION
Hd("Discussion")
Hd("What Leps-P Adds", 2)
P("Leps-P does not predict the central strength of an untested material better than Leps 2.0, and with a few site tests it "
  "performs similarly to a careful least-squares fit with the curvature fixed at 0.86. Its contributions are the population "
  "curvature itself, which corrects the bias of fits with *m* = 0.9; a calibrated, stress-dependent uncertainty without site tests; "
  "a validated updating rule with a conflict check; the separation of publication and material variability; and a framework in "
  "which placement density, specimen size, correlation between lithologies, and pore pressure enter a single reliability analysis.")
Hd("Design Format", 2)
P("A design procedure consistent with these results is: (1) assign rock quality and placement void ratio to each lithology and "
  "check the conflict diagnostic for any site tests; (2) update with tests that pass, at the level only; (3) shift the envelopes "
  "to field particle size and placement density; (4) adopt as characteristic value for a global mechanism the 5th percentile of "
  "the mean envelope, a cautious estimate of the mean in the sense of Eurocode 7 (CEN 2004; Prästings et al. 2019); (5) check the "
  "probability of failure by sampling the envelopes with stated correlations between lithologies, including the weakest-lithology "
  "and pore-pressure scenarios; and (6) calibrate the target factor of safety for the 5th-percentile envelope to the consequence "
  "class rather than applying factors calibrated for best-estimate parameters (Read and Stacey 2009), because the strength "
  "uncertainty would otherwise be counted twice.")
Hd("Limitations", 2)
P("The calibration database contains no tests above σ′n = 7.9 MPa, its high-stress tests come from eight publications, and its "
  "lower tail is truncated at φ′1 = 31°. Class assignment for lithologies not listed by Cylwik et al. (2025) is a judgment, and "
  "class labels may carry information from measured strength. The size effect was estimated from few materials and depends on "
  "gradation. The within-material residual is laboratory scatter and does not represent spatial variability in the dump. The "
  "pore-pressure analysis is one-dimensional and the undrained strength ratios were taken from lower-stress studies. These gaps "
  "can be closed by re-estimating Leps-P as data appear: drained tests at σ′3 of 5–10 MPa on carbonate and argillaceous waste "
  "rock, undrained tests at high pressure on crushed gradations, tests of one material at two sizes with parallel gradation, and "
  "in situ density and permeability measurements in high dumps.")

# ================================================================== CONCLUSIONS
Hd("Conclusions")
for t in [
    f"Without site tests, the Leps-P 80% interval covered {pc(A_all['m']['cob80'])} of tests from held-out publications, with accuracy "
    "equal to that of Leps 2.0 with constant dispersion and conservative intervals above 2.75 MPa.",
    f"With a few site tests at low stress, updating only the level of the envelope reduced the CRPS at higher stress from "
    f"{f3(BN['crps'] + BN['mejora_vs_previa'][0])} to {f3(BN['crps'])}; updating the curvature was overconfident, and a least-squares fit "
    "with *m* fixed at 0.86 performed similarly.",
    f"Between-publication differences accounted for {pc(fl_th)} of the between-set variance of the reference angle and most of that of "
    "the curvature.",
    "The specimen-size effect was heterogeneous; for nearly parallel gradations it was negative and compatible with Weibull scaling.",
    "For a 700 m dump of skarn waste rock, placement density, untested lithologies, and pore pressure in a saturated basal layer "
    "under rapid loading controlled the factor of safety more than the extrapolation of the envelope.",
]:
    p = doc.add_paragraph(style="List Number"); runs(p, t)

# ================================================================== BACK MATTER
Hd("Data Availability Statement")
P("Some or all data, models, or code generated or used during the study are available in a repository or online in accordance "
  "with funder data retention policies (Leps-P code, parameters, and evaluation scripts: [[repository DOI]]; calibration "
  "database: Arrieta 2025). Some or all data, models, or code generated or used during the study are proprietary or "
  "confidential in nature and may only be provided with restrictions (the site test results, which may be provided in "
  "anonymized form with the permission of their owner).", indent=False)
Hd("Acknowledgments")
P("[[Acknowledgments, funding, and the permission of the owner of the site test data to be completed.]] [[Declaration of the "
  "use of AI tools, as required by ASCE policy: tool, purpose, and confirmation that the authors reviewed and take "
  "responsibility for the content.]]", indent=False)
Hd("Notation")
P("*The following symbols are used in this paper:*", indent=False)
NOT = [("C_β", "covariance matrix of the fixed effects"), ("*D*f, *D*l", "maximum particle size in the field and in the laboratory"),
       ("$D^2$", "prior predictive Mahalanobis distance of the site tests"), ("G_p, G_s", "covariance matrices of the publication and set random effects"),
       ("*K*0", "coefficient of earth pressure at rest"), ("*k*", "permeability of the basal layer"), ("*R*", "loading rate of the dump over the basal layer"),
       ("Su", "undrained shear strength"), ("a0, a1", "parameters of the stress-dependent residual dispersion"),
       ("b^(p), b^(s)", "publication and set random effects on (θ, *m*)"), ("c2", "quadratic coefficient of the log–log envelope"),
       ("*e*", "initial void ratio"), ("*m*", "curvature exponent of the power envelope"), ("*p*", "*p*-value of the conflict check"),
       ("p_a", "atmospheric pressure (0.1013 MPa)"), ("*x*", "ln(σ′n/σ0)"), ("*x̄*", "mean of *x* over the site tests of a set"),
       ("*y*", "ln(τ/σ0)"), ("β_c", "intercept of rock quality class *c*"), ("β_D", "specimen-size coefficient"),
       ("β_e", "effect of void ratio on θ"), ("β_m", "mean curvature"), ("θ", "ln tan φ′1"), ("κ", "ratio of size exponent to Weibull modulus"),
       ("σ0", "reference stress, 10p_a"), ("σ′n", "effective normal stress on the failure plane"), ("σ′v, σv", "effective and total vertical stress"),
       ("σ′3", "effective confining stress"), ("σε", "standard deviation of test-to-test scatter"), ("τ", "shear strength"),
       ("φ′1", "secant friction angle at σ′n = σ0"), ("φ′e=0", "φ′1 extrapolated to *e* = 0 in Leps 2.0"), ("φ′s", "secant friction angle"),
       ("ω", "decrease of φ′1 per unit void ratio")]
for s_, d_ in NOT:
    P(f"{s_} = {d_};", indent=False)

br(); Hd("References")
REFS = [
    "Arrieta, M. 2025. \"Rockfill test database. Version 1.\" Zenodo. Accessed October 6, 2026. https://doi.org/10.5281/zenodo.14665504.",
    "Bandini, V., and M. R. Coop. 2011. \"The influence of particle breakage on the location of the critical state line of sands.\" *Soils Found.* 51 (4): 591–600. https://doi.org/10.3208/sandf.51.591.",
    "Barton, N., and B. Kjærnsli. 1981. \"Shear strength of rockfill.\" *J. Geotech. Eng. Div.* 107 (7): 873–891. https://doi.org/10.1061/AJGEB6.0001167.",
    "Bishop, A. W. 1955. \"The use of the slip circle in the stability analysis of slopes.\" *Géotechnique* 5 (1): 7–17. https://doi.org/10.1680/geot.1955.5.1.7.",
    "CEN (European Committee for Standardization). 2004. *Eurocode 7: Geotechnical design—Part 1: General rules.* EN 1997-1. Brussels, Belgium: CEN.",
    "Charles, J. A., and K. S. Watts. 1980. \"The influence of confining pressure on the shear strength of compacted rockfill.\" *Géotechnique* 30 (4): 353–367. https://doi.org/10.1680/geot.1980.30.4.353.",
    "Ching, J., and K.-K. Phoon. 2019. \"Constructing site-specific multivariate probability distribution model using Bayesian machine learning.\" *J. Eng. Mech.* 145 (1): 04018126. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001537.",
    "Ching, J., S. Wu, and K.-K. Phoon. 2021. \"Constructing quasi-site-specific multivariate probability distribution using hierarchical Bayesian model.\" *J. Eng. Mech.* 147 (10): 04021069. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001964.",
    "Christian, J. T., C. C. Ladd, and G. B. Baecher. 1994. \"Reliability applied to slope stability analysis.\" *J. Geotech. Eng.* 120 (12): 2180–2207. https://doi.org/10.1061/(ASCE)0733-9410(1994)120:12(2180).",
    "Contreras, L.-F., and S. Linero-Molina. 2025. \"Hierarchical Bayesian inference of rockfill strength parameters from large-scale triaxial test data.\" *Aust. Geomech.* 60 (2): 91–106. https://doi.org/10.56295/AGJ6024.",
    "Contreras, L.-F., S. Linero-Molina, and J. Dixon. 2022. \"Bayesian approach to improve the confidence of the estimation of the shear strength of coarse mine waste using Barton's empirical criterion.\" *Aust. Geomech.* 57 (2): 163–174. https://doi.org/10.56295/AGJ5728.",
    "Coop, M. R., K. K. Sørensen, T. Bodas Freitas, and G. Georgoutsos. 2004. \"Particle breakage during shearing of a carbonate sand.\" *Géotechnique* 54 (3): 157–163. https://doi.org/10.1680/geot.2004.54.3.157.",
    "Cylwik, S. D., M. Arrieta, J. C. Cunning, and P. M. Hawley. 2025. \"Update to the Leps strength envelopes for rockfill based on 900 large-scale triaxial tests.\" *Int. J. Geomech.* 25 (12): 04025285. https://doi.org/10.1061/IJGNAI.GMENG-12050.",
    "De Mello, V. F. B. 1977. \"Reflections on design decisions of practical significance to embankment dams.\" *Géotechnique* 27 (3): 281–355. https://doi.org/10.1680/geot.1977.27.3.281.",
    "DerSimonian, R., and N. Laird. 1986. \"Meta-analysis in clinical trials.\" *Controlled Clin. Trials* 7 (3): 177–188. https://doi.org/10.1016/0197-2456(86)90046-2.",
    "Frossard, E., W. Hu, C. Dano, and P.-Y. Hicher. 2012. \"Rockfill shear strength evaluation: A rational method based on size effects.\" *Géotechnique* 62 (5): 415–427. https://doi.org/10.1680/geot.10.P.079.",
    "Gelman, A., J. B. Carlin, H. S. Stern, D. B. Dunson, A. Vehtari, and D. B. Rubin. 2013. *Bayesian data analysis.* 3rd ed. Boca Raton, FL: CRC Press. https://doi.org/10.1201/b16018.",
    "Gneiting, T., and A. E. Raftery. 2007. \"Strictly proper scoring rules, prediction, and estimation.\" *J. Am. Stat. Assoc.* 102 (477): 359–378. https://doi.org/10.1198/016214506000001437.",
    "Hartung, J., and G. Knapp. 2001. \"A refined method for the meta-analysis of controlled clinical trials with binary outcome.\" *Stat. Med.* 20 (24): 3875–3889. https://doi.org/10.1002/sim.1009.",
    "Hawley, M., and J. Cunning, eds. 2017. *Guidelines for mine waste dump and stockpile design.* Clayton, Australia: CSIRO Publishing.",
    "Indraratna, B., L. S. S. Wijewardena, and A. S. Balasubramaniam. 1993. \"Large-scale triaxial testing of greywacke rockfill.\" *Géotechnique* 43 (1): 37–51. https://doi.org/10.1680/geot.1993.43.1.37.",
    "Laird, N. M., and J. H. Ware. 1982. \"Random-effects models for longitudinal data.\" *Biometrics* 38 (4): 963–974. https://doi.org/10.2307/2529876.",
    "Leps, T. M. 1970. \"Review of shearing strength of rockfill.\" *J. Soil Mech. Found. Div.* 96 (4): 1159–1170. https://doi.org/10.1061/JSFEAQ.0001433.",
    "Hungr, O., R. F. Dawson, A. Kent, D. Campbell, and N. R. Morgenstern. 2002. \"Rapid flow slides of coal-mine waste in British Columbia, Canada.\" In *Catastrophic landslides: Effects, occurrence, and mechanisms*, edited by S. G. Evans and J. V. DeGraff, 191–208. Boulder, CO: Geological Society of America. https://doi.org/10.1130/REG15-p191.",
    "Linero, S., C. Palma, and R. Apablaza. 2007. \"Geotechnical characterisation of waste material in very high dumps with large scale triaxial testing.\" In *Proc., Int. Symp. on Rock Slope Stability in Open Pit Mining and Civil Engineering*, edited by Y. Potvin, 59–75. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/708_2.",
    "Marachi, N. D., C. K. Chan, and H. B. Seed. 1972. \"Evaluation of properties of rockfill materials.\" *J. Soil Mech. Found. Div.* 98 (1): 95–114. https://doi.org/10.1061/JSFEAQ.0001735.",
    "Marsal, R. J. 1973. \"Mechanical properties of rockfill.\" In *Embankment-dam engineering: Casagrande volume*, edited by R. C. Hirschfeld and S. J. Poulos, 109–200. New York: Wiley.",
    "Morgenstern, N. R., and V. E. Price. 1965. \"The analysis of the stability of general slip surfaces.\" *Géotechnique* 15 (1): 79–93. https://doi.org/10.1680/geot.1965.15.1.79.",
    "Olson, R. E. 1977. \"Consolidation under time dependent loading.\" *J. Geotech. Eng. Div.* 103 (1): 55–60. https://doi.org/10.1061/AJGEB6.0000369.",
    "Olson, S. M., and T. D. Stark. 2002. \"Liquefied strength ratio from liquefaction flow failure case histories.\" *Can. Geotech. J.* 39 (3): 629–647. https://doi.org/10.1139/t02-001.",
    "Ovalle, C., E. Frossard, C. Dano, W. Hu, S. Maiolino, and P.-Y. Hicher. 2014. \"The effect of size on the strength of coarse rock aggregates and large rockfill samples through experimental data.\" *Acta Mech.* 225 (8): 2199–2216. https://doi.org/10.1007/s00707-014-1127-z.",
    "Ovalle, C., S. Linero, C. Dano, E. Bard, P.-Y. Hicher, and R. Osses. 2020. \"Data compilation from large drained compression triaxial tests on coarse crushable rockfill materials.\" *J. Geotech. Geoenviron. Eng.* 146 (9): 06020013. https://doi.org/10.1061/(ASCE)GT.1943-5606.0002314.",
    "Phoon, K.-K., and F. H. Kulhawy. 1999. \"Characterization of geotechnical variability.\" *Can. Geotech. J.* 36 (4): 612–624. https://doi.org/10.1139/t99-038.",
    "Prästings, A., J. Spross, and S. Larsson. 2019. \"Characteristic values of geotechnical parameters in Eurocode 7.\" *Proc. Inst. Civ. Eng. Geotech. Eng.* 172 (4): 301–311. https://doi.org/10.1680/jgeen.18.00057.",
    "Read, J., and P. Stacey, eds. 2009. *Guidelines for open pit slope design.* Collingwood, Australia: CSIRO Publishing.",
    "Rocscience. [[year]]. *Slide2.* Version 9.043. Toronto: Rocscience.",
    "Valenzuela, L., E. Bard, J. Campaña, and M. E. Anabalón. 2008. \"High waste rock dumps: Challenges and developments.\" In *Proc., 1st Int. Seminar on the Management of Rock Dumps, Stockpiles and Heap Leach Pads*, 65–78. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/802_6.",
    "Varadarajan, A., K. G. Sharma, K. Venkatachalam, and A. K. Gupta. 2003. \"Testing and modeling two rockfill materials.\" *J. Geotech. Geoenviron. Eng.* 129 (3): 206–218. https://doi.org/10.1061/(ASCE)1090-0241(2003)129:3(206).",
    "Ventini, R., S. Lirer, C. Mancuso, and A. Flora. 2021. \"Simplified procedure to identify the critical state line of crushable rockfills.\" *J. Geotech. Geoenviron. Eng.* 147 (12): 04021150. https://doi.org/10.1061/(ASCE)GT.1943-5606.0002668.",
    "Xiao, Y., H. Liu, Y. Chen, and J. Jiang. 2014a. \"Strength and deformation of rockfill material based on large-scale triaxial compression tests. I: Influences of density and pressure.\" *J. Geotech. Geoenviron. Eng.* 140 (12): 04014070. https://doi.org/10.1061/(ASCE)GT.1943-5606.0001176.",
    "Xiao, Y., H. Liu, Y. Chen, and J. Jiang. 2014b. \"Strength and deformation of rockfill material based on large-scale triaxial compression tests. II: Influence of particle breakage.\" *J. Geotech. Geoenviron. Eng.* 140 (12): 04014071. https://doi.org/10.1061/(ASCE)GT.1943-5606.0001177.",
]
for ref in REFS:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5); runs(p, ref)

br(); Hd("Figure Captions List")
for c in [
    "Fig. 1. Distribution of effective normal stress on the failure plane in the calibration database, compared with stresses in a 700 m waste dump.",
    "Fig. 2. Predictive distribution of the secant friction angle of a new material of moderate rock quality with e = 0.45 tested by a new laboratory, for the quadratic and power forms. The shaded zone exceeds the largest normal stress in the database.",
    "Fig. 3. Out-of-sample calibration without site tests, leaving out each publication: (a) coverage of the 80% interval; and (b) fraction of tests below the 10th percentile, by range of normal stress. Dashed lines are nominal values.",
    "Fig. 4. Mean CRPS of ln τ for held-out tests at σ′n ≥ 2 MPa (64 tests, 13 publications) for the prior and for procedures that use site tests at σ′n ≤ 1.5 MPa; bars are 95% bootstrap confidence intervals.",
    "Fig. 5. Within-material specimen-size coefficient for 19 materials tested at more than one maximum particle size (bars: 95% confidence intervals; one value outside the axis is annotated). Hatched bands are Hartung–Knapp 95% confidence intervals of the random-effects means.",
    "Fig. 6. Waste rock of the skarn deposit: (a) laboratory-scale envelopes by lithology with the site tests (open symbols: tests flagged by the conflict check); and (b) envelopes of placed waste rock (void ratio 0.65, maximum particle size 1 m) by lithology and dump-average envelope. Shaded zones are beyond the stresses tested at the site.",
    "Fig. 7. Factor of safety of a 700 m dump on a 10° foundation slope with a 30 m saturated basal layer, as a function of the permeability of the basal layer and the loading rate (Morgenstern–Price method): (a) median envelope; and (b) 5th-percentile envelope.",
]:
    P(c, indent=False)


def tabla(n, titulo, enc, filas, nota=None):
    if n == 1:
        br()
    else:
        doc.add_paragraph()
    p = doc.add_paragraph(); rr = p.add_run(f"Table {n}. "); rr.bold = True; p.add_run(titulo); p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(filas), cols=len(enc)); M_ = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for j, h in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; runs(c.paragraphs[0], h)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f):
            c = t.rows[i].cells[j]; c.text = ""; pp = c.paragraphs[0]; runs(pp, str(v))
            if not any(r_.text.strip() for r_ in pp.runs):
                for om in pp._p.findall(f"{{{M_}}}oMath"):
                    para = etree.SubElement(pp._p, f"{{{M_}}}oMathPara"); ppr = etree.SubElement(para, f"{{{M_}}}oMathParaPr")
                    jc = etree.SubElement(ppr, f"{{{M_}}}jc"); jc.set(f"{{{M_}}}val", "left"); pp._p.remove(om); para.append(om)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing = 1.0
                for r_ in pp.runs:
                    r_.font.size = Pt(10)
    # ASCE: líneas horizontales arriba, bajo los encabezados y al final
    tblPr = t._tbl.tblPr; bd = OxmlElement("w:tblBorders")
    for side in ("top", "bottom"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); bd.append(e)
    tblPr.append(bd)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); cb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); cb.append(e); tcPr.append(cb)
    if nota:
        p = doc.add_paragraph(); runs(p, "Note: " + nota)


tabla(1, "Parameters of Leps-P (quadratic form)", ["Parameter", "Estimate", "Standard error or 90% credible interval"],
      [["β_c, good rock quality", f3(b["good"]), f3(SE["good"])], ["β_c, moderate rock quality", f3(b["moderate"]), f3(SE["moderate"])],
       ["β_c, low rock quality", f3(b["low"]), f3(SE["low"])], ["β_e", neg(f3(b["e"])), f3(SE["e"])], ["β_m", f3(b["x"]), f3(SE["x"])],
       ["c2", f4(b["x2"]), f4(SE["x2"])],
       ["SD of θ, publication", f3(HP[0][0]), f"{f3(HP[0][1])}–{f3(HP[0][2])}"], ["SD of m, publication", f3(HP[1][0]), f"{f3(HP[1][1])}–{f3(HP[1][2])}"],
       ["Correlation of θ and m, publication", neg(f2(HP[2][0])), neg(f"{f2(HP[2][1])} to {f2(HP[2][2])}")],
       ["SD of θ, set", f3(HP[3][0]), f"{f3(HP[3][1])}–{f3(HP[3][2])}"], ["SD of m, set", f3(HP[4][0]), f"{f3(HP[4][1])}–{f3(HP[4][2])}"],
       ["Correlation of θ and m, set", neg(f2(HP[5][0])), neg(f"{f2(HP[5][1])} to {f2(HP[5][2])}")],
       ["σε at σ′n = σ0", f3(HP[6][0]), f"{f3(HP[6][1])}–{f3(HP[6][2])}"], ["a1", neg(f3(HP[7][0])), neg(f"{f3(HP[7][1])} to {f3(HP[7][2])}")]],
      f"{N['ensayos']} tests, {N['sets']} sets, {N['pub']} publications. Fixed effects: restricted maximum likelihood; variance parameters: posterior median and 90% credible interval. "
      f"φ′e=0 = {phi_e0['good']:.1f}°, {phi_e0['moderate']:.1f}°, and {phi_e0['low']:.1f}° for good, moderate, and low rock quality.")
tabla(2, "Out-of-sample evaluation (leave one publication out)", ["Procedure", "Tests", "80% coverage", "Below P10", "CRPS"],
      [["Mode A, Leps-P, all stresses", A_all["m"]["n"], pc(A_all["m"]["cob80"]), pc(A_all["m"]["bajo10"]), f4(A_all["m"]["crps"])],
       ["Mode A, Leps 2.0 with constant dispersion", A_all["m0"]["n"], pc(A_all["m0"]["cob80"]), pc(A_all["m0"]["bajo10"]), f4(A_all["m0"]["crps"])],
       ["Mode A, Leps-P, σ′n ≥ 2.75 MPa", A_hi["m"]["n"], pc(A_hi["m"]["cob80"]), pc(A_hi["m"]["bajo10"]), f4(A_hi["m"]["crps"])],
       ["Mode A, Leps-P, class unknown", AMIX["n"], pc(AMIX["cob80"]), pc(AMIX["bajo10"]), f4(AMIX["crps"])],
       ["Mode B, prior (no site tests)", BN["n"], pc(PRB["cob80"]), pc(PRB["bajo10"]), f4(PRB["crps"])],
       ["Mode B, fit of φ′1 with m = 0.90", BI90["n"], pc(BI90["cob80"]), pc(BI90["bajo10"]), f4(BI90["crps"])],
       ["Mode B, fit of φ′1 with m = 0.86", BI86["n"], pc(BI86["cob80"]), pc(BI86["bajo10"]), f4(BI86["crps"])],
       ["Mode B, Leps-P joint update", BJ["n"], pc(BJ["cob80"]), pc(BJ["bajo10"]), f4(BJ["crps"])],
       ["Mode B, Leps-P level-only update", BN["n"], pc(BN["cob80"]), pc(BN["bajo10"]), f4(BN["crps"])]],
      "Nominal values: 80% coverage and 10% below P10. Mode B predicts tests at σ′n ≥ 2 MPa from tests at σ′n ≤ 1.5 MPa of the same set. CRPS in units of ln τ.")
filas = []
lit = SC["e=0.65|general"]["litologias"]; litl = SC["e=lab|ninguno"]["litologias"]
for k in lit:
    a_ = litl.get(k); f_ = lit[k]
    filas.append([k, f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}" if a_ else "—",
                  f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}", f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}",
                  f"{f_['P50'][isn(10)]:.1f} / {f_['P05'][isn(10)]:.1f}"])
for nm, k in (("Dump average, model correlation", "model"), ("Dump average, full correlation", "comonotonic")):
    a_ = SC["e=lab|ninguno"]["botadero"][k]; f_ = SC["e=0.65|general"]["botadero"][k]
    filas.append([nm, f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}",
                  f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(10)]:.1f} / {f_['P05'][isn(10)]:.1f}"])
tabla(3, "Secant friction angle φ′s (°) of the skarn waste rock, median / 5th percentile of the mean envelope",
      ["Material", "Laboratory, 3 MPa", "Placed, 1 MPa", "Placed, 3 MPa", "Placed, 10 MPa"], filas,
      "Laboratory: specimen void ratio and size. Placed: void ratio 0.65 (marl 0.80, degraded), maximum particle size 1 m, size effect from all groups. No site tests above σ′n ≈ 4.6 MPa.")
filas = []
for k, nm in ((KMOD, "Dump average, model correlation"), (KCOM, "Dump average, full correlation"), (kweakt, "Weakest tested lithology only"), (kweak, "Untested degraded marl only")):
    filas.append([nm] + [f"{f2(DE('e=0.65', h, k)['P50']['F'])} / {f2(DE('e=0.65', h, k)['P05']['F'])}" for h in (150, 400, 700)]
                 + [f"{MCx('e=0.65', k)['fallas']}/{MCx('e=0.65', k)['n']}"])
tabla(4, "Factor of safety of dumps on a competent level foundation, placed waste rock (median / 5th-percentile envelope)",
      ["Envelope", "150 m", "400 m", "700 m", "Failures, 700 m"], filas,
      "Bishop method. Monte Carlo with 600 envelope samples for 700 m. At laboratory density the 700 m dump gave "
      f"{f2(d7l['P50']['F'])} / {f2(d7l['P05']['F'])} and no failures.")
filas = [[f"{R_} m/year"] + [f2(s62(R_, k)) for k in ks] for R_ in Rs]
tabla(5, "Factor of safety of a 700 m dump on a 10° foundation slope with a 30 m saturated basal layer (5th-percentile envelope)",
      ["Loading rate", "k = 10⁻⁶ m/s", "k = 10⁻⁷ m/s", "k = 10⁻⁸ m/s", "k = 10⁻⁹ m/s"], filas,
      "Morgenstern–Price method; pore pressure from one-dimensional consolidation under linearly increasing load (constrained modulus 50 MPa).")
out = AQ / "LepsP_IJGM_manuscript_r2.docx"; doc.save(str(out)); print("written", out, "| equations:", NEQ[0], "| Slide2 62:", len(S62))
