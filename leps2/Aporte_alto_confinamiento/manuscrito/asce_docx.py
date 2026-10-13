"""Formato ASCE (Int. J. Geomech.) para los generadores del manuscrito HiBER: doble espacio, numeración continua de
líneas y de páginas, ecuaciones OMML, títulos sin numerar y tablas solo con líneas horizontales."""
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from lxml import etree
import latex2mathml.converter as l2m
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))

# ------------------------------------------------------------------ document setup
doc = Document()
cp = doc.core_properties; cp.author = ""; cp.title = "HiBER"; cp.comments = ""; cp.last_modified_by = ""; cp.subject = ""; cp.keywords = ""
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.line_spacing = 2.0; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
for m_ in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m_, Inches(1))
ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); sec._sectPr.append(ln)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    r_ = fp.add_run()
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r_._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r_._r.append(e)


def omml(latex):
    return XSL(etree.fromstring(l2m.convert(latex).encode())).getroot()


SYM = [("b^(p)", r"\mathbf{b}^{(p)}"), ("b^(s)", r"\mathbf{b}^{(s)}"), ("C_β", r"\mathbf{C}_\beta"), ("β̂", r"\hat{\boldsymbol{\beta}}"),
       ("β_c", r"\beta_c"), ("β_e", r"\beta_e"), ("β_D", r"\beta_D"), ("G_p", r"\mathbf{G}_p"), ("G_s", r"\mathbf{G}_s"),
       ("σε²", r"\sigma_\varepsilon^2"), ("σε", r"\sigma_\varepsilon"), ("φ′cs", r"\phi'_{cs}"), ("φ′e=0", r"\phi'_{e=0}"),
       ("φ′1", r"\phi'_1"), ("φ′s", r"\phi'_s"), ("σ′n", r"\sigma'_n"), ("σ′3", r"\sigma'_3"), ("σ′v", r"\sigma'_v"),
       ("σ0", r"\sigma_0"), ("p_a", r"p_a"), ("*D*max", r"D_{max}"), ("*D*f", r"D_f"), ("*D*l", r"D_l"), ("q_0", r"q_0"), ("g_0", r"g_0"),
       ("*x*i", r"x_i"), ("*r*i", r"r_i"), ("*x̄*", r"\bar{x}"), ("x̃", r"\tilde{\mathbf{x}}"), ("**z**", r"\mathbf{z}"), ("a0, a1", r"a_0, a_1"),
       ("D², ", r"D^2, ")]


def simbolos(text):
    parts = re.split(r"(\$[^$]+\$|https?://\S+)", text); out = []
    for s in parts:
        if s.startswith("$") or s.startswith("http"):
            out.append(s); continue
        for a_, l_ in SYM:
            s = s.replace(a_, f"${l_}$")
        out.append(s)
    return "".join(out)


def runs(p, text, _sym=True):
    if _sym:
        text = simbolos(text)
    for tok in re.split(r"(\$[^$]+\$|\*[^*]+\*|\[\[[^\]]+\]\])", text):
        if not tok:
            continue
        if tok.startswith("$"):
            p._p.append(parse_xml(etree.tostring(omml(tok[1:-1]))))
        elif tok.startswith("*") and "$" in tok:
            runs(p, tok[1:-1], False)
        elif tok.startswith("*"):
            p.add_run(tok[1:-1]).italic = True
        elif tok.startswith("[["):
            rr = p.add_run(tok[2:-2]); rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p.add_run(tok)
    return p


def P(text, indent=True):
    TEX.append(("p", text, indent))
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    return runs(p, text)


def Hd(text, level=1):
    TEX.append(("h", text, level))
    p = doc.add_paragraph(); rr = p.add_run(text); rr.bold = level == 1; rr.italic = level == 2; p.paragraph_format.keep_with_next = True


NEQ = [0]


def EQ(latex):
    TEX.append(("eq", latex))
    NEQ[0] += 1; p = doc.add_paragraph(); pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(3.25), alignment=1); pf.tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    p.add_run("\t"); p._p.append(parse_xml(etree.tostring(omml(latex)))); p.add_run(f"\t({NEQ[0]})")


def br():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ salida LaTeX paralela (ascelike-new, oficial ASCE)
TEX = []                                    # (tipo, ...) en el orden del documento; lo escribe escribir_tex()
META = {}


def titulo(text):
    META["titulo"] = text
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text).bold = True


def autores(nombre, afil):
    META["autores"] = nombre; META["afil"] = afil
    for t in (nombre, afil):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, t)


def item(text):
    TEX.append(("item", text)); p = doc.add_paragraph(style="List Number"); runs(p, text)


def ref(text):
    TEX.append(("ref", text))
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5); runs(p, text, False)


def figura(n, archivo, leyenda):
    """Leyenda en la lista de Word; figura con \\includegraphics en LaTeX (endfloat la lleva al final)."""
    P(f"Fig. {n}. {leyenda}", indent=False); TEX.pop(); TEX.append(("fig", n, archivo, leyenda))


_CH = {"α": r"\alpha", "β": r"\beta", "γ": r"\gamma", "δ": r"\delta", "Δ": r"\Delta", "ε": r"\varepsilon", "ζ": r"\zeta", "κ": r"\kappa",
       "η": r"\eta", "θ": r"\theta", "λ": r"\lambda", "μ": r"\mu", "π": r"\pi", "ρ": r"\rho", "σ": r"\sigma", "τ": r"\tau",
       "φ": r"\phi", "χ": r"\chi", "ψ": r"\psi", "ω": r"\omega", "′": "'", "≤": r"\le", "≥": r"\ge", "≈": r"\approx",
       "×": r"\times", "·": r"\cdot", "±": r"\pm", "∞": r"\infty", "−": "-", "∼": r"\sim", "→": r"\rightarrow", "≠": r"\neq",
       "∝": r"\propto", "√": r"\surd", "ᵃ": "^{a}", "ᵇ": "^{b}"}
_TXT = {"°": r"\textdegree{}", "–": "--", "—": "---", "’": "'", "‘": "`", "“": "``", "”": "''", "…": r"\ldots{}",
        "\u00a0": "~", "\u2009": r"\,", "\u202f": r"\,", "³": r"\textsuperscript{3}"}
_SUP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻", "0123456789-"); _SUB = str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")


def _esc(s):
    out = []
    for tok in re.split(r"(https?://\S+?)(?=\.?(?:\s|$))", s):
        if tok.startswith("http"):
            out.append(r"\url{" + tok + "}"); continue
        tok = re.sub(r"([\\{}%&#_^~$])", lambda m: {"\\": r"\textbackslash{}", "^": r"\^{}", "~": r"\~{}"}.get(m.group(1), "\\" + m.group(1)), tok)
        tok = re.sub(r"(\w)\u0302", r"$\\hat{\1}$", tok); tok = re.sub(r"(\w)\u0304", r"$\\bar{\1}$", tok); tok = re.sub(r"(\w)\u0303", r"$\\tilde{\1}$", tok)
        tok = re.sub(r"[⁰¹²⁴⁵⁶⁷⁸⁹⁻]+(?:³)?|²", lambda m: "$^{" + m.group(0).translate(_SUP) + "}$", tok)
        tok = re.sub(r"[₀-₉]+", lambda m: "$_{" + m.group(0).translate(_SUB) + "}$", tok)
        tok = "".join(f"${_CH[c]}$" if c in _CH else _TXT.get(c, c) for c in tok)
        out.append(tok)
    return "".join(out)


def tex(text, sym=True):
    if sym:
        text = simbolos(text)
    out = []
    for tok in re.split(r"(\$[^$]+\$|\*[^*]+\*|\[\[[^\]]+\]\])", text):
        if not tok:
            continue
        if tok.startswith("$"):
            out.append(tok)
        elif tok.startswith("*"):
            inner = tex(tok[1:-1], False)
            out.append(inner if inner.startswith("$") and inner.endswith("$") else r"\textit{" + inner + "}")
        elif tok.startswith("[["):
            out.append(r"\textbf{[" + _esc(tok[2:-2]) + "]}")
        else:
            out.append(_esc(tok))
    s = "".join(out)
    while "$$" in s:
        s = s.replace("$$", "")
    return re.sub(r'"([^"]*)"', r"``\1''", s)                    # comillas tipográficas


def _tabla_tex(n, titulo_, enc, filas, nota):
    cols = "l" + "c" * (len(enc) - 1)
    ancho = len(enc) >= 6 or max(len(str(v)) for f in filas for v in f) > 16   # tablas anchas: letra menor
    L = [r"\begin{table}", r"\caption{" + tex(titulo_) + "}", r"\label{tab:" + str(n) + "}", r"\centering" + (r"\footnotesize" if ancho else r"\small"),
         r"\renewcommand{\arraystretch}{1.15}", (r"\setlength{\tabcolsep}{3pt}" if ancho else ""), r"\begin{tabular}{" + cols + "}", r"\hline",
         " & ".join(tex(h) for h in enc) + r" \\", r"\hline"]
    for f in filas:
        L.append(" & ".join(tex(str(v)) for v in f) + r" \\")
    L += [r"\hline", r"\end{tabular}"]
    if nota:
        L.append(r"\par\vspace{4pt}\parbox{\linewidth}{\footnotesize Note: " + tex(nota) + "}")
    L.append(r"\end{table}")
    return "\n".join(L)


# ---- citas autor-año → \cite (claves de latex/lepsp.bib)
CITAS = {("Alonso et al.", "2012"): "alonso2012", ("Arrieta", "2025"): "arrieta2025", ("Barton and Kjærnsli", "1981"): "barton1981",
         ("Been and Jefferies", "1985"): "been1985", ("Bolton", "1986"): "bolton1986", ("CEN", "2004"): "cen2004", ("CEN", "2024"): "cen2024",
         ("Charles and Watts", "1980"): "charles1980", ("Ching and Phoon", "2019"): "ching2019", ("Ching et al.", "2021"): "ching2021",
         ("Christian et al.", "1994"): "christian1994", ("Clopper and Pearson", "1934"): "clopper1934",
         ("Contreras and Linero-Molina", "2025"): "contreras2025", ("Contreras et al.", "2022"): "contreras2022", ("Coop et al.", "2004"): "coop2004",
         ("Cylwik et al.", "2025"): "cylwik2025", ("De Mello", "1977"): "demello1977", ("Einav", "2007"): "einav2007",
         ("Frossard et al.", "2012"): "frossard2012", ("Gneiting and Raftery", "2007"): "gneiting2007", ("Hardin", "1985"): "hardin1985",
         ("Hawley and Cunning", "2017"): "hawley2017", ("Holm", "1979"): "holm1979", ("Jefferies and Been", "2016"): "jefferies2016",
         ("Knapp and Hartung", "2003"): "knapp2003", ("Lade et al.", "1996"): "lade1996", ("Laird and Ware", "1982"): "laird1982",
         ("Leps", "1970"): "leps1970", ("Lindstrom and Bates", "1990"): "lindstrom1990", ("Linero et al.", "2007"): "linero2007",
         ("Marachi et al.", "1972"): "marachi1972", ("Marsal", "1973"): "marsal1973", ("McDowell and Bolton", "1998"): "mcdowell1998",
         ("Morgenstern and Price", "1965"): "morgenstern1965", ("Oldecop and Alonso", "2001"): "oldecop2001", ("Ovalle et al.", "2014"): "ovalle2014",
         ("Ovalle et al.", "2020"): "ovalle2020", ("Prästings et al.", "2019"): "prastings2019", ("Read and Stacey", "2009"): "read2009",
         ("Rocscience", "[year]"): "rocscience", ("Russell and Khalili", "2004"): "russell2004", ("Valenzuela et al.", "2008"): "valenzuela2008",
         ("Varadarajan et al.", "2003"): "varadarajan2003", ("Vehtari et al.", "2021"): "vehtari2021", ("Viechtbauer", "2007"): "viechtbauer2007",
         # ronda 6
         ("Bozorgzadeh et al.", "2019"): "bozorgzadeh2019", ("Ching and Phoon", "2020"): "ching2020", ("Ovalle and Dano", "2020"): "ovalle2020b",
         ("Girumugisha et al.", "2026"): "girumugisha2026", ("Ventini et al.", "2021"): "ventini2021", ("Xiao et al.", "2014a"): "xiao2014a",
         ("Xiao et al.", "2014b"): "xiao2014b", ("Indraratna et al.", "1993"): "indraratna1993", ("Bard et al.", "2011"): "bard2011",
         ("Higgins et al.", "2009"): "higgins2009", ("Plummer", "2015"): "plummer2015", ("Vanmarcke", "1977"): "vanmarcke1977",
         ("Pinheiro and Bates", "2000"): "pinheiro2000", ("CEN", "2002"): "cen2002"}
_AUT = "|".join(sorted({re.escape(a) for a, _ in CITAS}, key=len, reverse=True))
_ANO = r"(?:(?:1[89]\d\d|20\d\d)[ab]?|\\textbf\{\[year\]\})"
_UNA = re.compile(rf"(?:(e\.g\.,)\s*)?({_AUT}) ({_ANO}(?:, {_ANO})*)(?:, p\. (\d+))?")


def _claves(aut, anos):
    return [CITAS[(aut, "[year]" if "year" in y else y)] for y in re.findall(_ANO, anos)]


def citas(s):
    """Convierte las citas del texto: «Autor (año)» → \\citeN; «(A año; B año)» → \\cite{a,b}; citas dentro de un paréntesis
    con otro texto → \\citeNP. Devuelve el texto y la lista de claves usadas."""
    usadas = []

    def grupo(m):
        cont = m.group(1); partes = cont.split("; ")
        llenas = [_UNA.fullmatch(p_) for p_ in partes]
        if all(llenas) and not any(x.group(1) or x.group(4) for x in llenas):
            ks = [k for x in llenas for k in _claves(x.group(2), x.group(3))]; usadas.extend(ks)
            return r"\cite{" + ",".join(ks) + "}"

        def una(x):
            ks = _claves(x.group(2), x.group(3)); usadas.extend(ks)
            nota = f"[p.~{x.group(4)}]" if x.group(4) else ""
            return (x.group(1) + " " if x.group(1) else "") + r"\citeNP" + nota + "{" + ",".join(ks) + "}"
        return "(" + _UNA.sub(una, cont) + ")"
    s = re.sub(r"\(([^()]*)\)", grupo, s)

    def narr(m):
        ks = _claves(m.group(1), m.group(2)); usadas.extend(ks); return r"\citeN{" + ",".join(ks) + "}"
    s = re.sub(rf"({_AUT}) \(({_ANO})\)", narr, s)
    return s, usadas


def notacion(items):
    """Lista de notación en el orden de ASCE (latinas, mayúscula antes que minúscula; luego griegas; luego otros
    símbolos), con «;» al final de cada definición, «; and» en la penúltima y punto en la última."""
    GR = "αβγδεζηθικλμνξοπρστυφχψω"; GRL = {"Delta": "δ", "beta": "β", "tau": "τ", "phi": "φ", "sigma": "σ"}

    def clave(sym):
        t = re.sub(r"\\mathbf\{(\w)\}", r"\1", sym).replace("$", "").replace("*", "").lstrip("\\")
        for nm, g in GRL.items():
            if t.startswith(nm):
                t = g + t[len(nm):]
        c = t[0]
        if c.isascii() and c.isalpha():
            return (0, c.lower(), 0 if c.isupper() else 1, t)
        if c.lower() in GR or c == "Δ":
            return (1, GR.index("δ" if c == "Δ" else c.lower()), t)
        return (2, t)
    items = sorted(items, key=lambda it: clave(it[0]))
    for i, (s_, d_) in enumerate(items):
        fin = "." if i == len(items) - 1 else ("; and" if i == len(items) - 2 else ";")
        P(f"{s_} = {d_}{fin}", indent=False); TEX.pop()
    TEX.append(("notation", items))


def escribir_tex(ruta, figdir, bib="lepsp", nametag="[Author]"):
    """Escribe el manuscrito con la clase y el preámbulo de la plantilla oficial ascelike-new (Journal: doble espacio,
    líneas numeradas, tablas y figuras al final con sus listas; referencias con BibTeX y ascelike-new.bst)."""
    H = [r"\documentclass[Journal,letterpaper]{ascelike-new}", r"\WarningFilter{caption}{Unknown document class}",
         r"\usepackage[utf8]{inputenc}", r"\usepackage[T1]{fontenc}", r"\usepackage{lmodern}", r"\usepackage{graphicx}",
         r"\usepackage[style=base,figurename=Fig.,labelfont=bf,labelsep=period]{caption}", r"\usepackage{subcaption}",
         r"\usepackage{amsmath}", r"\usepackage{newtxtext,newtxmath}",
         r"\usepackage[colorlinks=true,citecolor=red,linkcolor=black]{hyperref}",
         "% adicionales a la plantilla", r"\usepackage{bm,textcomp,url,longtable}", r"\urlstyle{same}",
         r"\renewcommand{\boldsymbol}[1]{\bm{#1}}", r"\graphicspath{{" + figdir + "}}",
         r"\NameTag{" + nametag + r", \today}", r"\begin{document}",
         r"\title{" + tex(META.get("titulo", ""), False) + "}", r"\author[1]{" + tex(META.get("autores", "")) + "}",
         r"\affil[1]{" + tex(META.get("afil", "")) + "}", r"\maketitle"]
    B = []; modo = None; tablas = []
    for e in TEX:
        k = e[0]
        if modo == "abstract" and k != "p":
            B.append(r"\end{abstract}"); modo = None
        if modo == "enum" and k != "item":
            B.append(r"\end{enumerate}"); modo = None
        if k == "h":
            t = e[1]
            if t == "Abstract":
                B.append(r"\begin{abstract}"); modo = "abstract"; continue
            if t == "Figure Captions List":
                modo = "figs"; continue
            if t == "References":
                B.append(r"\bibliography{" + bib + "}"); modo = "refs"; continue
            B.append((r"\section{" if e[2] == 1 else r"\subsection{") + tex(t, False) + "}")
        elif k == "p":
            if modo == "figs":
                continue
            if e[1].startswith("*Author keywords:*"):
                if modo == "abstract":
                    B.append(r"\end{abstract}"); modo = None
                B.append(r"\KeyWords{" + tex(e[1].replace("*Author keywords:* ", "")) + "}\n"); continue
            B.append(("" if e[2] else r"\noindent ") + tex(e[1]) + "\n")
        elif k == "eq":
            B.append(r"\begin{equation}" + "\n" + e[1] + "\n" + r"\end{equation}")
        elif k == "item":
            if modo != "enum":
                B.append(r"\begin{enumerate}"); modo = "enum"
            B.append(r"\item " + tex(e[1]))
        elif k == "notation":
            n_ = len(e[1])
            B.append(r"\begin{longtable}{@{}r@{\hspace{1em}=\hspace{1em}}p{5.0in}@{}}")
            for i, (s_, d_) in enumerate(e[1]):
                fin = "." if i == n_ - 1 else ("; and" if i == n_ - 2 else ";")
                B.append(tex(s_) + " & " + tex(d_) + fin + r" \\")
            B.append(r"\end{longtable}"); B.append(r"\addtocounter{table}{-1}")   # longtable sin título no debe consumir un número de tabla
        elif k == "fig":
            B.append(r"\begin{figure}" + "\n" + r"\centering\includegraphics[width=\linewidth]{" + e[2] + "}\n" + r"\caption{"
                     + tex(e[3]) + "}\n" + r"\label{fig:" + str(e[1]) + "}\n" + r"\end{figure}")
        elif k == "table":
            tablas.append(_tabla_tex(*e[1:]))
    s, usadas = citas("\n".join(B))
    open(ruta, "w", encoding="utf-8").write("\n".join(H) + "\n\n" + s + "\n\n" + "\n\n".join(tablas) + "\n\n" + r"\end{document}" + "\n")
    return sorted(set(usadas))



def tabla(n, titulo, enc, filas, nota=None):
    TEX.append(("table", n, titulo, enc, filas, nota))
    if n == 1:
        br()
    else:
        doc.add_paragraph()
    p = doc.add_paragraph(); rr = p.add_run(f"Table {n}. "); rr.bold = True; runs(p, titulo); p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(filas), cols=len(enc)); M_ = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for j, h in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; runs(c.paragraphs[0], h)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f):
            c = t.rows[i].cells[j]; c.text = ""; pp = c.paragraphs[0]; runs(pp, str(v))
            if not any(r_.text.strip() for r_ in pp.runs):
                for om in pp._p.findall(f"{{{M_}}}oMath"):
                    para = etree.SubElement(pp._p, f"{{{M_}}}oMathPara"); ppr = etree.SubElement(para, f"{{{M_}}}oMathParaPr")
                    jc = etree.SubElement(ppr, f"{{{M_}}}jc"); jc.set(f"{{{M_}}}val", "left"); pp._p.remove(om); para.append(om)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing = 1.0
                for r_ in pp.runs:
                    r_.font.size = Pt(10)
    tblPr = t._tbl.tblPr; bd = OxmlElement("w:tblBorders")
    for side in ("top", "bottom"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); bd.append(e)
    tblPr.append(bd)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); cb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); cb.append(e); tcPr.append(cb)
    if nota:
        p = doc.add_paragraph(); runs(p, "Note: " + nota)


