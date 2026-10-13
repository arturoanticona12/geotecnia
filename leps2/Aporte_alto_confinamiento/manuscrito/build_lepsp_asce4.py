"""Leps-P manuscript, round 4 (ASCE Int. J. Geomech., Technical Paper). Bounded envelope with material-level random
effects on the dilatant level and the decay rate; stability with Morgenstern–Price in Slide2. All numbers are read
from result files; formatting in asce_docx.py."""
import json, math
from pathlib import Path
import numpy as np, pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
import asce_docx as A
from asce_docx import doc, P, Hd, EQ, br, runs, tabla

AQ = Path(__file__).parent; R = AQ.parent; AN = R / "analisis"; AA = R / "analisis_antamina"
J = lambda p: json.load(open(p, encoding="utf-8"))
PAR = J(R / "leps3" / "lepsp_parametros.json"); FA = PAR["formas"]["acotada"]; N = PAR["n"]; ESCP = PAR["escala"]
M4 = J(AN / "lepsp_mcmc4.json"); M3 = J(AN / "lepsp_mcmc3.json"); CV = J(AN / "lepsp_cv4_metricas.json")
ESC = J(AN / "lepsp_escala3.json"); ESC1 = J(AN / "lepsp_escala3_f1.json"); VAR = J(AN / "lepsp_varianza.json"); DPL = J(AN / "lepsp_desplegada.json")
TD = J(AN / "lepsp_tabla_diseno.json"); DG = J(AN / "lepsp_diag_alta.json"); G2 = J(AN / "lepsp_acotada2_grilla.json")
SIT = J(AA / "lepsp_sitio3.json"); SITQ = J(AA / "lepsp_sitio3_x2.json"); SITL = J(AA / "lepsp_sitio3_labindep.json"); CF = J(AA / "a7_criterio_falla.json")
EST = J(AN / "lepsp_estabilidad_slide2.json")
A.SYM[:0] = [("φ′∞", r"\phi'_\infty"), ("ζ0", r"\zeta_0"), ("u_η", r"u_\eta"), ("u_ζ", r"u_\zeta")]

# ------------------------------------------------------------------ numbers
CL = ["good", "moderate", "low"]; b = dict(zip(FA["nombres"], FA["beta"])); SE = dict(zip(FA["nombres"], np.sqrt(np.diag(FA["cov_beta"]))))
HP = FA["hiper"]; deg = lambda th: math.degrees(math.atan(math.exp(th))); phi1 = {c: deg(b[c]) for c in CL}
pc = lambda v: f"{100 * v:.0f}%"; f2 = lambda v: f"{v:.2f}"; f3 = lambda v: f"{v:.3f}"; f4 = lambda v: f"{v:.4f}"
neg = lambda s: s.replace("-", "−"); sg = lambda v, dd=4: neg(f"{v:+.{dd}f}")
ML_ac = M4["cresta"]["ml_opt"]; ML_q = M3["ml"]["x2"]
mlg = np.array(G2["ml"]); dML0 = float(np.max(mlg) - np.max(mlg[G2["phi"].index(0.0)]))
ridge = [c for c in FA["cresta"] if c["ml"] >= max(x["ml"] for x in FA["cresta"]) - 3]
phi_lo, phi_hi = min(c["phi_cs"] for c in ridge), max(c["phi_cs"] for c in ridge)
PHM, ZEM, PHT = M4["cresta"]["phi_media"], M4["cresta"]["zeta_media"], M4["cresta"]["phi_media_tan"]
sh = [math.exp(2 * s["psi"][0]) / (math.exp(2 * s["psi"][0]) + math.exp(2 * s["psi"][3])) for s in FA["muestras"]]
SHARE = (np.median(sh), np.percentile(sh, 5), np.percentile(sh, 95))
A_, Ah, Al = CV["A|todo"], CV["A|sn>2.75"], CV["A|sn<0.5"]
BN, BS, BD, BJ, BP, B90, BTR = (CV[f"B|{k}"] for k in ("nivel", "sitio", "descarte", "conjunta", "previa", "ingenuo_m090", "ingenuo_mtr"))
dg = lambda o: f"{sg(o['media'])} (95% CI {sg(o['ic95'][0])} to {sg(o['ic95'][1])}; sign-flip *p* = {o['p_signflip']:.2f}, Holm-adjusted {o['p_holm']:.2f})"
EP, EP1 = ESC["publicaciones R ≥ 2"], ESC1["publicaciones R ≥ 2"]
V3, V10 = VAR["low e=0.65 σ′n=3"], VAR["low e=0.65 σ′n=10"]; vs = lambda V, k: 100 * V["comp"][k] / V["total"]
sdφ = lambda V: math.degrees(math.sqrt(V["total"]) * math.sin(math.radians(V["phi_mediana"])) * math.cos(math.radians(V["phi_mediana"])))
SNs = np.array(SIT["sn"]); isn = lambda s: int(np.argmin(np.abs(SNs - s)))
notused = [s for s in SIT["sets"] if not s["usado"]]
D_ = lambda caso, esc, H, q: EST["det"][f"{caso}|{esc}|{H}|{q}"]
MCm, MCe = EST["mc"]["modelo"], EST["mc"]["E"]
cpi = lambda m: f"{100 * m['pf_ic95'][0]:.1f}–{100 * m['pf_ic95'][1]:.1f}%"
GM = J(AN / "lepsp_estabilidad4.json") if (AN / "lepsp_estabilidad4.json").exists() else None
fb = SIT["escenarios"]["e=0.65|general"]["botadero"]; fbn = SIT["escenarios"]["e=0.65|ninguno"]["botadero"]; lbn = SIT["escenarios"]["e=lab|ninguno"]["botadero"]
weak = lambda dct, t: [v for k, v in dct.items() if k.startswith(t)][0]

# ================================================================== TITLE, ABSTRACT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Leps-P: Probabilistic Strength Envelopes with Stress-Decaying Curvature for Rockfill and Mine Waste Rock").bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, "[[Author names, ASCE grade, affiliations, and corresponding author e-mail to be completed.]]")
Hd("Abstract")
P("Reliability-based design of high waste dumps needs the drained shear strength of rockfill and mine waste rock as a "
  "distribution, with its dependence on stress and a rule to combine published experience with site tests. Leps-P is a "
  "three-level hierarchical model of the secant friction angle in which the dilatant part of strength decays with stress "
  "toward an asymptote; each material has its own dilatant level and decay rate. It was calibrated on "
  f"{N['ensayos']} large-scale triaxial tests from {N['pub']} publications of the public Leps 2.0 database. The likelihood "
  f"favored this form over a quadratic log–log envelope with the same number of parameters (log-likelihood difference "
  f"{ML_ac - ML_q:.1f}) and over a form without an asymptote ({dML0:.1f}). The asymptote, about {PHM:.0f}°, is a population "
  "decay parameter well below the angles measured at the highest stresses, not a critical-state angle. Leaving out one "
  "publication at a time, the bounded, quadratic, and Leps 2.0 envelopes predicted the strength of new materials with "
  "statistically indistinguishable accuracy, but only Leps-P had calibrated intervals at low stress. With a few site tests "
  f"at low stress, updating the level of each material gave calibrated predictions at higher stress ({pc(BN['cob80'])} "
  "coverage of the 80% interval); differences from simpler fits were not significant with 13 publications. At field scale, "
  f"about {vs(V3, 'escala_campo'):.0f}% of the variance of the strength of an untested waste rock at 3 MPa came from the "
  "assumed heterogeneity of the specimen-size effect. For a 700 m dump on a competent foundation, analyzed with the "
  f"Morgenstern–Price method, the factor of safety was {f2(D_('e065', 'model', 700, 'P50'))} with the median and "
  f"{f2(D_('e065', 'model', 700, 'P05'))} with the 5th-percentile envelope of placed waste rock. Design charts and open "
  "code are provided.", indent=False)
Hd("Practical Applications")
P("Designers of high waste dumps and rockfill dams often have few or no large-scale shear tests and rely on published "
  "charts of rockfill strength. This work turns the Leps chart into a tool that gives a best estimate of strength and a "
  "cautious value for design, for any rock quality and density, together with the range of uncertainty. Its curves "
  "flatten at high pressure, as expected when crushed rock stops dilating, and they are supported by tests up to about "
  "8 MPa on the failure plane. A few tests of the site material at low pressure sharpen the estimate at higher pressure, "
  "and the tool warns when site tests disagree with published experience. For untested waste rock, most of the remaining "
  "uncertainty comes, under the adopted model, from the size of the rock particles and from differences between "
  "materials and laboratories; testing the actual waste rock at large scale and at two particle sizes is therefore the "
  "most valuable investigation. Free software accompanies the paper.", indent=False)
P("*Author keywords:* Rockfill; Mine waste rock; Shear strength; Nonlinear failure envelope; Dilatancy; Hierarchical model; "
  "Reliability-based design; High waste dumps.", indent=False)

# ================================================================== INTRODUCTION
Hd("Introduction")
P("Leps (1970) compiled 109 large-scale triaxial tests and proposed three nonlinear envelopes for rockfill that have been "
  "used for dams and waste dumps whenever site-specific tests were unavailable. Cylwik et al. (2025) updated them with a "
  "public database of 938 tests (Arrieta 2025), adopted the power criterion of De Mello (1977) and Charles and Watts (1980) "
  "with a single fitting parameter, the secant friction angle at σ′n = 10p_a, and related that angle to rock quality and "
  "void ratio. That update is referred to here as Leps 2.0.")
P("Two developments motivate a further extension. First, design is moving toward reliability-based criteria (Read and Stacey "
  "2009; Hawley and Cunning 2017), which require a strength distribution and a rule to combine published experience with "
  "site tests. Second, waste dumps now reach 500–700 m (Valenzuela et al. 2008). For a uniform dump on a competent "
  "foundation the critical surface stays at normal stresses of a few megapascals, but mechanisms controlled by the "
  "foundation, by a basal layer, or by the weakest material reach higher stresses, where waste rock crushes and loses "
  "dilatancy (Marsal 1973; Hardin 1985; Lade et al. 1996; Linero et al. 2007; Ovalle et al. 2020). How an envelope "
  "behaves at the edge of and beyond the data then matters: a power law with a constant exponent keeps curving downward, "
  "whereas the excess of strength over the critical state, which is the contribution of dilatancy, is suppressed by "
  "confinement (Bolton 1986; Jefferies and Been 2016).")
P("Hierarchical models that borrow strength from generic databases have been developed for soil properties (Ching and Phoon "
  "2019; Ching et al. 2021), and a hierarchical Bayesian analysis of the Barton and Kjærnsli (1981) parameters has been applied "
  "to the rockfill groups of a single project (Contreras et al. 2022; Contreras and Linero-Molina 2025). Leps-P differs in "
  "pooling 36 publications, in its envelope form, and in its out-of-sample evaluation. Its contributions are (1) an envelope "
  "whose curvature decays with stress for every material, compared with fixed-form alternatives by likelihood and out of "
  "sample; (2) the separation of publication and material variability; (3) a site-updating rule with a prior–data conflict "
  "check, compared with the simple fits used in practice; (4) a decomposition of the uncertainty of an untested waste rock at "
  "field scale; and (5) design charts and an open implementation.")

# ================================================================== MODEL
Hd("Database and Model Formulation")
Hd("Calibration Database", 2)
P(f"The public database (Arrieta 2025) was filtered to tests with σ′3 ≥ 60 kPa, reported initial void ratio, and at least two tests "
  f"per set, giving {N['ensayos']} tests on {N['sets']} sets from {N['pub']} publications (Fig. S1). A set is a material with a given "
  "gradation, density, and rock type tested in one program; failure is the maximum deviator stress, as in Leps 2.0. Rock quality "
  "classes follow Table 5 of Cylwik et al. (2025). Tests with σ′n > 3 MPa (90 tests) come from eight publications, seven of them "
  "published between 1963 and 1983 (e.g., Marachi et al. 1972; Marsal 1973), and sets with φ′1 < 31° were excluded by Cylwik "
  f"et al. (2025), so the lower tail of the population is truncated. At the maximum deviator the specimens were contracting: "
  f"the median net volumetric strain was {DG['bandas'][0]['evol_mediana']:.1f}% below 0.5 MPa and {DG['bandas'][-1]['evol_mediana']:.1f}% "
  f"above 3 MPa, where {pc(DG['bandas'][-1]['frac_contractiva'])} of the specimens contracted and the median breakage index of "
  f"Marsal (1973) was {DG['bandas'][-1]['bg_mediana']:.0f}% (Fig. S2). Above 5 MPa the measured secant angles ranged from "
  f"{DG['phis_sobre_5'][0]:.1f}° to {DG['phis_sobre_5'][2]:.1f}° (5th to 95th percentile, {DG['n_sobre_5']} tests).")
Hd("Envelope", 2)
P("The Leps 2.0 criterion is the power law")
EQ(r"\tau=\tan\phi'_1\,\sigma_0^{1-m}\,\sigma_n'^{\,m}")
P("where σ0 = 10p_a, p_a = 0.1013 MPa, φ′1 is the secant friction angle at σ′n = σ0, and *m* is a constant exponent; the secant "
  "friction angle is φ′s = arctan(τ/σ′n). In granular materials the excess of strength over the critical state is governed by "
  "dilatancy, which is controlled by the state of the material relative to its critical state and is suppressed by "
  "confinement (Bolton 1986; Jefferies and Been 2016, p. 43); particle crushing reinforces this trend (Hardin 1985; McDowell "
  "and Bolton 1998; Coop et al. 2004; Einav 2007). A minimal form with this property is", indent=False)
EQ(r"\tan\phi'_s=T+\left(\tan\phi'_1-T\right)\exp(-\zeta x),\qquad x=\ln(\sigma'_n/\sigma_0),\quad T=\tan\phi'_\infty")
P("in which the dilatant part decays at rate ζ toward the asymptotic angle φ′∞. For φ′∞ = 0 it reduces to a power law with "
  "*m* = 1 − ζ; Bolton (1986) proposed instead a dilatant part that vanishes at a finite stress. Because the specimens at high "
  "stress were contracting at the maximum deviator, part of the curvature reflects strain-limited, under-mobilized tests, so "
  "φ′∞ is a population decay parameter of the maximum-deviator envelope and not a critical-state angle.", indent=False)
Hd("Hierarchical Model", 2)
P("With *y* = ln(τ/σ0), the strength of test *i* of set *j* of publication *k* is")
EQ(r"y_{ijk}=x_{ijk}+\ln\!\left[T+\left(\exp\eta_{jk}-T\right)\exp\!\left(-\zeta_{jk}x_{ijk}\right)\right]+\varepsilon_{ijk}")
EQ(r"\eta_{jk}=\beta_{c(jk)}+\beta_e e_{jk}+u_{\eta,jk},\quad \zeta_{jk}=\zeta_0+u_{\zeta,jk},\quad \mathbf{u}_{jk}=\mathbf{u}^{(p)}_k+\mathbf{u}^{(s)}_{jk}")
EQ(r"\mathbf{u}^{(p)}\sim N(\mathbf{0},\mathbf{G}_p),\quad \mathbf{u}^{(s)}\sim N(\mathbf{0},\mathbf{G}_s),\quad \ln\sigma_\varepsilon(x)=a_0+a_1x")
P("where η = ln tan φ′1 is the dilatant level of the material, β_c the coefficient of rock quality class *c*, β_e the effect of "
  "initial void ratio *e*, ζ0 the population decay rate, u = (u_η, u_ζ) the publication and set random effects with covariance "
  "matrices G_p and G_s, and ε the test-to-test deviation with standard deviation σε (Laird and Ware 1982). Every material "
  "thus has its own envelope that decays toward the common asymptote, and its local slope is below one when ζ > 0. The "
  "publication effect combines laboratory, procedure, era, and material selection; because true strength and laboratory bias "
  "cannot be distinguished, it is treated as uncertainty of strength. Through β_c and β_e the median φ′1 reproduces the "
  "relation φ′1 = φ′e=0 − ωe of Leps 2.0. A quadratic log–log form, *y* = θ + *m*·*x* + *c*₂*x*², and a power form "
  "(*c*₂ = 0), with random effects on θ and *m*, were fitted for comparison.", indent=False)
P("The asymptote is weakly identified: lower values of φ′∞ with slower decay fit nearly as well as higher values with faster "
  "decay. The profile likelihood of (φ′∞, ζ0) was evaluated on a grid (14–32° and 0.175–0.475), the points holding 99.5% of the "
  "normalized likelihood were refitted, and predictions average over them with weights proportional to the likelihood (flat "
  "prior on the grid); a prior uniform in tan φ′∞ was used as a sensitivity.")
Hd("Predictive Distribution and Design Values", 2)
P("For a new material of class *c* and void ratio *e* tested by a new laboratory, *y* is approximately normal with mean "
  "*f*(*x*) given by Eq. (3) with u = 0 and variance")
EQ(r"\mathrm{Var}[y]=\mathbf{z}^T(\mathbf{G}_p+\mathbf{G}_s)\mathbf{z}+\mathbf{j}^T\mathbf{C}_\beta\mathbf{j}+\sigma_\varepsilon^2(x)")
P(r"where $\mathbf{z}$ and $\mathbf{j}$ are the gradients of *f* with respect to (η, ζ) and to the fixed effects and C_β is the "
  "covariance of the fixed effects (first-order linearization; Lindstrom and Bates 1990). The first two terms describe the "
  "uncertainty of the mean envelope of the material, which does not average out along a failure surface within that "
  "material (Christian et al. 1994); the last describes test-to-test scatter. The implementation samples the variance "
  "parameters from their posterior, (φ′∞, ζ0) from the ridge weights, the fixed effects, and the random effects. Fig. 1 and "
  "Table 3 give the median and 5th percentile of the mean envelope of untested materials at laboratory specimen size.", indent=False)
Hd("Updating with Site Tests", 2)
P("Site triaxial tests are reduced with the maximum-deviator criterion of the database and converted to (σ′n, τ) with Eqs. "
  "(13), (14), and (16) of Cylwik et al. (2025). Consistency with the prior is checked with the prior predictive Mahalanobis distance")
EQ(r"D^2=\mathbf{r}^T\mathbf{V}^{-1}\mathbf{r}\sim\chi^2_n,\qquad \mathbf{V}=\mathbf{Z}(\mathbf{G}_p+\mathbf{G}_s)\mathbf{Z}^T+\mathbf{J}\mathbf{C}_\beta\mathbf{J}^T+\mathbf{S}")
P(r"where $\mathbf{r}$ holds the residuals *r*i of the *n* site tests from the prior mean, $\mathbf{Z}$ and $\mathbf{J}$ have rows "
  r"$\mathbf{z}_i^T$ and $\mathbf{j}_i^T$, and $\mathbf{S}=\mathrm{diag}\,\sigma_\varepsilon^2(x_i)$; a *p*-value below 0.05 flags the "
  "tests for review. The update treats the residuals as measurements of the level of the envelope at the centroid *x̄* of "
  r"the tests. With $\mathbf{v}$ the random effects of the material and the deviation of the fixed effects, $\mathbf{P}$ their "
  r"prior covariance, $\mathbf{a}$ the gradient at *x̄*, and g_0 = $\mathbf{a}^T\mathbf{P}\mathbf{a}$, the level q_0 = "
  r"$\mathbf{a}^T\mathbf{v}$ has the posterior", indent=False)
EQ(r"\hat q_0=\frac{\sum_i r_i/\sigma_{\varepsilon,i}^2}{1/g_0+\sum_i 1/\sigma_{\varepsilon,i}^2},\qquad \mathrm{Var}(q_0)=\left(\frac{1}{g_0}+\sum_i\frac{1}{\sigma_{\varepsilon,i}^2}\right)^{-1}")
P(r"and $\mathbf{v}$ is shifted by $\mathbf{P}\mathbf{a}\,\hat q_0/g_0$. Eq. (8) is a cut (modular) update that can be applied "
  "in a spreadsheet; it does not estimate the decay of a single material from a few tests. When several materials of one "
  "site were tested by the same laboratory, flagged sets are excluded and the remaining sets are updated jointly with a "
  "shared publication effect, which flagged and untested materials then inherit.", indent=False)
Hd("Placement Density and Specimen Size", 2)
P("A change of void ratio Δ*e* from the specimens to the placed material enters through η, so its effect fades with stress. "
  "Because β_e was estimated between sets, in which void ratio also reflects gradation and particle shape, applying it to a "
  "change of state of one material is an assumption. Specimen size enters as a shift of level")
EQ(r"\Delta y=\beta_D\ln\left(D_f/D_l\right)")
P("where *D*f and *D*l are the maximum particle sizes in the field and in the laboratory; Weibull scaling predicts β_D < 0 "
  "(Frossard et al. 2012; Ovalle et al. 2014; Alonso et al. 2012).", indent=False)

# ================================================================== ESTIMATION & EVALUATION
Hd("Estimation and Out-of-Sample Evaluation")
P("Fixed effects were estimated by Gauss–Newton iteration on the linearized mean, with restricted maximum likelihood (REML) "
  "for the variance parameters. Variance parameters were sampled with an adaptive Metropolis algorithm under weakly "
  "informative priors (half-normal with scale 0.25 on standard deviations, uniform on correlations, normal on a0, a1), with "
  "four chains from dispersed starting values; the fixed effects are thus estimated conditionally (an empirical Bayes scheme). "
  "Convergence was checked with the rank-normalized potential scale reduction factor and bulk and tail effective sample "
  "sizes (Vehtari et al. 2021).")
P("Each publication was left out in turn; in each fold (φ′∞, ζ0) were selected on the training data only and the model was "
  "refitted. In Mode A (no site tests), every test of the held-out publication was predicted. In Mode B, each held-out set "
  "with at least two tests at σ′n ≤ 1.5 MPa and one at σ′n ≥ 2 MPa was updated with the former, and the latter were predicted "
  "with: the prior; the level update of each set (Eq. 8); the same after discarding flagged sets; the recommended "
  "procedure (flagged sets discarded and all sets of the publication updated jointly with a shared laboratory effect); a "
  "joint update of both random effects; and least-squares fits of φ′1 with *m* = 0.90 or with *m* equal to the median "
  "within-set slope of the training data, whose predictive standard deviation was their error on the training "
  "publications. These procedures were developed while examining earlier versions of the model on the same held-out "
  "sets, so the Mode B comparisons are exploratory. Scores were the coverage of the 80% interval, the fraction of tests below "
  "the 10th percentile, the continuous ranked probability score (CRPS) of ln τ (Gneiting and Raftery 2007), and, for the 10th "
  "percentile, the quantile (pinball) loss. Differences in publication means were tested with exact sign-flip tests, with "
  "Holm (1979) adjustment within each mode, and summarized with bootstrap intervals over publications. The Mode A "
  "reference was the Leps 2.0 envelope (*m* = 0.9) with a constant dispersion fitted to the training data; its constants were "
  "fitted to the whole database, which favors it.")
P(f"The folds used plug-in predictive distributions. With all data, the distribution of the implementation differed from "
  f"the plug-in one by at most {DPL['dmed_abs_max']:.3f} in median ln τ and had a standard deviation {DPL['ratio_sd_min']:.2f}–"
  f"{DPL['ratio_sd_max']:.2f} times as large; the implementation uses the same update rule as Mode B.")

# ================================================================== RESULTS
Hd("Results")
Hd("Parameters and Evidence for Stress-Decaying Curvature", 2)
P(f"Table 1 lists the parameters. The rank-normalized potential scale reduction factor was at most {max(M4['acotada']['rhat']):.3f}, "
  f"and the bulk and tail effective sample sizes were at least {min(M4['acotada']['ess_bulk']):.0f} and {min(M4['acotada']['ess_tail']):.0f}. "
  f"Half-normal prior scales of 0.1 and 0.5 changed the median between-publication standard deviation of η from "
  f"{M4['previa_0.25']['hiper'][0][0]:.3f} to {M4['previa_0.1']['hiper'][0][0]:.3f} and {M4['previa_0.5']['hiper'][0][0]:.3f}. At σ0, "
  f"publication effects account for {pc(SHARE[0])} (90% credible interval {pc(SHARE[1])}–{pc(SHARE[2])}) of the between-material "
  "variance of η.")
P(f"The bounded form exceeded the quadratic form, which has the same number of parameters, by {ML_ac - ML_q:.1f} in maximum "
  f"log-likelihood, and removing the asymptote (φ′∞ = 0, a power law in tan φ′s) reduced it by {dML0:.1f}. Within 3 "
  f"log-likelihood units of the maximum, φ′∞ ranged from {phi_lo:.0f}° to {phi_hi:.0f}° (Fig. 2); the weighted mean was "
  f"{PHM:.1f}° ({PHT:.1f}° with a prior uniform in tan φ′∞) and ζ0 = {ZEM:.2f}. All 36 folds selected φ′∞ of 24–26° and ζ0 = 0.30. "
  "These are in-sample comparisons; out of sample the forms were indistinguishable (next section).")
Hd("Mode A: No Site Tests", 2)
P(f"Across all held-out tests (Table 2, Fig. 3a), the bounded and quadratic forms had CRPS {f4(A_['acotada']['crps'])} and "
  f"{f4(A_['cuadratica']['crps'])} ({dg(A_['acotada-cuadratica|crps'])}), and Leps-P covered {pc(A_['acotada']['cob80'])} of tests with "
  f"its 80% interval. Above 2.75 MPa ({Ah['n']} tests, {Ah['n_pub']} publications) both forms underpredicted the median (mean "
  f"standardized residual {sg(Ah['acotada']['z_medio'], 2)} and {sg(Ah['cuadratica']['z_medio'], 2)}), and the Leps 2.0 reference had a "
  f"lower 10% quantile loss ({dg(Ah['acotada-leps2_ref|pin10'])}). Below 0.5 MPa the constant-dispersion interval covered only "
  f"{pc(Al['leps2_ref']['cob80'])} of tests, with the median biased low (mean standardized residual {sg(Al['leps2_ref']['z_medio'], 2)}), "
  f"whereas Leps-P covered {pc(Al['acotada']['cob80'])}; this reflects the constant dispersion of the reference, not the Leps 2.0 "
  "envelope itself.")
Hd("Mode B: Site Tests at Low Stress", 2)
P(f"With the level update of each set, {pc(BN['cob80'])} of the held-out tests fell within the 80% interval and the standard "
  f"deviation of standardized residuals was {f2(BN['sd_z'])}, so the updated intervals were calibrated (Table 2, Fig. 3b). The "
  f"CRPS was {f4(BN['crps'])}, against {f4(BP['crps'])} for the prior, {f4(BTR['crps'])} for the least-squares fit with *m* from the "
  f"training data, and {f4(B90['crps'])} for the fit with *m* = 0.90, which placed {pc(B90['bajo10'])} of tests below its 10th "
  f"percentile. The recommended procedure scored {f4(BS['crps'])} and the joint update {f4(BJ['crps'])}. None of these differences "
  f"was significant after adjustment (smallest Holm-adjusted *p* = {min(CV[f'B|{k}']['vs_nivel']['p_holm'] for k in ('sitio', 'descarte', 'conjunta', 'previa', 'ingenuo_m090', 'ingenuo_mtr')):.2f}); "
  f"with {BN['n']} tests from 13 publications the comparison can detect only large differences. The conflict check flagged "
  f"{pc(CV['B|conflicto_frac'])} of the sets.")
Hd("Specimen Size", 2)
P(f"Sixteen materials of {EP['k']} publications were tested at maximum particle sizes differing by a factor of two or more. With "
  "set intercepts adjusted for void ratio and materials combined within publications, the random-effects mean of β_D was "
  f"{sg(EP['b'][0], 3)} (Knapp–Hartung 95% CI {sg(EP['ci'][0][0], 3)} to {sg(EP['ci'][0][1], 3)}; Knapp and Hartung 2003) with a "
  f"between-publication standard deviation τ = {f3(EP['tau'])} (Q-profile 95% CI {f3(EP['tau_ci'][0])}–{f3(EP['tau_ci'][1])}; "
  f"Viechtbauer 2007; Fig. S4). Adding the between-set variance of the model to each estimate gave τ = {f3(EP1['tau'])}. The interval "
  "for the mean excludes the magnitude of about −0.05 to −0.15 implied by Weibull scaling of homothetic gradations, while "
  "the database mixes scalped and parallel gradations and rounded and angular particles, for which opposite trends have "
  "been reported (Varadarajan et al. 2003). For design, Leps-P uses a zero mean with the larger τ as uncertainty.")
Hd("Sources of Uncertainty at Field Scale", 2)
P(f"For an untested material of low rock quality placed at *e* = 0.65 and scaled from 0.18 m to 1 m particles, the standard "
  f"deviation of the mean-envelope ln τ at 3 MPa corresponds to about {sdφ(V3):.1f}° in φ′s. Under the adopted size model, "
  f"specimen size contributed {vs(V3, 'escala_campo'):.0f}% of the variance, between-set variability {vs(V3, 'material'):.0f}%, "
  f"publication {vs(V3, 'publicacion'):.0f}%, the asymptote and decay rate {vs(V3, 'cresta'):.0f}%, and the fixed effects "
  f"{vs(V3, 'efectos_fijos'):.0f}% (Fig. S3); at 10 MPa the asymptote contributed {vs(V10, 'cresta'):.0f}%. The size share scales "
  "with ln²(*D*f/*D*l). Because all materials share the asymptote, between-material variance decreases at very high stress; "
  "this is a property of the model, not an observation.")

# ================================================================== APPLICATION
Hd("Application: Waste Rock from a Large Open-Pit Mine")
Hd("Site Data and Conflict Check", 2)
pB = [s["p_conflicto"] for s in notused if s["lit"] == "B"]; pC = [s["p_conflicto"] for s in notused if s["lit"] == "C"]
mB = [CF[c][k]["m_libre"] for c in CF for k in CF[c] if "HF" in k]
P(f"The waste rock of a large open-pit mine comprises five rock types, A to E. Drained triaxial tests on large specimens of "
  f"scalped gradation by one laboratory were available for A, B, and C ({len(SIT['sets'])} sets); D and E were untested, and E is "
  "prone to degradation. Rock quality was taken as low for A, C, D, and E and moderate for B; illustrative mass fractions of "
  f"35%, 25%, 20%, 10%, and 10% were used. For rock type A the conflict check gave *p* = {TD['clase_A']['baja']:.2f} with the low "
  f"class and {TD['clase_A']['buena']:.3f} with the good class, so the tests were consistent with the class assigned. The check "
  f"flagged {len(notused)} of the {len(SIT['sets'])} sets (Fig. 4a): both sets of B (*p* ≤ {max(pB):.4f}) and one of C (*p* = "
  f"{pC[0]:.3f}). In the B sets the secant angle increased with stress and the specimens were still contracting at the maximum "
  "deviator, reached at 13–19% axial strain. Re-reducing them from the tabulated stress–strain curves with the deviator at "
  f"10% and 15% axial strain gave free log–log slopes of {min(mB):.2f}–{max(mB):.2f}, each from three tests, so the increase "
  "is not an artifact of the failure definition, although slopes near one are not significantly above it. A slope above one "
  "is excluded at the population level by the model, not for individual materials. Fixed-strain criteria created new "
  "conflicts in the other rock types, which confirms that site tests must be reduced with the criterion of the calibration "
  "database. Flagged sets did not update the envelope, and the remaining sets were updated jointly with a common laboratory "
  "effect; each rock type is represented by one material.")
Hd("Placement Density, Specimen Size, and Dump-Average Envelope", 2)
lE = weak(fb, "weakest (")
P(f"Laboratory specimens had void ratios of {min(s['e'] for s in SIT['sets']):.2f}–{max(s['e'] for s in SIT['sets']):.2f}; a placement "
  "void ratio of 0.65 was adopted for end-dumped waste rock, with an additional increase of 0.15 for the degradation of E. A "
  "failure surface through a mixed dump mobilizes the mass-weighted strength of the rock types it crosses; this assumes mixing "
  "along the surface, and the weakest-rock-type case bounds zoned or segregated placement. At σ′n = 3 MPa the median and 5th "
  f"percentile of the dump-average φ′s were {lbn['model']['P50'][isn(3)]:.1f}° and {lbn['model']['P05'][isn(3)]:.1f}° at laboratory "
  f"density and size, {fbn['model']['P50'][isn(3)]:.1f}° and {fbn['model']['P05'][isn(3)]:.1f}° at placement density without size "
  f"shift, and {fb['model']['P50'][isn(3)]:.1f}° and {fb['model']['P05'][isn(3)]:.1f}° with both (Table 4; Fig. 4b). Density thus "
  "lowered the median and size uncertainty the 5th percentile. A surface confined to E lowered the 5th percentile by "
  f"{fb['model']['P05'][isn(3)] - lE['P05'][isn(3)]:.1f}°, and giving D and E a laboratory effect independent of the site changed "
  f"it by {abs(weak(SITL['escenarios']['e=0.65|general']['botadero'], 'weakest (')['P05'][isn(3)] - lE['P05'][isn(3)]):.1f}°.")
Hd("Stability of a High Dump", 2)
P("Dumps 150–700 m high with an overall slope of 2H:1V on a competent, drained, level foundation were analyzed in Slide2 "
  "(Rocscience [[year]]) with the Morgenstern and Price (1965) method (half-sine interslice function) and a noncircular "
  "search optimized by simulated annealing, using the tabulated envelopes (Table 5). The unit weight was consistent with "
  "the void ratio: 17.2 kN/m³ for *e* = 0.65 (specific gravity 2.75, water content 5%), 15.7 kN/m³ for E (*e* = 0.80), and "
  f"21 kN/m³ at laboratory density. For 700 m, the factor of safety was {f2(D_('e065', 'model', 700, 'P50'))} with the median and "
  f"{f2(D_('e065', 'model', 700, 'P05'))} with the 5th-percentile dump-average envelope; at placement density without the size "
  f"shift it was {f2(D_('e065 sin escala', 'model', 700, 'P50'))} and {f2(D_('e065 sin escala', 'model', 700, 'P05'))}, and at laboratory "
  f"density and size {f2(D_('lab sin escala', 'model', 700, 'P50'))} and {f2(D_('lab sin escala', 'model', 700, 'P05'))}. The quadratic "
  f"form gave {f2(D_('x2 e065', 'model', 700, 'P50'))} and {f2(D_('x2 e065', 'model', 700, 'P05'))}: on a competent foundation the "
  "critical surface stays at a few megapascals, where the forms agree. In a Monte Carlo analysis, each of "
  f"{MCm['n']} envelope samples was analyzed in Slide2 with its own search of the critical surface. The model correlation gave "
  f"{MCm['fallas']} failure{'s' if MCm['fallas'] != 1 else ''} (probability of failure {cpi(MCm)}, Clopper–Pearson 95% interval; Clopper and Pearson 1934; lognormal "
  f"reliability index {MCm['beta_ln']:.1f}), and a surface confined to E gave {MCe['fallas']} of {MCe['n']} "
  f"({cpi(MCe)}; index {MCe['beta_ln']:.1f}). These are conditional, strength-only probabilities for a drained foundation and a "
  "two-dimensional section, which exclude placement density, unit weight, and foundation as random variables.")

# ================================================================== DISCUSSION
Hd("Discussion")
Hd("What Leps-P Adds", 2)
P("Leps-P does not predict the central strength of an untested material better than Leps 2.0, and with few site tests its "
  "advantage over simple fits is not statistically demonstrable with the available publications. Its contributions are a "
  "material-level envelope whose curvature decays with stress, strongly favored by the likelihood and consistent with the "
  "suppression of dilatancy; intervals that are calibrated where a constant dispersion is not, both without and with site "
  "tests; an updating rule with a conflict check; and a decomposition showing that, under the adopted size model, particle "
  "size and the identity of the material dominate the uncertainty of an untested waste rock. Large-scale tests of the "
  "actual waste rock, and of one material at two particle sizes, are therefore the most valuable investigations.")
Hd("Design Format", 2)
P("A procedure consistent with these results is: (1) assign rock quality and placement void ratio to each rock type, reduce "
  "any site tests with the maximum-deviator criterion, and check the conflict diagnostic; (2) update with the tests that "
  "pass (Eq. 8); (3) shift the envelopes to placement density and field particle size; (4) adopt as characteristic value for "
  "a global mechanism the 5th percentile of the mean envelope, a Bayesian cautious estimate of the mean (Prästings et al. "
  "2019) in the spirit of Eurocode 7 (CEN 2004, 2024); and (5) check the probability of failure by sampling the envelopes with "
  "stated correlations between rock types, including the weakest-rock-type case. Because the factor of safety in limit "
  "equilibrium divides the shear strength, the factor computed with the characteristic envelope equals the strength factor "
  "that the design would still tolerate; applying mining acceptance criteria calibrated for best-estimate strength (Read and "
  "Stacey 2009; Hawley and Cunning 2017) on top of it would partly count the strength uncertainty twice, since those "
  "criteria also cover model, pore-pressure, and geometric uncertainty. No target values are proposed here; the pairing of "
  "factor of safety and conditional probability of failure in Table 5 is an example to be calibrated site by site.")
Hd("Limitations", 2)
P("The validity domain is σ′n ≤ 8 MPa and, above 3 MPa, *e* ≤ 0.55, where 83 of the 90 high-stress tests lie; percentiles beyond "
  "it are model extrapolations. Tests above 3 MPa come from eight publications, the lower tail of the population is "
  "truncated at φ′1 = 31°, class labels may carry information from measured strength, and the within-material residual is "
  "laboratory scatter rather than spatial variability. The size-effect heterogeneity, the largest term at field scale, rests "
  "on 12 publications. The site application uses one laboratory and illustrative proportions. Pore pressure in saturated "
  "basal layers, wetting collapse (Oldecop and Alonso 2001), seismic loading, three-dimensional geometry, and static "
  "liquefaction are outside the scope. Leps-P can be re-estimated as drained tests at σ′3 of 5–10 MPa, tests of one material at "
  "two sizes, and in situ density data become available.")

# ================================================================== CONCLUSIONS
Hd("Conclusions")
for t in [
    "An envelope in which the dilatant part of tan φ′s decays with stress toward an asymptote, with a material-level dilatant "
    f"level and decay rate, was strongly favored by the likelihood over fixed-form alternatives; the asymptote (about {PHM:.0f}°) "
    "is a decay parameter, not a critical-state angle.",
    "Out of sample, the bounded, quadratic, and Leps 2.0 envelopes predicted new materials with indistinguishable accuracy; "
    "Leps-P intervals were calibrated at low stress, where a constant dispersion was not.",
    "Updating the level of each material with a few site tests at low stress gave calibrated predictions at higher stress; "
    "its advantage over simpler fits was not statistically significant with 13 publications.",
    "No systematic specimen-size effect was detected within materials; under the adopted model its heterogeneity dominates "
    "the uncertainty of an untested waste rock at field scale.",
    "For a 700 m dump on a competent foundation, analyzed with the Morgenstern–Price method, placement density lowered the "
    "median strength and size uncertainty the cautious strength; the envelope form did not affect the factor of safety.",
]:
    p = doc.add_paragraph(style="List Number"); runs(p, t)

# ================================================================== BACK MATTER
Hd("Data Availability Statement")
P("Some or all data, models, or code generated or used during the study are available in a repository or online in accordance "
  "with funder data retention policies (Leps-P code, parameters, and evaluation scripts: [[repository DOI]]; calibration "
  "database: Arrieta 2025). Some or all data, models, or code generated or used during the study are proprietary or "
  "confidential in nature and may only be provided with restrictions (the site test results, which may be provided in "
  "anonymized form with the permission of their owner).", indent=False)
Hd("Acknowledgments")
P("[[Acknowledgments, funding, and the permission of the owner of the site test data to be completed.]] [[Declaration of the "
  "use of AI tools, as required by ASCE policy: tool, purpose, and confirmation that the authors reviewed and take "
  "responsibility for the content.]]", indent=False)
Hd("Notation")
P("*The following symbols are used in this paper:*", indent=False)
NOT = [(r"$\mathbf{C}_\beta$", "covariance matrix of the fixed effects"), ("*D*f, *D*l", "maximum particle size in the field and in the laboratory"),
       (r"$D^2$", "prior predictive Mahalanobis distance of the site tests"), (r"$\mathbf{G}_p$, $\mathbf{G}_s$", "covariance matrices of the publication and set random effects"),
       (r"$\mathbf{J}$, $\mathbf{Z}$", "matrices of the gradients of the envelope with respect to the fixed and random effects"),
       (r"$\mathbf{P}$", "prior covariance of the random effects and fixed-effect deviations of a material"), (r"$\mathbf{S}$", "diagonal matrix of test-to-test variances"),
       ("*T*", "tan φ′∞"), (r"$\mathbf{V}$", "prior predictive covariance of the site tests"),
       (r"$\mathbf{a}$", "gradient of the envelope at the centroid of the site tests"), ("a0, a1", "parameters of the test-to-test dispersion"),
       ("*c*, *c*₂", "rock quality class; quadratic coefficient of the comparison form"), ("*e*", "initial void ratio"),
       ("*f*", "mean envelope of a new material in terms of *y*"), ("g_0", "prior variance of q_0"), (r"$\mathbf{j}$, $\mathbf{z}$", "gradients of *f* with respect to the fixed and random effects"),
       ("*m*", "exponent of the power envelope"), ("*n*", "number of site tests"), ("p_a", "atmospheric pressure (0.1013 MPa)"),
       ("q_0", "level of the envelope at the centroid of the site tests"), (r"$\mathbf{r}$, *r*i", "residuals of the site tests from the prior mean"),
       (r"$\mathbf{u}$, $\mathbf{u}^{(p)}$, $\mathbf{u}^{(s)}$", "random effects on (η, ζ): total, publication, and set"), (r"$\mathbf{v}$", "random effects and fixed-effect deviations of a material"),
       ("*x*", "ln(σ′n/σ0)"), ("*x̄*", "mean of *x* over the site tests of a set"), ("*y*", "ln(τ/σ0)"),
       ("β_c", "coefficient of rock quality class *c*"), ("β_D", "specimen-size coefficient"), ("β_e", "effect of void ratio on η"),
       ("Δ*e*, Δ*y*", "change of void ratio; shift of *y* for specimen size"), ("ε", "test-to-test deviation"), ("ζ, ζ0", "decay rate of a material and of the population"),
       ("η", "dilatant level of a material, ln tan φ′1"), ("θ", "level of the comparison forms"), ("σ0", "reference stress, 10p_a"),
       ("σ′3", "effective confining stress"), ("σ′n", "effective normal stress on the failure plane"), ("σε", "standard deviation of test-to-test scatter"),
       ("τ", "shear strength; in the size analysis, between-publication standard deviation of β_D"), ("φ′1", "secant friction angle at σ′n = σ0"),
       ("φ′∞", "asymptotic secant friction angle"), ("φ′e=0", "φ′1 extrapolated to *e* = 0 in Leps 2.0"), ("φ′s", "secant friction angle"),
       ("ω", "decrease of φ′1 per unit void ratio in Leps 2.0")]
for s_, d_ in NOT:
    P(f"{s_} = {d_};", indent=False)
Hd("Supplemental Materials")
P("Figs. S1–S4 (stress coverage of the database; volumetric strain and breakage at failure; variance decomposition; "
  "specimen-size effect by material) and Table S1 (every Mode B procedure evaluated) are available online in the ASCE Library "
  "(www.ascelibrary.org).", indent=False)

br(); Hd("References")
REFS = [
    "Alonso, E. E., M. Tapias, and J. Gili. 2012. \"Scale effects in rockfill behaviour.\" *Géotech. Lett.* 2 (3): 155–160. https://doi.org/10.1680/geolett.12.00025.",
    "Arrieta, M. 2025. \"Rockfill test database. Version 1.\" Zenodo. Accessed October 6, 2026. https://doi.org/10.5281/zenodo.14665504.",
    "Barton, N., and B. Kjærnsli. 1981. \"Shear strength of rockfill.\" *J. Geotech. Eng. Div.* 107 (7): 873–891. https://doi.org/10.1061/AJGEB6.0001167.",
    "Bolton, M. D. 1986. \"The strength and dilatancy of sands.\" *Géotechnique* 36 (1): 65–78. https://doi.org/10.1680/geot.1986.36.1.65.",
    "CEN (European Committee for Standardization). 2004. *Eurocode 7: Geotechnical design—Part 1: General rules.* EN 1997-1. Brussels, Belgium: CEN.",
    "CEN (European Committee for Standardization). 2024. *Eurocode 7: Geotechnical design—Part 1: General rules.* EN 1997-1:2024. Brussels, Belgium: CEN.",
    "Charles, J. A., and K. S. Watts. 1980. \"The influence of confining pressure on the shear strength of compacted rockfill.\" *Géotechnique* 30 (4): 353–367. https://doi.org/10.1680/geot.1980.30.4.353.",
    "Ching, J., and K.-K. Phoon. 2019. \"Constructing site-specific multivariate probability distribution model using Bayesian machine learning.\" *J. Eng. Mech.* 145 (1): 04018126. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001537.",
    "Ching, J., S. Wu, and K.-K. Phoon. 2021. \"Constructing quasi-site-specific multivariate probability distribution using hierarchical Bayesian model.\" *J. Eng. Mech.* 147 (10): 04021069. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001964.",
    "Christian, J. T., C. C. Ladd, and G. B. Baecher. 1994. \"Reliability applied to slope stability analysis.\" *J. Geotech. Eng.* 120 (12): 2180–2207. https://doi.org/10.1061/(ASCE)0733-9410(1994)120:12(2180).",
    "Clopper, C. J., and E. S. Pearson. 1934. \"The use of confidence or fiducial limits illustrated in the case of the binomial.\" *Biometrika* 26 (4): 404–413. https://doi.org/10.1093/biomet/26.4.404.",
    "Contreras, L.-F., and S. Linero-Molina. 2025. \"Hierarchical Bayesian inference of rockfill strength parameters from large-scale triaxial test data.\" *Aust. Geomech.* 60 (2): 91–106. https://doi.org/10.56295/AGJ6024.",
    "Contreras, L.-F., S. Linero-Molina, and J. Dixon. 2022. \"Bayesian approach to improve the confidence of the estimation of the shear strength of coarse mine waste using Barton's empirical criterion.\" *Aust. Geomech.* 57 (2): 163–174. https://doi.org/10.56295/AGJ5728.",
    "Coop, M. R., K. K. Sørensen, T. Bodas Freitas, and G. Georgoutsos. 2004. \"Particle breakage during shearing of a carbonate sand.\" *Géotechnique* 54 (3): 157–163. https://doi.org/10.1680/geot.2004.54.3.157.",
    "Cylwik, S. D., M. Arrieta, J. C. Cunning, and P. M. Hawley. 2025. \"Update to the Leps strength envelopes for rockfill based on 900 large-scale triaxial tests.\" *Int. J. Geomech.* 25 (12): 04025285. https://doi.org/10.1061/IJGNAI.GMENG-12050.",
    "De Mello, V. F. B. 1977. \"Reflections on design decisions of practical significance to embankment dams.\" *Géotechnique* 27 (3): 281–355. https://doi.org/10.1680/geot.1977.27.3.281.",
    "Einav, I. 2007. \"Breakage mechanics—Part I: Theory.\" *J. Mech. Phys. Solids* 55 (6): 1274–1297. https://doi.org/10.1016/j.jmps.2006.11.003.",
    "Frossard, E., W. Hu, C. Dano, and P.-Y. Hicher. 2012. \"Rockfill shear strength evaluation: A rational method based on size effects.\" *Géotechnique* 62 (5): 415–427. https://doi.org/10.1680/geot.10.P.079.",
    "Gneiting, T., and A. E. Raftery. 2007. \"Strictly proper scoring rules, prediction, and estimation.\" *J. Am. Stat. Assoc.* 102 (477): 359–378. https://doi.org/10.1198/016214506000001437.",
    "Hardin, B. O. 1985. \"Crushing of soil particles.\" *J. Geotech. Eng.* 111 (10): 1177–1192. https://doi.org/10.1061/(ASCE)0733-9410(1985)111:10(1177).",
    "Hawley, M., and J. Cunning, eds. 2017. *Guidelines for mine waste dump and stockpile design.* Clayton, Australia: CSIRO Publishing.",
    "Holm, S. 1979. \"A simple sequentially rejective multiple test procedure.\" *Scand. J. Stat.* 6 (2): 65–70.",
    "Jefferies, M., and K. Been. 2016. *Soil liquefaction: A critical state approach.* 2nd ed. Boca Raton, FL: CRC Press. https://doi.org/10.1201/b19114.",
    "Knapp, G., and J. Hartung. 2003. \"Improved tests for a random effects meta-regression with a single covariate.\" *Stat. Med.* 22 (17): 2693–2710. https://doi.org/10.1002/sim.1482.",
    "Lade, P. V., J. A. Yamamuro, and P. A. Bopp. 1996. \"Significance of particle crushing in granular materials.\" *J. Geotech. Eng.* 122 (4): 309–316. https://doi.org/10.1061/(ASCE)0733-9410(1996)122:4(309).",
    "Laird, N. M., and J. H. Ware. 1982. \"Random-effects models for longitudinal data.\" *Biometrics* 38 (4): 963–974. https://doi.org/10.2307/2529876.",
    "Leps, T. M. 1970. \"Review of shearing strength of rockfill.\" *J. Soil Mech. Found. Div.* 96 (4): 1159–1170. https://doi.org/10.1061/JSFEAQ.0001433.",
    "Lindstrom, M. J., and D. M. Bates. 1990. \"Nonlinear mixed effects models for repeated measures data.\" *Biometrics* 46 (3): 673–687. https://doi.org/10.2307/2532087.",
    "Linero, S., C. Palma, and R. Apablaza. 2007. \"Geotechnical characterisation of waste material in very high dumps with large scale triaxial testing.\" In *Proc., Int. Symp. on Rock Slope Stability in Open Pit Mining and Civil Engineering*, edited by Y. Potvin, 59–75. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/708_2.",
    "Marachi, N. D., C. K. Chan, and H. B. Seed. 1972. \"Evaluation of properties of rockfill materials.\" *J. Soil Mech. Found. Div.* 98 (1): 95–114. https://doi.org/10.1061/JSFEAQ.0001735.",
    "Marsal, R. J. 1973. \"Mechanical properties of rockfill.\" In *Embankment-dam engineering: Casagrande volume*, edited by R. C. Hirschfeld and S. J. Poulos, 109–200. New York: Wiley.",
    "McDowell, G. R., and M. D. Bolton. 1998. \"On the micromechanics of crushable aggregates.\" *Géotechnique* 48 (5): 667–679. https://doi.org/10.1680/geot.1998.48.5.667.",
    "Morgenstern, N. R., and V. E. Price. 1965. \"The analysis of the stability of general slip surfaces.\" *Géotechnique* 15 (1): 79–93. https://doi.org/10.1680/geot.1965.15.1.79.",
    "Oldecop, L. A., and E. E. Alonso. 2001. \"A model for rockfill compressibility.\" *Géotechnique* 51 (2): 127–139. https://doi.org/10.1680/geot.2001.51.2.127.",
    "Ovalle, C., E. Frossard, C. Dano, W. Hu, S. Maiolino, and P.-Y. Hicher. 2014. \"The effect of size on the strength of coarse rock aggregates and large rockfill samples through experimental data.\" *Acta Mech.* 225 (8): 2199–2216. https://doi.org/10.1007/s00707-014-1127-z.",
    "Ovalle, C., S. Linero, C. Dano, E. Bard, P.-Y. Hicher, and R. Osses. 2020. \"Data compilation from large drained compression triaxial tests on coarse crushable rockfill materials.\" *J. Geotech. Geoenviron. Eng.* 146 (9): 06020013. https://doi.org/10.1061/(ASCE)GT.1943-5606.0002314.",
    "Prästings, A., J. Spross, and S. Larsson. 2019. \"Characteristic values of geotechnical parameters in Eurocode 7.\" *Proc. Inst. Civ. Eng. Geotech. Eng.* 172 (4): 301–311. https://doi.org/10.1680/jgeen.18.00057.",
    "Read, J., and P. Stacey, eds. 2009. *Guidelines for open pit slope design.* Collingwood, Australia: CSIRO Publishing.",
    "Rocscience. [[year]]. *Slide2.* Version 9.043. Toronto: Rocscience.",
    "Valenzuela, L., E. Bard, J. Campaña, and M. E. Anabalón. 2008. \"High waste rock dumps: Challenges and developments.\" In *Proc., 1st Int. Seminar on the Management of Rock Dumps, Stockpiles and Heap Leach Pads*, 65–78. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/802_6.",
    "Varadarajan, A., K. G. Sharma, K. Venkatachalam, and A. K. Gupta. 2003. \"Testing and modeling two rockfill materials.\" *J. Geotech. Geoenviron. Eng.* 129 (3): 206–218. https://doi.org/10.1061/(ASCE)1090-0241(2003)129:3(206).",
    "Vehtari, A., A. Gelman, D. Simpson, B. Carpenter, and P.-C. Bürkner. 2021. \"Rank-normalization, folding, and localization: An improved R̂ for assessing convergence of MCMC.\" *Bayesian Anal.* 16 (2): 667–718. https://doi.org/10.1214/20-BA1221.",
    "Viechtbauer, W. 2007. \"Confidence intervals for the amount of heterogeneity in meta-analysis.\" *Stat. Med.* 26 (1): 37–52. https://doi.org/10.1002/sim.2514.",
]
for ref in REFS:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5); runs(p, ref, False)

br(); Hd("Figure Captions List")
for c in [
    "Fig. 1. Design envelopes of Leps-P for a new material tested by a new laboratory, at laboratory specimen size, by rock quality and void ratio: median (thick lines) and 5th percentile (thin lines) of the mean envelope. The shaded zone exceeds the largest normal stress in the database.",
    "Fig. 2. Profile log-likelihood of the asymptotic angle and the population decay rate relative to its maximum (dot).",
    "Fig. 3. Out-of-sample evaluation, leaving out each publication: (a) coverage of the 80% interval by range of normal stress without site tests; and (b) mean CRPS of ln τ for held-out tests at σ′n ≥ 2 MPa with site tests at σ′n ≤ 1.5 MPa (bars: 95% bootstrap intervals over publications; filled: recommended procedure).",
    "Fig. 4. Waste rock of the example: (a) laboratory-scale envelopes by rock type with the site tests (open symbols: tests flagged by the conflict check); and (b) envelopes of placed waste rock (void ratio 0.65, maximum particle size 1 m) by rock type and dump-average envelope. Dotted vertical lines mark the largest normal stress of the site tests.",
]:
    P(c, indent=False)

rg = lambda i, d=3: f"{HP[i][1]:.{d}f}–{HP[i][2]:.{d}f}"
tabla(1, "Parameters of Leps-P", ["Parameter", "Estimate", "Standard error or 90% credible interval"],
      [["β_c, good rock quality", f3(b["good"]), f3(SE["good"])], ["β_c, moderate rock quality", f3(b["moderate"]), f3(SE["moderate"])],
       ["β_c, low rock quality", f3(b["low"]), f3(SE["low"])], ["β_e", neg(f3(b["e"])), f3(SE["e"])],
       ["φ′∞ (°), weighted mean over the ridge", f"{PHM:.1f}", f"{phi_lo:.0f}–{phi_hi:.0f} (within 3 log-likelihood units)"], ["ζ0, weighted mean over the ridge", f2(ZEM), "—"],
       ["SD of η, publication", f3(HP[0][0]), rg(0)], ["SD of ζ, publication", f3(HP[1][0]), rg(1)],
       ["Correlation of η and ζ, publication", neg(f2(HP[2][0])), neg(f"{HP[2][1]:.2f} to {HP[2][2]:.2f}")],
       ["SD of η, set", f3(HP[3][0]), rg(3)], ["SD of ζ, set", f3(HP[4][0]), rg(4)],
       ["Correlation of η and ζ, set", neg(f2(HP[5][0])), neg(f"{HP[5][1]:.2f} to {HP[5][2]:.2f}")],
       ["σε at σ′n = σ0", f3(HP[6][0]), rg(6)], ["a1", neg(f3(HP[7][0])), neg(f"{HP[7][1]:.3f} to {HP[7][2]:.3f}")]],
      f"{N['ensayos']} tests, {N['sets']} sets, {N['pub']} publications. Fixed effects at the maximum-likelihood point of the ridge "
      f"(φ′∞ = {FA['phi_cs']:.0f}°, ζ0 = {FA['zeta']}); variance parameters: posterior median and 90% credible interval. Median φ′1 at e = 0: "
      f"{phi1['good']:.1f}°, {phi1['moderate']:.1f}°, and {phi1['low']:.1f}° for good, moderate, and low rock quality.")
r_ = lambda o: [pc(o["cob80"]), pc(o["bajo10"]), f4(o["crps"])]
tabla(2, "Out-of-sample evaluation (leave one publication out)", ["Procedure", "80% coverage", "Below P10", "CRPS"],
      [["Mode A, all stresses", "", "", ""], ["   Leps-P, bounded form"] + r_(A_["acotada"]), ["   Leps-P, quadratic form"] + r_(A_["cuadratica"]),
       ["   Leps 2.0 envelope, constant dispersion"] + r_(A_["leps2_ref"]), ["Mode A, σ′n > 2.75 MPa", "", "", ""],
       ["   Leps-P, bounded form"] + r_(Ah["acotada"]), ["   Leps-P, quadratic form"] + r_(Ah["cuadratica"]), ["   Leps 2.0 envelope, constant dispersion"] + r_(Ah["leps2_ref"]),
       ["Mode B, site tests at σ′n ≤ 1.5 MPa", "", "", ""], ["   Prior, no site tests"] + r_(BP), ["   Fit of φ′1, m = 0.90"] + r_(B90),
       ["   Fit of φ′1, m from training data"] + r_(BTR), ["   Joint update of both random effects"] + r_(BJ),
       ["   Level update of each set"] + r_(BN), ["   Recommended procedure"] + r_(BS)],
      f"Mode A: {A_['n']} tests ({A_['n_pub']} publications); above 2.75 MPa, {Ah['n']} tests ({Ah['n_pub']} publications). Mode B: {BN['n']} tests at "
      "σ′n ≥ 2 MPa (46 sets, 13 publications). Nominal values: 80% coverage and 10% below P10. CRPS in units of ln τ. No difference was significant after Holm adjustment.")
fil = []
for cal, nom in (("buena", "Good"), ("moderada", "Moderate"), ("baja", "Low")):
    for e in (0.3, 0.45, 0.65):
        v = TD["tabla"][f"{cal}|{e}"]; fil.append([nom if e == 0.3 else "", f"{e:.2f}"] + [f"{a:.1f} / {c:.1f}" + ("ᵃ" if (e == 0.65 and s_ > 3) or s_ > 7.9 else "")
                                                                                      for a, c, s_ in zip(v["P50"], v["P05"], TD["sn"])])
tabla(3, "Design values of the secant friction angle φ′s (°) for an untested material: median / 5th percentile of the mean envelope",
      ["Rock quality", "e"] + [f"{s:g} MPa" for s in TD["sn"]], fil,
      "Laboratory specimen size, new laboratory, no site tests, drained, maximum-deviator criterion; for global mechanisms. "
      "Field values are lower when the specimen-size uncertainty is added (about 3° at the 5th percentile for a size ratio of 5.6). "
      "Not applicable to degradable, argillaceous, or fines-rich waste below the database range (φ′1 ≥ 31°). "
      "ᵃExtrapolated beyond the validity domain (σ′n ≤ 8 MPa; e ≤ 0.55 above 3 MPa).")
fil = []; lab = SIT["escenarios"]["e=lab|ninguno"]["litologias"]; fld = SIT["escenarios"]["e=0.65|general"]["litologias"]
for k in "ABCDE":
    a_ = lab[k]; f_ = fld[k]
    fil.append([f"Rock type {k}{' (untested)' if k in 'DE' else ''}", f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}",
                f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}", f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}"])
for nm, k in (("Dump average, model correlation", "model"), ("Dump average, full correlation", "comonotonic")):
    a_ = lbn[k]; f_ = fb[k]
    fil.append([nm, f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}", f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}"])
tabla(4, "Secant friction angle φ′s (°) of the waste rock of the example, median / 5th percentile of the mean envelope",
      ["Material", "Laboratory, 3 MPa", "Placed, 1 MPa", "Placed, 3 MPa"], fil,
      "Laboratory: specimen void ratio and size. Placed: void ratio 0.65 (rock type E, 0.80), maximum particle size 1 m, size effect with zero mean. "
      "The full-correlation case sorts the samples of each rock type at every stress, an upper bound on dependence. No site tests above σ′n ≈ 4.6 MPa.")
fil = []
for esc, nm in (("model", "Dump average, model correlation"), ("comonotonic", "Dump average, full correlation"), ("weakest tested A", "Weakest tested rock type only"), ("weakest E", "Untested, degraded rock type E only")):
    fil.append([nm] + [f"{f2(D_('e065', esc, h, 'P50'))} / {f2(D_('e065', esc, h, 'P05'))}" if h == 700 or esc == "model" or True else "" for h in (150, 400, 700)])
tabla(5, "Factor of safety of dumps on a competent level foundation (Morgenstern–Price, Slide2), placed waste rock: median / 5th-percentile envelope",
      ["Envelope", "150 m", "400 m", "700 m"], fil,
      f"Unit weight 17.2 kN/m³ (rock type E only: 15.7 kN/m³). Monte Carlo at 700 m with {MCm['n']} envelope samples, each with its own search: "
      f"model correlation {MCm['fallas']} failure{'s' if MCm['fallas'] != 1 else ''} ({cpi(MCm)}; lognormal reliability index {MCm['beta_ln']:.1f}); rock type E only {MCe['fallas']} failure{'s' if MCe['fallas'] != 1 else ''} ({cpi(MCe)}; index {MCe['beta_ln']:.1f}). "
      f"700 m, model correlation: without size shift {f2(D_('e065 sin escala', 'model', 700, 'P50'))} / {f2(D_('e065 sin escala', 'model', 700, 'P05'))}; "
      f"laboratory density and size (21 kN/m³) {f2(D_('lab sin escala', 'model', 700, 'P50'))} / {f2(D_('lab sin escala', 'model', 700, 'P05'))}; "
      f"quadratic form {f2(D_('x2 e065', 'model', 700, 'P50'))} / {f2(D_('x2 e065', 'model', 700, 'P05'))}.")
out = AQ / "LepsP_IJGM_manuscript_r4.docx"; doc.save(str(out)); print("written", out, "| equations:", A.NEQ[0])
