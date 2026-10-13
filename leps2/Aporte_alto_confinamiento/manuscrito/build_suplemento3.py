"""Supplemental Materials of the Leps-P manuscript (round 3): Fig. S1 (specimen-size effect) and Table S1 (every Mode B
procedure evaluated). Numbers read from result files."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

AQ = Path(__file__).parent; AN = AQ.parent / "analisis"; FIG = AQ.parent / "figuras" / "lepsp3"
J = lambda p: json.load(open(p, encoding="utf-8"))
C3 = J(AN / "lepsp_cv3_metricas.json"); C2 = J(AN / "lepsp_cv2_metricas.json"); E3 = J(AN / "lepsp_escala3.json")
doc = Document(); st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
doc.add_paragraph().add_run("Supplemental Materials — Leps-P: Probabilistic Strength Envelopes with Dilatancy-Bounded Curvature for Rockfill and Mine Waste Rock").bold = True
doc.add_picture(str(FIG / "FigS1_size_effect.png"), width=Inches(5.2))
P = E3["publicaciones R ≥ 2"]
doc.add_paragraph(f"Fig. S1. Within-material specimen-size coefficient βD for materials of the database tested at more than one maximum particle size "
                  f"(bars: 95% confidence intervals; set intercepts adjusted for void ratio, with the between-set variance unrelated to size added). "
                  f"Hatched band: Knapp–Hartung 95% confidence interval of the random-effects mean across {P['k']} publications with size ratio ≥ 2 "
                  f"({P['b'][0]:+.3f}; {P['ci'][0][0]:+.3f} to {P['ci'][0][1]:+.3f}). One value outside the axis is annotated.")
doc.add_paragraph().add_run("Table S1. Mode B: every procedure evaluated (held-out tests at σ′n ≥ 2 MPa predicted from tests at σ′n ≤ 1.5 MPa of the same set)").bold = True
filas = []
NOM3 = {"nivel": "Level-only update (prespecified primary)", "previa": "Prior, no site tests", "conjunta": "Joint update of level and slope",
        "procedimiento": "Level-only update after discarding flagged sets", "ingenuo_m090": "Least-squares fit of φ′1, m = 0.90",
        "ingenuo_mtr": "Least-squares fit of φ′1, m = median within-set slope of training data"}
for k, nm in NOM3.items():
    o = C3[f"B|{k}"]; filas.append(("Bounded (round 3)", nm, o["cob80"], o["bajo10"], o["crps"]))
NOM2 = {"conjunta": "Joint update of level and slope", "nivel": "Level-only update", "ingenuo m=0.86": "Least-squares fit of φ′1, m = 0.86 (full-data estimate)",
        "ingenuo m=0.9": "Least-squares fit of φ′1, m = 0.90"}
for k, nm in NOM2.items():
    o = C2[f"B|{k}"]; filas.append(("Quadratic (earlier analysis)", nm, o["cob80"], o["bajo10"], o["crps"]))
for r in (2, 3, 4):
    for p in (0.01, 0.05, 0.1):
        o = C2[f"B|regla r{r} p{p}"]
        filas.append(("Quadratic (earlier analysis)", f"Mixed rule: joint update if stress ratio ≥ {r} and p ≥ {p}, else level-only", o["cob80"], o["bajo10"], o["crps"]))
t = doc.add_table(rows=1 + len(filas), cols=5)
for j, h in enumerate(["Envelope form", "Procedure", "80% coverage", "Below P10", "CRPS"]):
    t.rows[0].cells[j].text = h
for i, f in enumerate(filas, 1):
    for j, v in enumerate((f[0], f[1], f"{100 * f[2]:.0f}%", f"{100 * f[3]:.0f}%", f"{f[4]:.4f}")):
        t.rows[i].cells[j].text = v
for row in t.rows:
    for c in row.cells:
        for pp in c.paragraphs:
            for r_ in pp.runs:
                r_.font.size = Pt(9)
doc.add_paragraph("Note: The bounded-form procedures were specified before the bounded form was evaluated; the level-only update was the primary "
                  "procedure. The quadratic-form results are from an earlier analysis in which the rule thresholds were explored; they are reported "
                  "for completeness and were not used to select the primary procedure of the bounded form. CRPS in units of ln τ.")
cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = "Leps-P supplemental materials"
doc.save(str(AQ / "LepsP_IJGM_supplement_r3.docx")); print("written supplement")
