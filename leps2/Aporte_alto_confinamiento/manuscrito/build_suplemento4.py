"""Supplemental Materials of the Leps-P manuscript (round 4): Figs. S1–S4 and Table S1 (every Mode B procedure).
Numbers read from result files."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

AQ = Path(__file__).parent; AN = AQ.parent / "analisis"; FIG = AQ.parent / "figuras" / "lepsp4"
J = lambda p: json.load(open(p, encoding="utf-8"))
CV = J(AN / "lepsp_cv4_metricas.json"); E3 = J(AN / "lepsp_escala3.json")["publicaciones R ≥ 2"]; DG = J(AN / "lepsp_diag_alta.json")
doc = Document(); st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
doc.add_paragraph().add_run("Supplemental Materials — Leps-P: Probabilistic Strength Envelopes with Stress-Decaying Curvature for Rockfill and Mine Waste Rock").bold = True
neg = lambda s: s.replace("-", "−")
NBG = DG["bg_excluidos_mayor_100"]; RHO = neg(f"{DG['res_vs_bg_alta']['spearman']:+.2f}"); NRB = DG["res_vs_bg_alta"]["n"]
EB, EL, EH = (neg(f"{v:+.3f}") for v in (E3["b"][0], E3["ci"][0][0], E3["ci"][0][1]))
FIGS = [("FigS1_stress_coverage.png", 6.0, "Fig. S1. Distribution of effective normal stress on the failure plane in the calibration database."),
        ("FigS2_volumetric_breakage.png", 6.0, f"Fig. S2. Net volumetric strain at the maximum deviator (negative: contraction) and breakage index of Marsal (1973) versus σ′n in the calibration database; shaded: σ′n > 2.75 MPa. Reported breakage indices above 100%, which are impossible by definition, were excluded ({NBG} tests). Above 2.75 MPa the residuals of Leps-P were unrelated to the breakage index (Spearman ρ = {RHO}, n = {NRB})."),
        ("FigS3_variance.png", 3.5, "Fig. S3. Shares of the variance of the mean-envelope ln τ for an untested material of low rock quality placed at e = 0.65, scaled from 0.18 m to 1 m maximum particle size."),
        ("FigS4_size_effect.png", 5.6, f"Fig. S4. Within-material specimen-size coefficient βD (bars: 95% confidence intervals; set intercepts adjusted for void ratio). Hatched band: Knapp–Hartung 95% confidence interval of the random-effects mean across {E3['k']} publications with size ratio ≥ 2 ({EB}; {EL} to {EH}). One value outside the axis is annotated.")]
for f, w, cap in FIGS:
    doc.add_picture(str(FIG / f), width=Inches(w)); doc.add_paragraph(cap)
doc.add_paragraph().add_run("Table S1. Mode B: every procedure evaluated (held-out tests at σ′n ≥ 2 MPa predicted from tests at σ′n ≤ 1.5 MPa of the same set)").bold = True
NOM = {"previa": "Prior, no site tests", "ingenuo_m090": "Least-squares fit of φ′1, m = 0.90", "ingenuo_mtr": "Least-squares fit of φ′1, m = median within-set slope of training data",
       "conjunta": "Joint update of both random effects", "nivel": "Level update of each set (Eq. 8)", "descarte": "Level update after discarding flagged sets",
       "sitio": "Recommended procedure: discard flagged sets, joint update with shared laboratory effect"}
t = doc.add_table(rows=1 + len(NOM), cols=7)
for j, h in enumerate(["Procedure", "80% coverage", "Below P10", "SD of z", "CRPS", "Δ vs level update (95% CI)", "Holm-adjusted p"]):
    t.rows[0].cells[j].text = h
for i, (k, nm) in enumerate(NOM.items(), 1):
    o = CV[f"B|{k}"]; c = o.get("vs_nivel")
    vals = [nm, f"{100 * o['cob80']:.0f}%", f"{100 * o['bajo10']:.0f}%", f"{o['sd_z']:.2f}", f"{o['crps']:.4f}",
            neg(f"{c['media']:+.4f} ({c['ic95'][0]:+.4f} to {c['ic95'][1]:+.4f})") if c else "—", f"{c['p_holm']:.2f}" if c else "—"]
    for j, v in enumerate(vals):
        t.rows[i].cells[j].text = v
for row in t.rows:
    for c in row.cells:
        for pp in c.paragraphs:
            for r_ in pp.runs:
                r_.font.size = Pt(9)
doc.add_paragraph("Note: 64 tests, 46 sets, 13 publications. The procedures were developed while examining earlier versions of the model on the "
                  "same held-out sets; the comparisons are exploratory. p-values from exact sign-flip tests on publication means, Holm-adjusted within "
                  "Mode B. CRPS in units of ln τ.")
cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = "Leps-P supplemental materials"
doc.save(str(AQ / "LepsP_IJGM_supplement_r4.docx")); print("written supplement r4")
