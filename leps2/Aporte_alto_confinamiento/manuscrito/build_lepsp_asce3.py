"""Leps-P manuscript, round 3 (ASCE Int. J. Geomech., Technical Paper). Central model: dilatancy-bounded envelope.
All numbers are read from result files. ASCE format: double spacing, continuous line numbers, unnumbered headings,
equations as Word equation objects (OMML), figure caption list after the references, tables at the end."""
import json, math, re
from pathlib import Path
import numpy as np, pandas as pd
from scipy.stats import norm, wilcoxon
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from lxml import etree
import latex2mathml.converter as l2m

AQ = Path(__file__).parent; R = AQ.parent; AN = R / "analisis"; AA = R / "analisis_antamina"
J = lambda p: json.load(open(p, encoding="utf-8"))
PAR = J(R / "leps3" / "lepsp_parametros.json"); FA = PAR["formas"]["acotada"]; N = PAR["n"]
M3 = J(AN / "lepsp_mcmc3.json"); CV = J(AN / "lepsp_cv3_metricas.json"); ESC = J(AN / "lepsp_escala3.json"); VAR = J(AN / "lepsp_varianza.json")
GR = J(AN / "lepsp_acotada_grilla.json"); DPL = J(AN / "lepsp_desplegada.json"); TD = J(AN / "lepsp_tabla_diseno.json"); SX = J(AN / "lepsp_acotada_sitio.json")
SIT = J(AA / "lepsp_sitio3.json"); SITQ = J(AA / "lepsp_sitio3_x2.json"); CF = J(AA / "a7_criterio_falla.json"); EST = J(AN / "lepsp_estabilidad3.json")
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))

# ------------------------------------------------------------------ numbers
CL = ["good", "moderate", "low"]; b = dict(zip(FA["nombres"], FA["beta"])); SE = dict(zip(FA["nombres"], np.sqrt(np.diag(FA["cov_beta"]))))
HP = FA["hiper"]                       # [median, p5, p95] for sdθp, sdmp, ρp, sdθs, sdms, ρs, σε(0), a1
deg = lambda th: math.degrees(math.atan(math.exp(th)))
phi1 = {c: deg(b[c] + b["e"] * 0) for c in CL}
pc = lambda v: f"{100 * v:.0f}%"; f2 = lambda v: f"{v:.2f}"; f3 = lambda v: f"{v:.3f}"; f4 = lambda v: f"{v:.4f}"
neg = lambda s: s.replace("-", "−"); sg = lambda v, dd=4: neg(f"{v:+.{dd}f}")
# share of between-publication variance at σ0 (posterior)
sh = []
for s in FA["muestras"]:
    p = s["psi"]; sh.append(math.exp(2 * p[0]) / (math.exp(2 * p[0]) + math.exp(2 * p[3])))
SHARE = (np.median(sh), np.percentile(sh, 5), np.percentile(sh, 95))
# evidence for bounded curvature
ML_ac, ML_q = M3["ml"]["acotada"], M3["ml"]["x2"]
mlg = np.array(GR["ml"]); i0 = GR["phi"].index(0.0); dML0 = float(np.max(mlg) - np.max(mlg[i0]))
from scipy.stats import chi2  # noqa: E402
p_LR = float(1 - chi2.cdf(2 * dML0, 1))
ridge = [c for c in FA["cresta"] if c["ml"] >= max(x["ml"] for x in FA["cresta"]) - 3]
phi_lo, phi_hi = min(c["phi_cs"] for c in ridge), max(c["phi_cs"] for c in ridge)
PHM, ZEM = M3["acotada_cresta"]["phi_media"], M3["acotada_cresta"]["zeta_media"]
Bs = pd.read_csv(AN / "lepsp_acotada_B.csv")
crps = lambda mu, s, y: s * ((y - mu) / s * (2 * norm.cdf((y - mu) / s) - 1) + 2 * norm.pdf((y - mu) / s) - 1 / np.sqrt(np.pi))
Bs["d"] = crps(Bs.mu_b, Bs.sd_b, Bs.y) - crps(Bs.mu_q, Bs.sd_q, Bs.y); pm_ = Bs.groupby("pub").d.mean()
SHAPE = dict(n=len(Bs), sets=Bs.Set.nunique(), pubs=len(pm_), mejora=int((pm_ < 0).sum()), p=float(wilcoxon(pm_).pvalue),
             cq=float(crps(Bs.mu_q, Bs.sd_q, Bs.y).mean()), cb=float(crps(Bs.mu_b, Bs.sd_b, Bs.y).mean()),
             bq=float((Bs.y - Bs.mu_q).mean()), bb=float((Bs.y - Bs.mu_b).mean()))
A_ = CV["A|todo"]; Ah = CV["A|sn>2.75"]; Al = CV["A|sn<0.5"]
BN, BP, BJ, BPR, B90, BTR = (CV[f"B|{k}"] for k in ("nivel", "previa", "conjunta", "procedimiento", "ingenuo_m090", "ingenuo_mtr"))
dg = lambda o: f"{sg(o['media'])} (95% CI {sg(o['ic95'][0])} to {sg(o['ic95'][1])}; Wilcoxon *p* = {o['wilcoxon_p']:.3f}, {o['pubs_mejora']} of {o['n_pub']} publications)"
EP = ESC["publicaciones R ≥ 2"]; EG = ESC["grupos R ≥ 2"]; MR = ESC["metaregresion"]
V3 = VAR["low e=0.65 σ′n=3"]; V10 = VAR["low e=0.65 σ′n=10"]; vs = lambda V, k: 100 * V["comp"][k] / V["total"]
SNs = np.array(SIT["sn"]); isn = lambda s: int(np.argmin(np.abs(SNs - s)))
used = [s for s in SIT["sets"] if s["usado"]]; notused = [s for s in SIT["sets"] if not s["usado"]]
DE = lambda k: EST[k]["det"]; MC = lambda k: EST[k]["mc"]
kA = "acotada e=0.65 γ17|700|"; kL = "acotada e=lab γ21|700|"; kQ = "x2 e=0.65 γ17|700|"
weak = lambda pref, t: [k for k in EST if k.startswith(pref + t)][0]
K_MOD, K_COM, K_WT, K_W = kA + "model", kA + "comonotonic", weak(kA, "weakest tested"), weak(kA, "weakest (")
_bn = pd.read_csv(AN / "lepsp_cv3_B.csv").query("variante == 'nivel'"); NB, NBP = len(_bn), _bn.pub.nunique()
cpi = lambda m: f"{100 * m['pf_ic95'][0]:.1f}–{100 * m['pf_ic95'][1]:.1f}%"

# ------------------------------------------------------------------ document setup
doc = Document()
cp = doc.core_properties; cp.author = ""; cp.title = "Leps-P"; cp.comments = ""; cp.last_modified_by = ""; cp.subject = ""; cp.keywords = ""
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.line_spacing = 2.0; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
for m_ in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m_, Inches(1))
ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); sec._sectPr.append(ln)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tag, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    r_ = fp.add_run()
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r_._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r_._r.append(e)


def omml(latex):
    return XSL(etree.fromstring(l2m.convert(latex).encode())).getroot()


SYM = [("b^(p)", r"\mathbf{b}^{(p)}"), ("b^(s)", r"\mathbf{b}^{(s)}"), ("C_β", r"\mathbf{C}_\beta"), ("β̂", r"\hat{\boldsymbol{\beta}}"),
       ("β_c", r"\beta_c"), ("β_e", r"\beta_e"), ("β_D", r"\beta_D"), ("G_p", r"\mathbf{G}_p"), ("G_s", r"\mathbf{G}_s"),
       ("σε²", r"\sigma_\varepsilon^2"), ("σε", r"\sigma_\varepsilon"), ("φ′cs", r"\phi'_{cs}"), ("φ′e=0", r"\phi'_{e=0}"),
       ("φ′1", r"\phi'_1"), ("φ′s", r"\phi'_s"), ("σ′n", r"\sigma'_n"), ("σ′3", r"\sigma'_3"), ("σ′v", r"\sigma'_v"),
       ("σ0", r"\sigma_0"), ("p_a", r"p_a"), ("*D*max", r"D_{max}"), ("*D*f", r"D_f"), ("*D*l", r"D_l"), ("q_0", r"q_0"), ("g_0", r"g_0"),
       ("*x*i", r"x_i"), ("*r*i", r"r_i"), ("*x̄*", r"\bar{x}"), ("x̃", r"\tilde{\mathbf{x}}"), ("**z**", r"\mathbf{z}"), ("a0, a1", r"a_0, a_1"),
       ("D², ", r"D^2, ")]


def simbolos(text):
    parts = re.split(r"(\$[^$]+\$|https?://\S+)", text); out = []
    for s in parts:
        if s.startswith("$") or s.startswith("http"):
            out.append(s); continue
        for a_, l_ in SYM:
            s = s.replace(a_, f"${l_}$")
        out.append(s)
    return "".join(out)


def runs(p, text, _sym=True):
    if _sym:
        text = simbolos(text)
    for tok in re.split(r"(\$[^$]+\$|\*[^*]+\*|\[\[[^\]]+\]\])", text):
        if not tok:
            continue
        if tok.startswith("$"):
            p._p.append(parse_xml(etree.tostring(omml(tok[1:-1]))))
        elif tok.startswith("*") and "$" in tok:
            runs(p, tok[1:-1], False)
        elif tok.startswith("*"):
            p.add_run(tok[1:-1]).italic = True
        elif tok.startswith("[["):
            rr = p.add_run(tok[2:-2]); rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p.add_run(tok)
    return p


def P(text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    return runs(p, text)


def Hd(text, level=1):
    p = doc.add_paragraph(); rr = p.add_run(text); rr.bold = level == 1; rr.italic = level == 2; p.paragraph_format.keep_with_next = True


NEQ = [0]


def EQ(latex):
    NEQ[0] += 1; p = doc.add_paragraph(); pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(3.25), alignment=1); pf.tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    p.add_run("\t"); p._p.append(parse_xml(etree.tostring(omml(latex)))); p.add_run(f"\t({NEQ[0]})")


def br():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================== TITLE, ABSTRACT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Leps-P: Probabilistic Strength Envelopes with Dilatancy-Bounded Curvature for Rockfill and Mine Waste Rock").bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, "[[Author names, ASCE grade, affiliations, and corresponding author e-mail to be completed.]]")
Hd("Abstract")
P("Reliability-based design of high waste dumps needs the drained shear strength of rockfill and mine waste rock as a "
  "distribution, together with its dependence on stress and a rule to combine published experience with site tests. "
  "Leps-P is a three-level hierarchical model of the secant friction angle in which the dilatant part of strength decays "
  "with stress toward an asymptote, so that the curvature of the envelope decreases as dilatancy is suppressed. It was "
  f"calibrated on {N['ensayos']} large-scale triaxial tests from {N['pub']} publications in the public Leps 2.0 database and "
  "evaluated by leaving out one publication at a time. The data supported a curvature that decays with stress (likelihood "
  f"ratio test, *p* < 0.001) but did not identify the asymptotic angle, which was therefore averaged over the likelihood "
  f"ridge. Out of sample, the bounded form scored better than a quadratic log–log form (continuous ranked probability score "
  f"{f4(A_['acotada']['crps'])} versus {f4(A_['cuadratica']['crps'])}, Wilcoxon *p* = {A_['acotada-cuadratica|crps']['wilcoxon_p']:.2f}) "
  "and removed the low bias of the median above 2.75 MPa. With a few site tests at low stress, a prespecified level-only "
  f"update predicted the strength at higher stress better than a least-squares fit with a fixed curvature (score "
  f"{f4(BN['crps'])} versus {f4(BTR['crps'])}). At field scale, the uncertainty of an untested waste rock came mainly from "
  f"specimen size and from material and laboratory variability; the unidentified asymptote contributed about "
  f"{vs(V10, 'cresta'):.0f}% of the variance at 10 MPa. Within-material comparisons did not detect a systematic specimen-size effect. "
  "For a 700 m dump of waste rock from an open-pit mine, placement density and untested rock types controlled the margin "
  "of safety. The model, design charts, and code allow re-estimation as tests at higher stress become available.", indent=False)
Hd("Practical Applications")
P("Designers of high waste dumps and rockfill dams often have few or no large-scale shear tests of their own and rely on "
  "published charts of rockfill strength. This work turns the most widely used chart into a tool that gives a best "
  "estimate of strength together with a cautious value for design, for any rock quality and density. It shows that the "
  "strength of crushed rock flattens at high pressure in the way expected from the loss of dilation, and that this "
  "flattening can be used safely within the range of available tests. A few laboratory tests of the site material at low "
  "pressure sharpen the estimate at higher pressure, and the tool warns the user when site tests disagree with published "
  "experience. Most of the remaining uncertainty for an untested material comes from the size of the rock particles and "
  "from differences between materials and laboratories, so testing the actual waste rock at large scale is the most "
  "valuable investigation. For a very high dump, the loose state of placed waste rock and rock types without tests "
  "governed the safety margin. Free software accompanies the paper.", indent=False)
P("*Author keywords:* Rockfill; Mine waste rock; Shear strength; Nonlinear failure envelope; Dilatancy; Hierarchical model; "
  "Reliability-based design; High waste dumps.", indent=False)

# ================================================================== INTRODUCTION
Hd("Introduction")
P("Leps (1970) compiled 109 large-scale triaxial tests and proposed three nonlinear envelopes for rockfill that have been "
  "used for dams and waste dumps whenever site-specific tests were unavailable. Cylwik et al. (2025) updated them with a "
  "public database of 938 tests (Arrieta 2025), adopted the power criterion of De Mello (1977) and Charles and Watts (1980) "
  "with a single fitting parameter, the secant friction angle at σ′n = 10p_a, and related that angle to rock quality and "
  "void ratio. That recent and widely adopted update is referred to as Leps 2.0.")
P("Two developments motivate a further extension. First, waste dumps now reach 500–700 m (Valenzuela et al. 2008), so "
  "effective normal stresses on deep failure surfaces approach or exceed the 7.9 MPa reached in the database, in a material "
  "that crushes and loses dilatancy (Marsal 1973; Hardin 1985; Lade et al. 1996; Linero et al. 2007; Ovalle et al. 2020). How "
  "the envelope behaves beyond the data then matters, and a power law with constant exponent keeps curving downward "
  "indefinitely, whereas the strength of a granular material is bounded below by its critical state (Bolton 1986; Jefferies "
  "and Been 2016). Second, design is moving toward reliability-based criteria (Read and Stacey 2009; Hawley and Cunning 2017), "
  "which require a strength distribution and a rule to combine published experience with site tests.")
P("Hierarchical models that borrow strength from generic databases have been developed for soil properties (Ching and Phoon "
  "2019; Ching et al. 2021), and a hierarchical Bayesian analysis of the Barton and Kjærnsli (1981) parameters has been applied "
  "to rockfill groups of a single project (Contreras et al. 2022; Contreras and Linero-Molina 2025). The contributions of this "
  "work are (1) an envelope whose curvature decays with stress as dilatancy is suppressed, tested against a fixed-form "
  "alternative out of sample and within materials; (2) the separation of publication and material variability; (3) a "
  "prespecified site-updating rule with a prior–data conflict check, compared with the simple fits used in practice; (4) a "
  "decomposition of the uncertainty of an untested waste rock at field scale; and (5) design charts and an open "
  "implementation. Leps-P is an independent extension built on the public Leps 2.0 database; it is not affiliated with the "
  "authors of Leps 2.0.")

# ================================================================== MODEL
Hd("Database and Model Formulation")
Hd("Calibration Database", 2)
P(f"The public database (Arrieta 2025) was filtered to tests with σ′3 ≥ 60 kPa, reported initial void ratio, and at least two tests "
  f"per set, giving {N['ensayos']} tests on {N['sets']} sets from {N['pub']} publications (Fig. 1). A set is a material with a given "
  "gradation, density, and rock type tested in one program; failure is the maximum deviator stress, as in Leps 2.0. Rock quality "
  "classes follow Table 5 of Cylwik et al. (2025). Tests with σ′n > 3 MPa (90 tests) come from eight publications, seven of them "
  "published between 1963 and 1983 (e.g., Marachi et al. 1972; Marsal 1973), and sets with φ′1 < 31° were excluded by Cylwik "
  "et al. (2025), so the lower tail of the population is truncated.")
Hd("Mechanical Basis of the Envelope", 2)
P("The Leps 2.0 criterion is the power law")
EQ(r"\tau=\tan\phi'_1\,\sigma_0^{1-m}\,\sigma_n'^{\,m}")
P("where σ0 = 10p_a, p_a = 0.1013 MPa, φ′1 is the secant friction angle at σ′n = σ0, and *m* is a constant exponent; the secant "
  "friction angle is φ′s = arctan(τ/σ′n). In granular materials the excess of strength over the critical-state value is the "
  "contribution of dilatancy, which is governed by the state of the material relative to its critical state and is "
  "suppressed by high confining stress (Bolton 1986; Jefferies and Been 2016, p. 43); particle crushing reinforces this "
  "trend (Hardin 1985; McDowell and Bolton 1998; Einav 2007). A minimal form with this property is", indent=False)
EQ(r"\tan\phi'_s=\tan\phi'_{cs}+\left(\tan\phi'_1-\tan\phi'_{cs}\right)\exp(-\zeta x),\quad x=\ln(\sigma'_n/\sigma_0)")
P("in which the dilatant part decays at rate ζ and φ′s tends to an asymptotic angle φ′cs. For φ′cs = 0 it reduces to a power "
  "law with *m* = 1 − ζ. Because ψ-type state measures (Jefferies and Been 2016) combine void ratio and stress, the effect of "
  "void ratio enters through tan φ′1 and therefore fades with stress, as observed in large triaxial tests (Marachi et al. "
  "1972; Xiao et al. 2014a). Here φ′cs is a population asymptote, not a measured critical-state angle.", indent=False)
Hd("Hierarchical Model", 2)
P("With *y* = ln(τ/σ0), each set *j* of publication *k* is a realization of a population:")
EQ(r"y_{ijk}=x_{ijk}+\ln\!\left[T+\left(e^{\eta_{jk}}-T\right)e^{-\zeta x_{ijk}}\right]+b_{0,jk}+b_{1,jk}\,x_{ijk}+\varepsilon_{ijk},\quad \eta_{jk}=\beta_{c(jk)}+\beta_e e_{jk}")
EQ(r"(b_0,b_1)_{jk}=\mathbf{b}^{(p)}_k+\mathbf{b}^{(s)}_{jk},\quad\mathbf{b}^{(p)}\sim N(\mathbf{0},\mathbf{G}_p),\quad\mathbf{b}^{(s)}\sim N(\mathbf{0},\mathbf{G}_s),\quad\ln\sigma_\varepsilon(x)=a_0+a_1x")
P("where T = tan φ′cs, β_c is the coefficient of rock quality class *c*, β_e the effect of initial void ratio *e*, b^(p) and b^(s) "
  "the publication and set random effects on the level and slope of the log–log envelope with covariance matrices G_p and G_s, "
  "and ε the dispersion of tests on one material, with standard deviation σε (Laird and Ware 1982). The publication effect "
  "combines laboratory, procedure, era, and material selection, which cannot be separated with these data; it plays the role of "
  "measurement and transformation uncertainty in the sense of Phoon and Kulhawy (1999), but because true strength and "
  "laboratory bias cannot be distinguished it is conservatively treated as uncertainty of strength. Through β_c and β_e the "
  "median φ′1 reproduces the relation φ′1 = φ′e=0 − ωe of Leps 2.0. A quadratic log–log form, *y* = θ + *m**x* + *c*₂*x*², and a "
  "power form (*c*₂ = 0), with the same random effects, were fitted for comparison.", indent=False)
Hd("Non-Identified Asymptote", 2)
P("The likelihood of (φ′cs, ζ) has a ridge: a lower asymptote with a slower decay fits the data almost as well as a higher "
  "asymptote with a faster decay, because both shapes coincide within the stress range of the tests. Rather than fixing "
  "φ′cs, the profile likelihood was evaluated on a grid (12–32° and 0.15–0.45), the grid points holding 99.5% of the "
  "normalized likelihood (flat prior on the grid) were refitted, and predictions average over them with these weights.", indent=False)
Hd("Predictive Distribution", 2)
P("For a new material of class *c* and void ratio *e* tested by a new laboratory, *y* is approximately normal with mean "
  "*f*(*x*; β̂) given by the first two terms of Eq. (3) and variance")
EQ(r"\mathrm{Var}[y]=\mathbf{z}^T(\mathbf{G}_p+\mathbf{G}_s)\mathbf{z}+\mathbf{j}^T\mathbf{C}_\beta\mathbf{j}+\sigma_\varepsilon^2(x),\quad \mathbf{z}=(1,x)^T")
P("where **j** is the gradient of *f* with respect to the fixed effects and C_β their covariance. The first two terms describe the "
  "uncertainty of the mean envelope of the material, which does not average out along a failure surface within that material "
  "(Christian et al. 1994); the last describes test-to-test scatter. In the implementation, variance parameters are sampled from "
  "their posterior, (φ′cs, ζ) from the ridge weights, and fixed effects from N(β̂, C_β); samples whose local slope exceeds one "
  "for σ′n ≤ 20 MPa are discarded, because a convex envelope has no support in granular materials. Fig. 2 and Table 3 give "
  "the resulting median and 5th-percentile envelopes for untested materials.", indent=False)
Hd("Updating with Site Tests", 2)
P("Site triaxial tests are converted to (σ′n, τ) with Eqs. (13), (14), and (16) of Cylwik et al. (2025), using the same failure "
  "criterion as the database. Consistency with the prior is checked with the prior predictive Mahalanobis distance")
EQ(r"D^2=\mathbf{r}^T\left(\mathbf{Z}(\mathbf{G}_p+\mathbf{G}_s)\mathbf{Z}^T+\mathbf{J}\mathbf{C}_\beta\mathbf{J}^T+\mathbf{S}\right)^{-1}\mathbf{r}\sim\chi^2_n")
P(r"where $\mathbf{r}$ holds the residuals *r*i of the *n* site tests from the prior mean, $\mathbf{Z}$ has rows (1, *x*i), and "
  r"$\mathbf{S}=\mathrm{diag}\,\sigma_\varepsilon^2(x_i)$. A *p*-value below 0.05 flags the tests for review. The prespecified rule "
  "updates only the level of the envelope at the centroid *x̄* of the tests: with q_0 = *b*₀ + *b*₁*x̄* and prior variance "
  "g_0 = (1, *x̄*)(G_p + G_s)(1, *x̄*)ᵀ, the posterior of q_0 is normal with", indent=False)
EQ(r"\hat q_0=\frac{\sum_i r_i/\sigma_{\varepsilon,i}^2}{1/g_0+\sum_i 1/\sigma_{\varepsilon,i}^2},\qquad \mathrm{Var}(q_0)=\left(\frac{1}{g_0}+\sum_i\frac{1}{\sigma_{\varepsilon,i}^2}\right)^{-1}")
P("and the slope keeps its prior distribution conditional on q_0, so the envelope pivots about the tested range. Eq. (7) is a "
  "cut (modular) update, which treats each residual as a measurement of the level at *x̄*; it can be applied in a spreadsheet. "
  "When several materials of one site were tested by the same laboratory, their publication-level effect is shared and all "
  "sets are updated jointly.", indent=False)
Hd("Placement Density and Specimen Size", 2)
P("A change of void ratio Δ*e* from the specimens to the placed material enters through η (β_e Δ*e*), so its effect decreases "
  "with stress. Specimen size enters as a shift of level")
EQ(r"\Delta y=\beta_D\ln\left(D_f/D_l\right)")
P("where *D*f and *D*l are the maximum particle sizes in the field and in the laboratory. Weibull scaling predicts β_D < 0 "
  "(Frossard et al. 2012; Ovalle et al. 2014; Alonso et al. 2012). β_D was estimated within materials of the database tested "
  "at several sizes.", indent=False)

# ================================================================== ESTIMATION & EVALUATION
Hd("Estimation and Out-of-Sample Evaluation")
P("Fixed effects were estimated by Gauss–Newton iteration on the linearized mean with restricted maximum likelihood (REML) for "
  "the variance parameters; random effects were integrated analytically. Variance parameters were then sampled with an "
  "adaptive Metropolis algorithm under weakly informative priors (half-normal with scale 0.25 on standard deviations, uniform "
  "on correlations, normal on a0, a1), with four chains started at dispersed values; the fixed effects are thus estimated "
  "conditionally (an empirical Bayes scheme). Convergence was checked with the rank-normalized potential scale reduction factor "
  "and bulk and tail effective sample sizes (Vehtari et al. 2021).")
P("Each publication was left out in turn. In each fold, (φ′cs, ζ) were selected on the training data only and the model was "
  "refitted. In Mode A (no site tests), every test of the held-out publication was predicted. In Mode B, each held-out set with "
  "at least two tests at σ′n ≤ 1.5 MPa and one at σ′n ≥ 2 MPa was updated with the former and the latter were predicted. "
  "The primary Mode B procedure (level-only update of every set) and the secondary comparisons were fixed before the bounded "
  "form was evaluated: the prior, a joint update of level and slope, the level-only update after discarding flagged sets, and "
  "least-squares fits of φ′1 with *m* = 0.90 or with *m* equal to the median within-set slope of the training data; the "
  "predictive standard deviation of the fits was their out-of-sample error on the training publications. Scores were the "
  "coverage of the 80% interval, the fraction of tests below the 10th percentile, and the continuous ranked probability score "
  "(CRPS) of ln τ (Gneiting and Raftery 2007). Differences were assessed by bootstrapping publications and by the Wilcoxon "
  "signed-rank test on publication means (Wilcoxon 1945). The Mode A reference was the published Leps 2.0 envelope with a "
  "constant dispersion fitted to the training data; because its constants were fitted to the whole database, this comparison "
  "favors the reference. A within-material shape test repeated Mode B for the quadratic and bounded forms on the same sets. "
  "The folds used plug-in predictive distributions (point estimates of all parameters); with all data, the distribution of the "
  f"implementation, which also propagates the posterior of the variance parameters, the ridge, and the slope restriction, differed "
  f"from the plug-in distribution by at most {DPL['dmed_abs_max']:.3f} in median ln τ and had a standard deviation "
  f"{DPL['ratio_sd_min']:.2f}–{DPL['ratio_sd_max']:.2f} times as large, so the evaluation applies to the implementation.")

# ================================================================== RESULTS
Hd("Results")
Hd("Parameters", 2)
P(f"Table 1 lists the parameters. The rank-normalized potential scale reduction factor was at most {max(M3['acotada']['rhat']):.3f}, and "
  f"the bulk and tail effective sample sizes were at least {min(M3['acotada']['ess_bulk']):.0f} and {min(M3['acotada']['ess_tail']):.0f}. "
  f"Replacing the half-normal scale of 0.25 by 0.1 or 0.5 changed the median between-publication standard deviation of the level "
  f"from {M3['acotada_previa_0.25']['hiper'][0][0]:.3f} to {M3['acotada_previa_0.1']['hiper'][0][0]:.3f} and {M3['acotada_previa_0.5']['hiper'][0][0]:.3f}. "
  f"At σ0, publication effects account for {pc(SHARE[0])} (90% credible interval {pc(SHARE[1])}–{pc(SHARE[2])}) of the "
  "between-set and between-publication variance of the level.")
Hd("Evidence for Stress-Dependent Curvature", 2)
P(f"Setting φ′cs = 0, which leaves a power law in tan φ′s with constant exponent, reduced the maximum log-likelihood by {dML0:.1f} "
  f"(likelihood ratio test with one degree of freedom, *p* = {p_LR:.1e}). The bounded form exceeded the quadratic form, which has "
  f"the same number of parameters, by {ML_ac - ML_q:.1f} in log-likelihood. Within 3 log-likelihood units of the maximum, "
  f"φ′cs ranged from {phi_lo:.0f}° to {phi_hi:.0f}° (Fig. 3a); the weighted means were φ′cs = {PHM:.1f}° and ζ = {ZEM:.2f}. "
  f"In the within-material shape test ({SHAPE['n']} tests, {SHAPE['sets']} sets, {SHAPE['pubs']} publications), the bounded form "
  f"reduced the mean bias of ln τ at higher stress from {sg(SHAPE['bq'])} to {sg(SHAPE['bb'])} and the CRPS from {f4(SHAPE['cq'])} "
  f"to {f4(SHAPE['cb'])}, and scored better in {SHAPE['mejora']} of {SHAPE['pubs']} publications (Wilcoxon *p* = {SHAPE['p']:.3f}; Fig. 3b). "
  f"For the site tests of the application, which were not used in calibration, it scored better in {SX['sets_mejora']} of {SX['sets']} "
  "sets, although both forms underpredicted the strength of these tests at their highest stress.")
Hd("Mode A: No Site Tests", 2)
P(f"Across all held-out tests (Table 2, Fig. 4a), the bounded form had a lower CRPS than the quadratic form, {dg(A_['acotada-cuadratica|crps'])}. "
  f"The 80% interval covered {pc(A_['acotada']['cob80'])} of the tests and {pc(A_['acotada']['bajo10'])} fell below the 10th percentile. "
  f"Above 2.75 MPa ({Ah['n']} tests, {Ah['n_pub']} publications), the bounded form reduced the mean standardized residual from "
  f"{sg(Ah['cuadratica']['z_medio'], 2)} to {sg(Ah['acotada']['z_medio'], 2)}, but the difference in CRPS was not significant "
  f"({dg(Ah['acotada-cuadratica|crps'])}), and the predictive distribution was too wide (standard deviation of standardized "
  f"residuals {f2(Ah['acotada']['sd_z'])}). The bounded form did not predict the strength of a new material better than Leps 2.0 "
  f"with constant dispersion ({dg(A_['acotada-leps2_ref|crps'])}); above 2.75 MPa the reference had a lower 10% quantile loss. "
  f"Below 0.5 MPa, however, the reference interval covered only {pc(Al['leps2_ref']['cob80'])} of tests, with the median biased low "
  f"(mean standardized residual {sg(Al['leps2_ref']['z_medio'], 2)}), whereas Leps-P covered {pc(Al['acotada']['cob80'])}.")
Hd("Mode B: Value of Site Tests", 2)
P(f"In Mode B ({NB} held-out tests from {NBP} publications), the prespecified level-only update reduced the CRPS from {f4(BP['crps'])} for the "
  f"prior to {f4(BN['crps'])} ({dg(BP['vs_nivel'])} for prior minus update; Fig. 4b). It also scored better than a least-squares fit "
  f"of φ′1 with *m* from the training data ({f4(BTR['crps'])}; {dg(BTR['vs_nivel'])}) and with *m* = 0.90 ({f4(B90['crps'])}, with "
  f"{pc(B90['bajo10'])} of tests below the 10th percentile). Joint updating of level and slope ({f4(BJ['crps'])}) and discarding "
  f"flagged sets before updating ({f4(BPR['crps'])}) were not distinguishable from the primary rule. The level-only intervals were "
  f"too wide (coverage {pc(BN['cob80'])} for a nominal 80%; standard deviation of standardized residuals {f2(BN['sd_z'])}). The "
  f"conflict check flagged {pc(CV['B|conflicto_frac'])} of the Mode B sets, close to its nominal 5%.")
Hd("Specimen Size", 2)
P(f"Sixteen materials of {EP['k']} publications were tested at maximum particle sizes differing by a factor of two or more. With "
  "set intercepts adjusted for void ratio, the between-set variance unrelated to size added to each estimate, and materials "
  f"combined within publications, the random-effects mean of β_D was {sg(EP['b'][0], 3)} (Knapp–Hartung 95% CI {sg(EP['ci'][0][0], 3)} "
  f"to {sg(EP['ci'][0][1], 3)}; Knapp and Hartung 2003), with a between-publication standard deviation of {f3(EP['tau'])}. A "
  f"material-level analysis gave {sg(EG['b'][0], 3)} ({sg(EG['ci'][0][0], 3)} to {sg(EG['ci'][0][1], 3)}), and neither the change in "
  "uniformity coefficient nor angularity explained the heterogeneity (Fig. S1). A systematic size effect was therefore not "
  "detected; for design, Leps-P uses a zero mean with this heterogeneity as uncertainty, which avoids crediting the positive "
  "estimate that runs against Weibull scaling.")
Hd("Sources of Uncertainty at Field Scale", 2)
P(f"For an untested material of low rock quality placed at *e* = 0.65 and scaled from 0.18 m to 1 m particles, the standard "
  f"deviation of the mean-envelope ln τ at 3 MPa corresponds to about "
  f"{math.degrees(math.sqrt(V3['total']) * math.sin(math.radians(V3['phi_mediana'])) * math.cos(math.radians(V3['phi_mediana']))):.1f}° in φ′s. "
  f"Specimen size contributed {vs(V3, 'escala_campo'):.0f}%, between-set variability {vs(V3, 'material'):.0f}%, publication "
  f"{vs(V3, 'publicacion'):.0f}%, the unidentified asymptote {vs(V3, 'cresta'):.0f}%, and the fixed effects {vs(V3, 'efectos_fijos'):.0f}% "
  f"(Fig. 5). At 10 MPa, beyond the data, the asymptote contributed {vs(V10, 'cresta'):.0f}%. The extrapolation of the envelope "
  "form is therefore a smaller source of uncertainty than the identity of the material and the scale of the particles.")

# ================================================================== APPLICATION
Hd("Application: Waste Rock from a Large Open-Pit Mine")
Hd("Site Data", 2)
P(f"The waste rock of a large open-pit mine hosted in carbonate rocks comprises five rock types, A to E, of which A is a "
  f"metacarbonate, B a silicate metasedimentary rock, C an igneous rock, and D and E carbonate and argillaceous sedimentary rocks. "
  f"Drained triaxial tests on large specimens of scalped gradation by one laboratory were available for A, B, and C "
  f"({len(SIT['sets'])} sets). Illustrative mass fractions of 35%, 25%, 20%, 10%, and 10% were used. Rock quality was taken as low "
  f"for A, C, D, and E and moderate for B. For rock type A, the conflict check gave *p* = {TD['clase_A']['baja']:.2f} with the low class, "
  f"{TD['clase_A']['moderada']:.2f} with the moderate class, and {TD['clase_A']['buena']:.3f} with the good class, so the tests were "
  "consistent with the class assigned.")
Hd("Conflict Check and Failure Criterion", 2)
pB = [s["p_conflicto"] for s in notused if s["lit"] == "B"]; pC = [s["p_conflicto"] for s in notused if s["lit"] == "C"]
P(f"The conflict check flagged {len(notused)} of the {len(SIT['sets'])} sets (Fig. 6a): both sets of rock type B (*p* ≤ {max(pB):.4f}) and "
  f"one set of rock type C (*p* = {pC[0]:.3f}). In both B sets the secant angle increased with stress and the specimens were still "
  "contracting at the maximum deviator, reached at 13–19% axial strain. Because the full stress–strain curves were tabulated in "
  "the test reports, the sets were re-reduced with the deviator at 10% and 15% axial strain: the free log–log slope of the B "
  f"sets remained above one ({min(CF['q10'][k]['m_libre'] for k in CF['q10'] if 'HF' in k):.2f}–{max(CF['qmax'][k]['m_libre'] for k in CF['qmax'] if 'HF' in k):.2f}), "
  "so the increase is not an artifact of the failure definition. Fixed-strain criteria created new conflicts in the other "
  "rock types, which confirms that site tests must be reduced with the criterion of the calibration database. Flagged sets "
  "did not update the envelope; because they are stronger than the prior at high stress, excluding them is conservative here, "
  "although not as a general rule. The remaining sets were updated jointly with a common laboratory effect.")
Hd("Placement Density and Dump-Average Envelope", 2)
la = SIT["escenarios"]["e=lab|ninguno"]["botadero"]["model"]; fa = SIT["escenarios"]["e=0.65|general"]["botadero"]["model"]
fq = SITQ["escenarios"]["e=0.65|general"]["botadero"]["model"]
P(f"Laboratory specimens had void ratios of {min(s['e'] for s in SIT['sets']):.2f}–{max(s['e'] for s in SIT['sets']):.2f}, whereas "
  "end-dumped waste rock is looser; a placement void ratio of 0.65 was adopted, with degradation of rock type E represented by an "
  "additional increase of 0.15. A failure surface through a mixed dump mobilizes the mass-weighted strength of the rock types it "
  "crosses; the envelopes were sampled jointly, sharing hyperparameters, fixed effects, laboratory effect, and size shift. At "
  f"σ′n = 3 MPa the median and 5th percentile of the dump-average φ′s were {la['P50'][isn(3)]:.1f}° and {la['P05'][isn(3)]:.1f}° at "
  f"laboratory density and size, and {fa['P50'][isn(3)]:.1f}° and {fa['P05'][isn(3)]:.1f}° for placed waste rock at field scale (Table 4; "
  f"Fig. 6b). At 20 MPa the bounded form gave {fa['P50'][-1]:.1f}° and {fa['P05'][-1]:.1f}°, against {fq['P50'][-1]:.1f}° and "
  f"{fq['P05'][-1]:.1f}° for the quadratic form. Full correlation between rock types lowered the 5th percentile at 3 MPa by "
  f"{fa['P05'][isn(3)] - SIT['escenarios']['e=0.65|general']['botadero']['comonotonic']['P05'][isn(3)]:.1f}°, and a surface confined to "
  f"the weakest rock type (E) by {fa['P05'][isn(3)] - [v for k_, v in SIT['escenarios']['e=0.65|general']['botadero'].items() if k_.startswith('weakest (')][0]['P05'][isn(3)]:.1f}°. "
  "The full-correlation case sorts the samples of each rock type at every stress, so it also removes the uncertainty of the slope.")
Hd("Stability of High Dumps", 2)
m7 = MC(K_MOD); mw = MC(K_W); ml_ = MC(kL + "model")
P(f"Dumps 150–700 m high with an overall slope of 2H:1V on a competent, drained, level foundation were analyzed with the method of "
  "Bishop (1955) and the tabulated nonlinear envelopes (Table 5). The unit weight was made consistent with the void ratio: "
  f"{EST['gam_base']:.1f} kN/m³ for *e* = 0.65 (specific gravity 2.75, water content 5%) and 21 kN/m³ for the laboratory density. For "
  f"700 m, the critical circle reached σ′n ≈ {DE(K_MOD)['P50']['snmax']:.1f} MPa, within the calibration range. The factor of safety "
  f"was {f2(DE(K_MOD)['P50']['F'])} with the median and {f2(DE(K_MOD)['P05']['F'])} with the 5th-percentile dump-average envelope, compared "
  f"with {f2(DE(kL + 'model')['P50']['F'])} and {f2(DE(kL + 'model')['P05']['F'])} at laboratory density and specimen size; with a unit weight of 21 kN/m³ "
  f"at *e* = 0.65 the latter values were {f2(DE('acotada e=0.65 γ21|700|model')['P50']['F'])} and {f2(DE('acotada e=0.65 γ21|700|model')['P05']['F'])}. "
  f"In a Monte Carlo analysis with 600 envelope samples per correlation assumption, each sample was evaluated on a family of circles "
  f"and, when the factor of safety was below 1.3, by a full search of the critical surface. The model correlation gave "
  f"{m7['fallas']} failures (probability of failure {cpi(m7)}, Clopper–Pearson 95% interval; Clopper and Pearson 1934), and a surface "
  f"confined to the untested, degraded rock type E gave {mw['fallas']} ({100 * mw['pf']:.1f}%, {cpi(mw)}). These are conditional "
  "probabilities that reflect strength uncertainty only, for a drained foundation and a two-dimensional section. The factor of "
  f"safety computed with the 5th-percentile envelope was close to the 5th percentile of the factor of safety ({f2(DE(K_MOD)['P05']['F'])} "
  f"versus {f2(m7['FS_P05'])}), because shared level shifts dominate the uncertainty here; this need not hold when pore pressure, "
  "curvature uncertainty, or independent materials dominate.")

# ================================================================== DISCUSSION
Hd("Discussion")
Hd("What Leps-P Adds", 2)
P("Leps-P does not predict the central strength of an untested material better than Leps 2.0. Its contributions are an envelope "
  "whose curvature decays with stress, which removes the low bias of fixed-form fits at high stress and is supported within "
  "materials; an uncertainty that is calibrated at low stress, where the constant-dispersion reference is not, and wide but "
  "unbiased at high stress; a prespecified updating rule that outperformed the simple fits used in practice; and a decomposition "
  "showing that, for an untested waste rock, the material and the particle scale matter more than the extrapolation of the "
  "envelope form. The last result indicates where investigation is most valuable: large-scale tests of the actual waste rock "
  "remove most of the material and laboratory components, and tests of one material at two particle sizes would constrain the "
  "largest remaining term.")
Hd("Design Format", 2)
P("A design procedure consistent with these results is: (1) assign rock quality and placement void ratio to each rock type and "
  "check the conflict diagnostic for any site tests, reduced with the maximum-deviator criterion; (2) update with the tests that "
  "pass, at the level only (Eq. 7); (3) shift the envelopes to placement density and field particle size; (4) adopt as "
  "characteristic value for a global mechanism the 5th percentile of the mean envelope, a Bayesian cautious estimate of the "
  "mean in the sense of Eurocode 7 (CEN 2004, 2024; Prästings et al. 2019); (5) check the probability of failure by sampling the "
  "envelopes with stated correlations between rock types, including the weakest-rock-type scenario; and (6) choose the target "
  "factor of safety for the 5th-percentile envelope from the consequence class. Table 5 pairs both quantities for the example: "
  f"a factor of safety of about {f2(DE(K_MOD)['P05']['F'])} with the 5th-percentile envelope corresponded to no failures in 600 samples, "
  f"and about {f2(DE(K_W)['P05']['F'])} to a conditional probability of failure near {100 * mw['pf']:.0f}%. Applying factors calibrated "
  "for best-estimate parameters (Read and Stacey 2009; Hawley and Cunning 2017) to the 5th-percentile envelope would count the "
  "strength uncertainty twice.")
Hd("Limitations", 2)
P("The database contains no tests above σ′n = 7.9 MPa and its tests above 3 MPa come from eight publications, so the bounded "
  "form is supported within the data and only its consequences are explored beyond; the asymptotic angle is not identified. "
  "The lower tail of the population is truncated at φ′1 = 31°, class labels may carry information from measured strength, and "
  "the within-material residual is laboratory scatter rather than spatial variability in the dump. The site application uses "
  "one laboratory and illustrative proportions. Pore pressure in saturated basal layers, wetting collapse (Oldecop and Alonso "
  "2001), seismic loading, three-dimensional geometry, and static liquefaction of contractive waste are outside the scope; they "
  "require their own analyses, for which the drained envelopes and their uncertainty are an input. Leps-P can be re-estimated as "
  "data appear, in particular drained tests at σ′3 of 5–10 MPa on carbonate and argillaceous waste rock and tests of one "
  "material at two sizes with parallel gradation.")

# ================================================================== CONCLUSIONS
Hd("Conclusions")
for t in [
    f"The calibration data support an envelope whose curvature decays with stress toward an asymptote (*p* = {p_LR:.0e}), consistent "
    "with the suppression of dilatancy; the asymptotic angle is not identified and is averaged over the likelihood ridge.",
    "Out of sample, the bounded form scored better than a quadratic log–log form and removed the low bias of the median at high "
    "stress, but it did not predict the central strength of a new material better than Leps 2.0 with constant dispersion.",
    "With a few site tests at low stress, a prespecified level-only update predicted the strength at higher stress better than "
    "least-squares fits with a fixed curvature; its intervals were wider than nominal.",
    "No systematic specimen-size effect was detected within materials; the heterogeneity of the size effect is the largest "
    "source of uncertainty of an untested waste rock at field scale, followed by material and laboratory variability.",
    "For a 700 m dump on a competent foundation, placement density and untested rock types controlled the margin of safety more "
    "than the form of the envelope.",
]:
    p = doc.add_paragraph(style="List Number"); runs(p, t)

# ================================================================== BACK MATTER
Hd("Data Availability Statement")
P("Some or all data, models, or code generated or used during the study are available in a repository or online in accordance "
  "with funder data retention policies (Leps-P code, parameters, and evaluation scripts: [[repository DOI]]; calibration "
  "database: Arrieta 2025). Some or all data, models, or code generated or used during the study are proprietary or "
  "confidential in nature and may only be provided with restrictions (the site test results, which may be provided in "
  "anonymized form with the permission of their owner).", indent=False)
Hd("Acknowledgments")
P("[[Acknowledgments, funding, and the permission of the owner of the site test data to be completed.]] [[Declaration of the "
  "use of AI tools, as required by ASCE policy: tool, purpose, and confirmation that the authors reviewed and take "
  "responsibility for the content.]]", indent=False)
Hd("Supplemental Materials")
P("Fig. S1 (specimen-size effect by material) and Table S1 (every Mode B procedure evaluated, for the bounded form and, in an "
  "earlier analysis, for the quadratic form with mixed updating rules) are available online in the ASCE Library (www.ascelibrary.org).", indent=False)
Hd("Notation")
P("*The following symbols are used in this paper:*", indent=False)
NOT = [("C_β", "covariance matrix of the fixed effects"), ("*D*f, *D*l", "maximum particle size in the field and in the laboratory"),
       ("$D^2$", "prior predictive Mahalanobis distance of the site tests"), ("G_p, G_s", "covariance matrices of the publication and set random effects"),
       ("**J**, **j**", "gradient of the mean envelope with respect to the fixed effects"), ("**r**", "vector of residuals of the site tests from the prior mean"),
       ("**S**", "diagonal matrix of test-to-test variances"), ("*T*", "tan φ′cs"), ("**Z**, **z**", "design rows (1, *x*) of the random effects"),
       ("a0, a1", "parameters of the stress-dependent test-to-test dispersion"), ("b^(p), b^(s)", "publication and set random effects on level and slope"),
       ("*b*₀, *b*₁", "level and slope deviations of a material in log–log space"), ("*e*", "initial void ratio"), ("g_0", "prior variance of q_0"),
       ("*m*", "exponent of the power envelope"), ("*n*", "number of site tests"), ("*p*", "*p*-value"), ("p_a", "atmospheric pressure (0.1013 MPa)"),
       ("q_0", "level of the envelope at *x̄*"), ("*r*i", "residual of site test *i*"), ("*x*", "ln(σ′n/σ0)"), ("*x̄*", "mean of *x* over the site tests of a set"),
       ("*y*", "ln(τ/σ0)"), ("β_c", "coefficient of rock quality class *c*"), ("β_D", "specimen-size coefficient"), ("β_e", "effect of void ratio on η"),
       ("ε", "test-to-test deviation"), ("ζ", "decay rate of the dilatant part of strength"), ("η", "ln tan φ′1 of the population"),
       ("σ0", "reference stress, 10p_a"), ("σ′3", "effective confining stress"), ("σ′n", "effective normal stress on the failure plane"),
       ("σε", "standard deviation of test-to-test scatter"), ("τ", "shear strength"), ("φ′1", "secant friction angle at σ′n = σ0"),
       ("φ′cs", "asymptotic secant friction angle"), ("φ′e=0", "φ′1 extrapolated to *e* = 0 in Leps 2.0"), ("φ′s", "secant friction angle"),
       ("ω", "decrease of φ′1 per unit void ratio in Leps 2.0")]
for s_, d_ in NOT:
    P(f"{s_} = {d_};", indent=False)

br(); Hd("References")
REFS = [
    "Alonso, E. E., M. Tapias, and J. Gili. 2012. \"Scale effects in rockfill behaviour.\" *Géotech. Lett.* 2 (3): 155–160. https://doi.org/10.1680/geolett.12.00025.",
    "Arrieta, M. 2025. \"Rockfill test database. Version 1.\" Zenodo. Accessed October 6, 2026. https://doi.org/10.5281/zenodo.14665504.",
    "Barton, N., and B. Kjærnsli. 1981. \"Shear strength of rockfill.\" *J. Geotech. Eng. Div.* 107 (7): 873–891. https://doi.org/10.1061/AJGEB6.0001167.",
    "Bishop, A. W. 1955. \"The use of the slip circle in the stability analysis of slopes.\" *Géotechnique* 5 (1): 7–17. https://doi.org/10.1680/geot.1955.5.1.7.",
    "Bolton, M. D. 1986. \"The strength and dilatancy of sands.\" *Géotechnique* 36 (1): 65–78. https://doi.org/10.1680/geot.1986.36.1.65.",
    "CEN (European Committee for Standardization). 2004. *Eurocode 7: Geotechnical design—Part 1: General rules.* EN 1997-1. Brussels, Belgium: CEN.",
    "CEN (European Committee for Standardization). 2024. *Eurocode 7: Geotechnical design—Part 1: General rules.* EN 1997-1:2024. Brussels, Belgium: CEN.",
    "Charles, J. A., and K. S. Watts. 1980. \"The influence of confining pressure on the shear strength of compacted rockfill.\" *Géotechnique* 30 (4): 353–367. https://doi.org/10.1680/geot.1980.30.4.353.",
    "Ching, J., and K.-K. Phoon. 2019. \"Constructing site-specific multivariate probability distribution model using Bayesian machine learning.\" *J. Eng. Mech.* 145 (1): 04018126. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001537.",
    "Ching, J., S. Wu, and K.-K. Phoon. 2021. \"Constructing quasi-site-specific multivariate probability distribution using hierarchical Bayesian model.\" *J. Eng. Mech.* 147 (10): 04021069. https://doi.org/10.1061/(ASCE)EM.1943-7889.0001964.",
    "Christian, J. T., C. C. Ladd, and G. B. Baecher. 1994. \"Reliability applied to slope stability analysis.\" *J. Geotech. Eng.* 120 (12): 2180–2207. https://doi.org/10.1061/(ASCE)0733-9410(1994)120:12(2180).",
    "Clopper, C. J., and E. S. Pearson. 1934. \"The use of confidence or fiducial limits illustrated in the case of the binomial.\" *Biometrika* 26 (4): 404–413. https://doi.org/10.1093/biomet/26.4.404.",
    "Contreras, L.-F., and S. Linero-Molina. 2025. \"Hierarchical Bayesian inference of rockfill strength parameters from large-scale triaxial test data.\" *Aust. Geomech.* 60 (2): 91–106. https://doi.org/10.56295/AGJ6024.",
    "Contreras, L.-F., S. Linero-Molina, and J. Dixon. 2022. \"Bayesian approach to improve the confidence of the estimation of the shear strength of coarse mine waste using Barton's empirical criterion.\" *Aust. Geomech.* 57 (2): 163–174. https://doi.org/10.56295/AGJ5728.",
    "Cylwik, S. D., M. Arrieta, J. C. Cunning, and P. M. Hawley. 2025. \"Update to the Leps strength envelopes for rockfill based on 900 large-scale triaxial tests.\" *Int. J. Geomech.* 25 (12): 04025285. https://doi.org/10.1061/IJGNAI.GMENG-12050.",
    "De Mello, V. F. B. 1977. \"Reflections on design decisions of practical significance to embankment dams.\" *Géotechnique* 27 (3): 281–355. https://doi.org/10.1680/geot.1977.27.3.281.",
    "Einav, I. 2007. \"Breakage mechanics—Part I: Theory.\" *J. Mech. Phys. Solids* 55 (6): 1274–1297. https://doi.org/10.1016/j.jmps.2006.11.003.",
    "Frossard, E., W. Hu, C. Dano, and P.-Y. Hicher. 2012. \"Rockfill shear strength evaluation: A rational method based on size effects.\" *Géotechnique* 62 (5): 415–427. https://doi.org/10.1680/geot.10.P.079.",
    "Gneiting, T., and A. E. Raftery. 2007. \"Strictly proper scoring rules, prediction, and estimation.\" *J. Am. Stat. Assoc.* 102 (477): 359–378. https://doi.org/10.1198/016214506000001437.",
    "Hardin, B. O. 1985. \"Crushing of soil particles.\" *J. Geotech. Eng.* 111 (10): 1177–1192. https://doi.org/10.1061/(ASCE)0733-9410(1985)111:10(1177).",
    "Hawley, M., and J. Cunning, eds. 2017. *Guidelines for mine waste dump and stockpile design.* Clayton, Australia: CSIRO Publishing.",
    "Jefferies, M., and K. Been. 2016. *Soil liquefaction: A critical state approach.* 2nd ed. Boca Raton, FL: CRC Press. https://doi.org/10.1201/b19114.",
    "Knapp, G., and J. Hartung. 2003. \"Improved tests for a random effects meta-regression with a single covariate.\" *Stat. Med.* 22 (17): 2693–2710. https://doi.org/10.1002/sim.1482.",
    "Lade, P. V., J. A. Yamamuro, and P. A. Bopp. 1996. \"Significance of particle crushing in granular materials.\" *J. Geotech. Eng.* 122 (4): 309–316. https://doi.org/10.1061/(ASCE)0733-9410(1996)122:4(309).",
    "Laird, N. M., and J. H. Ware. 1982. \"Random-effects models for longitudinal data.\" *Biometrics* 38 (4): 963–974. https://doi.org/10.2307/2529876.",
    "Leps, T. M. 1970. \"Review of shearing strength of rockfill.\" *J. Soil Mech. Found. Div.* 96 (4): 1159–1170. https://doi.org/10.1061/JSFEAQ.0001433.",
    "Linero, S., C. Palma, and R. Apablaza. 2007. \"Geotechnical characterisation of waste material in very high dumps with large scale triaxial testing.\" In *Proc., Int. Symp. on Rock Slope Stability in Open Pit Mining and Civil Engineering*, edited by Y. Potvin, 59–75. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/708_2.",
    "Marachi, N. D., C. K. Chan, and H. B. Seed. 1972. \"Evaluation of properties of rockfill materials.\" *J. Soil Mech. Found. Div.* 98 (1): 95–114. https://doi.org/10.1061/JSFEAQ.0001735.",
    "Marsal, R. J. 1973. \"Mechanical properties of rockfill.\" In *Embankment-dam engineering: Casagrande volume*, edited by R. C. Hirschfeld and S. J. Poulos, 109–200. New York: Wiley.",
    "McDowell, G. R., and M. D. Bolton. 1998. \"On the micromechanics of crushable aggregates.\" *Géotechnique* 48 (5): 667–679. https://doi.org/10.1680/geot.1998.48.5.667.",
    "Oldecop, L. A., and E. E. Alonso. 2001. \"A model for rockfill compressibility.\" *Géotechnique* 51 (2): 127–139. https://doi.org/10.1680/geot.2001.51.2.127.",
    "Ovalle, C., E. Frossard, C. Dano, W. Hu, S. Maiolino, and P.-Y. Hicher. 2014. \"The effect of size on the strength of coarse rock aggregates and large rockfill samples through experimental data.\" *Acta Mech.* 225 (8): 2199–2216. https://doi.org/10.1007/s00707-014-1127-z.",
    "Ovalle, C., S. Linero, C. Dano, E. Bard, P.-Y. Hicher, and R. Osses. 2020. \"Data compilation from large drained compression triaxial tests on coarse crushable rockfill materials.\" *J. Geotech. Geoenviron. Eng.* 146 (9): 06020013. https://doi.org/10.1061/(ASCE)GT.1943-5606.0002314.",
    "Phoon, K.-K., and F. H. Kulhawy. 1999. \"Characterization of geotechnical variability.\" *Can. Geotech. J.* 36 (4): 612–624. https://doi.org/10.1139/t99-038.",
    "Prästings, A., J. Spross, and S. Larsson. 2019. \"Characteristic values of geotechnical parameters in Eurocode 7.\" *Proc. Inst. Civ. Eng. Geotech. Eng.* 172 (4): 301–311. https://doi.org/10.1680/jgeen.18.00057.",
    "Read, J., and P. Stacey, eds. 2009. *Guidelines for open pit slope design.* Collingwood, Australia: CSIRO Publishing.",
    "Valenzuela, L., E. Bard, J. Campaña, and M. E. Anabalón. 2008. \"High waste rock dumps: Challenges and developments.\" In *Proc., 1st Int. Seminar on the Management of Rock Dumps, Stockpiles and Heap Leach Pads*, 65–78. Perth, Australia: Australian Centre for Geomechanics. https://doi.org/10.36487/ACG_repo/802_6.",
    "Vehtari, A., A. Gelman, D. Simpson, B. Carpenter, and P.-C. Bürkner. 2021. \"Rank-normalization, folding, and localization: An improved R̂ for assessing convergence of MCMC.\" *Bayesian Anal.* 16 (2): 667–718. https://doi.org/10.1214/20-BA1221.",
    "Wilcoxon, F. 1945. \"Individual comparisons by ranking methods.\" *Biometrics Bull.* 1 (6): 80–83. https://doi.org/10.2307/3001968.",
    "Xiao, Y., H. Liu, Y. Chen, and J. Jiang. 2014a. \"Strength and deformation of rockfill material based on large-scale triaxial compression tests. I: Influences of density and pressure.\" *J. Geotech. Geoenviron. Eng.* 140 (12): 04014070. https://doi.org/10.1061/(ASCE)GT.1943-5606.0001176.",
]
for ref in REFS:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5); runs(p, ref, False)

br(); Hd("Figure Captions List")
for c in [
    "Fig. 1. Distribution of effective normal stress on the failure plane in the calibration database, compared with stresses in a 700 m waste dump.",
    "Fig. 2. Design envelopes of Leps-P for a new material tested by a new laboratory, by rock quality and void ratio: median and 5th percentile of the mean envelope (characteristic value). The shaded zone exceeds the largest normal stress in the database.",
    "Fig. 3. Evidence for stress-dependent curvature: (a) profile log-likelihood of the asymptotic angle and decay rate relative to its maximum (dot); and (b) difference in CRPS between the bounded and quadratic forms when each material is updated at low stress and predicted at higher stress, by publication (marker size increases with the number of tests; filled: bounded form better).",
    "Fig. 4. Out-of-sample evaluation, leaving out each publication: (a) coverage of the 80% interval by range of normal stress without site tests; and (b) mean CRPS of ln τ for held-out tests at σ′n ≥ 2 MPa with site tests at σ′n ≤ 1.5 MPa (bars: 95% bootstrap confidence intervals over publications; filled: prespecified procedure).",
    "Fig. 5. Shares of the variance of the mean-envelope ln τ for an untested material of low rock quality placed at e = 0.65, scaled from 0.18 m to 1 m maximum particle size.",
    "Fig. 6. Waste rock of the example: (a) laboratory-scale envelopes by rock type (thin lines: 5th percentile) with the site tests (open symbols: tests flagged by the conflict check); and (b) envelopes of placed waste rock (void ratio 0.65, maximum particle size 1 m) by rock type and dump-average envelope. Shaded zones are beyond the stresses tested at the site.",
]:
    P(c, indent=False)


def tabla(n, titulo, enc, filas, nota=None):
    if n == 1:
        br()
    else:
        doc.add_paragraph()
    p = doc.add_paragraph(); rr = p.add_run(f"Table {n}. "); rr.bold = True; runs(p, titulo); p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(filas), cols=len(enc)); M_ = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for j, h in enumerate(enc):
        c = t.rows[0].cells[j]; c.text = ""; runs(c.paragraphs[0], h)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f):
            c = t.rows[i].cells[j]; c.text = ""; pp = c.paragraphs[0]; runs(pp, str(v))
            if not any(r_.text.strip() for r_ in pp.runs):
                for om in pp._p.findall(f"{{{M_}}}oMath"):
                    para = etree.SubElement(pp._p, f"{{{M_}}}oMathPara"); ppr = etree.SubElement(para, f"{{{M_}}}oMathParaPr")
                    jc = etree.SubElement(ppr, f"{{{M_}}}jc"); jc.set(f"{{{M_}}}val", "left"); pp._p.remove(om); para.append(om)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing = 1.0
                for r_ in pp.runs:
                    r_.font.size = Pt(10)
    tblPr = t._tbl.tblPr; bd = OxmlElement("w:tblBorders")
    for side in ("top", "bottom"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); bd.append(e)
    tblPr.append(bd)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); cb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); cb.append(e); tcPr.append(cb)
    if nota:
        p = doc.add_paragraph(); runs(p, "Note: " + nota)


rg = lambda i, d=3: f"{HP[i][1]:.{d}f}–{HP[i][2]:.{d}f}"
tabla(1, "Parameters of Leps-P (bounded form)", ["Parameter", "Estimate", "Standard error or 90% credible interval"],
      [["β_c, good rock quality", f3(b["good"]), f3(SE["good"])], ["β_c, moderate rock quality", f3(b["moderate"]), f3(SE["moderate"])],
       ["β_c, low rock quality", f3(b["low"]), f3(SE["low"])], ["β_e", neg(f3(b["e"])), f3(SE["e"])],
       ["φ′cs (°), weighted mean over the ridge", f"{PHM:.1f}", f"{phi_lo:.0f}–{phi_hi:.0f} (ΔML ≤ 3)"], ["ζ, weighted mean over the ridge", f2(ZEM), "—"],
       ["SD of level, publication", f3(HP[0][0]), rg(0)], ["SD of slope, publication", f3(HP[1][0]), rg(1)],
       ["Correlation of level and slope, publication", neg(f2(HP[2][0])), neg(f"{HP[2][1]:.2f} to {HP[2][2]:.2f}")],
       ["SD of level, set", f3(HP[3][0]), rg(3)], ["SD of slope, set", f3(HP[4][0]), rg(4)],
       ["Correlation of level and slope, set", neg(f2(HP[5][0])), neg(f"{HP[5][1]:.2f} to {HP[5][2]:.2f}")],
       ["σε at σ′n = σ0", f3(HP[6][0]), rg(6)], ["*a*₁", neg(f3(HP[7][0])), neg(f"{HP[7][1]:.3f} to {HP[7][2]:.3f}")]],
      f"{N['ensayos']} tests, {N['sets']} sets, {N['pub']} publications. Fixed effects at the maximum-likelihood point of the ridge "
      f"(φ′cs = {FA['phi_cs']:.0f}°, ζ = {FA['zeta']}); variance parameters: posterior median and 90% credible interval. Median φ′1 at e = 0: "
      f"{phi1['good']:.1f}°, {phi1['moderate']:.1f}°, and {phi1['low']:.1f}° for good, moderate, and low rock quality.")
r_ = lambda o: [pc(o["cob80"]), pc(o["bajo10"]), f4(o["crps"])]
tabla(2, "Out-of-sample evaluation (leave one publication out)", ["Procedure", "80% coverage", "Below P10", "CRPS"],
      [["Mode A, all stresses: Leps-P bounded"] + r_(A_["acotada"]), ["Mode A, all stresses: Leps-P quadratic"] + r_(A_["cuadratica"]),
       ["Mode A, all stresses: Leps 2.0, constant dispersion"] + r_(A_["leps2_ref"]),
       ["Mode A, σ′n > 2.75 MPa: Leps-P bounded"] + r_(Ah["acotada"]), ["Mode A, σ′n > 2.75 MPa: Leps-P quadratic"] + r_(Ah["cuadratica"]),
       ["Mode A, σ′n > 2.75 MPa: Leps 2.0, constant dispersion"] + r_(Ah["leps2_ref"]),
       ["Mode B: prior, no site tests"] + r_(BP), ["Mode B: fit of φ′1, *m* = 0.90"] + r_(B90), ["Mode B: fit of φ′1, *m* from training data"] + r_(BTR),
       ["Mode B: joint update"] + r_(BJ), ["Mode B: level-only update after conflict check"] + r_(BPR), ["Mode B: level-only update (prespecified)"] + r_(BN)],
      f"Mode A: {A_['n']} tests ({A_['n_pub']} publications); above 2.75 MPa, {Ah['n']} tests ({Ah['n_pub']} publications). Mode B: tests at "
      "σ′n ≥ 2 MPa predicted from tests at σ′n ≤ 1.5 MPa of the same set. Nominal values: 80% coverage and 10% below P10. CRPS in units of ln τ.")
fil = []
for cal, nom in (("buena", "Good"), ("moderada", "Moderate"), ("baja", "Low")):
    for e in (0.3, 0.45, 0.65):
        v = TD["tabla"][f"{cal}|{e}"]; fil.append([nom if e == 0.3 else "", f"{e:.2f}"] + [f"{a:.1f} / {c:.1f}" for a, c in zip(v["P50"], v["P05"])])
tabla(3, "Design values of the secant friction angle φ′s (°) for an untested material: median / 5th percentile of the mean envelope",
      ["Rock quality", "*e*"] + [f"{s:g} MPa" for s in TD["sn"]], fil,
      "Laboratory specimen size, new laboratory, no site tests; bounded form averaged over the likelihood ridge. Values at 8 MPa are slightly beyond the database (7.9 MPa).")
fil = []
lab = SIT["escenarios"]["e=lab|ninguno"]["litologias"]; fld = SIT["escenarios"]["e=0.65|general"]["litologias"]
for k in "ABCDE":
    a_ = lab.get(k); f_ = fld[k]
    fil.append([f"Rock type {k}{' (untested)' if k in 'DE' else ''}", f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}",
                f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}", f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(10)]:.1f} / {f_['P05'][isn(10)]:.1f}"])
for nm, k in (("Dump average, model correlation", "model"), ("Dump average, full correlation", "comonotonic")):
    a_ = SIT["escenarios"]["e=lab|ninguno"]["botadero"][k]; f_ = SIT["escenarios"]["e=0.65|general"]["botadero"][k]
    fil.append([nm, f"{a_['P50'][isn(3)]:.1f} / {a_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(1)]:.1f} / {f_['P05'][isn(1)]:.1f}",
                f"{f_['P50'][isn(3)]:.1f} / {f_['P05'][isn(3)]:.1f}", f"{f_['P50'][isn(10)]:.1f} / {f_['P05'][isn(10)]:.1f}"])
tabla(4, "Secant friction angle φ′s (°) of the waste rock of the example, median / 5th percentile of the mean envelope",
      ["Material", "Laboratory, 3 MPa", "Placed, 1 MPa", "Placed, 3 MPa", "Placed, 10 MPa"], fil,
      "Laboratory: specimen void ratio and size. Placed: void ratio 0.65 (rock type E, 0.80), maximum particle size 1 m, size effect with zero mean. No site tests above σ′n ≈ 4.6 MPa.")
fil = []
for k, nm in ((K_MOD, "Dump average, model correlation"), (K_COM, "Dump average, full correlation"), (K_WT, "Weakest tested rock type only"), (K_W, "Untested, degraded rock type E only")):
    base = k.rsplit("|", 2)[0] + "|"; tail = k.split("|")[-1]
    m_ = MC(k)
    fil.append([nm] + [f"{f2(DE(base.replace('|700|', '|') + f'{h}|{tail}')['P50']['F'])} / {f2(DE(base.replace('|700|', '|') + f'{h}|{tail}')['P05']['F'])}" for h in (150, 400, 700)]
               + [f"{m_['fallas']}/{m_['n']} ({cpi(m_)})"])
mq_ = MC(kQ + "model")
tabla(5, "Factor of safety of dumps on a competent level foundation, placed waste rock (median / 5th-percentile envelope) and conditional probability of failure at 700 m",
      ["Envelope", "150 m", "400 m", "700 m", "Failures, 700 m (95% CI)"], fil,
      f"Bishop method; unit weight {EST['gam_base']:.1f} kN/m³. Monte Carlo with 600 envelope samples and a full search when FS < 1.3; Clopper–Pearson intervals. "
      f"At laboratory density and specimen size (21 kN/m³) the 700 m dump gave {f2(DE(kL + 'model')['P50']['F'])} / {f2(DE(kL + 'model')['P05']['F'])} and {ml_['fallas']} failures; "
      f"with the quadratic form at e = 0.65, {f2(DE(kQ + 'model')['P50']['F'])} / {f2(DE(kQ + 'model')['P05']['F'])} and {mq_['fallas']} failures.")
out = AQ / "LepsP_IJGM_manuscript_r3.docx"; doc.save(str(out)); print("written", out, "| equations:", NEQ[0])
