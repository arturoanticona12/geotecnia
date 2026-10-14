import sys, os
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
src, outdir, tag = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(outdir, exist_ok=True)
prs = Presentation(src)
out = []
def walk(shapes, sl, depth=0):
    for sh in shapes:
        t = sh.shape_type
        pre = "  "*depth + f"[{sh.shape_id} {t} '{sh.name}']"
        if t == MSO_SHAPE_TYPE.GROUP:
            out.append(pre); walk(sh.shapes, sl, depth+1); continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(pre + " TEXT: " + sh.text_frame.text.replace("\n"," | "))
        if getattr(sh, "has_table", False) and sh.has_table:
            out.append(pre + " TABLE:")
            for r in sh.table.rows:
                out.append("  "*depth + "   || " + " ; ".join(c.text.replace("\n"," ") for c in r.cells))
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            out.append(pre + f" CHART type={ch.chart_type}")
            try:
                cats = list(ch.plots[0].categories)
                out.append("     cats: " + str(cats))
            except Exception as e: out.append(f"     cats err {e}")
            for pl in ch.plots:
                for s in pl.series:
                    try: vals = list(s.values)
                    except Exception: vals = "?"
                    out.append(f"     series '{s.name}': {vals}")
        if t == MSO_SHAPE_TYPE.PICTURE:
            img = sh.image
            fn = f"{tag}_s{sl:02d}_{sh.shape_id}.{img.ext}"
            with open(os.path.join(outdir, fn), "wb") as f: f.write(img.blob)
            out.append(pre + f" PICTURE -> {fn} ({img.size}) {sh.width/914400:.1f}x{sh.height/914400:.1f} in")
for i, s in enumerate(prs.slides, 1):
    out.append(f"\n===== SLIDE {i} =====")
    walk(s.shapes, i)
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        out.append("NOTES: " + s.notes_slide.notes_text_frame.text)
open(os.path.join(outdir, tag + "_dump.txt"), "w", encoding="utf-8").write("\n".join(out))
print(len(prs.slides), "slides")
