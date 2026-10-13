import openpyxl, pathlib, re, json
import numpy as np
f=pathlib.Path(r"H:\Mi unidad\00_WORKSPACE\20_PROYECTOS_CONSULTORÍA\Antamina\2026_Gate3\Revisión de literatura\Leps 2.0\Data Antamina\250903 Waste Charact\Resultados Parametros  Geotecnicos LE2 (03092025).xlsx")
wbf=openpyxl.load_workbook(f,data_only=False)
wbv=openpyxl.load_workbook(f,data_only=True)
spec={"UCS":("D","E",["G","I","J"]),"PLTA":("B","C",["G","I","J"]),"PLTD":("B","C",["G","I","J"]),"Brazilian":("D","E",["G","I","J"])}
out={}
for sh,(cr,cv,fcols) in spec.items():
    wsf=wbf[sh]; wsv=wbv[sh]
    rows=[]; nonnum=[]
    for i in range(2,wsf.max_row+1):
        rmu=wsf[f"{cr}{i}"].value; val=wsf[f"{cv}{i}"].value
        if rmu is None and val is None: continue
        if isinstance(val,(int,float)): rows.append((i,str(rmu).strip(),float(val)))
        else: nonnum.append((i,rmu,val))
    # blocks of contiguous RMU
    blocks=[]
    for i,r,v in rows:
        if blocks and blocks[-1][0]==r and blocks[-1][2]==i-1: blocks[-1][2]=i
        elif blocks and blocks[-1][0]==r: blocks[-1][2]=i  # allow gaps from nonnum
        else: blocks.append([r,i,i])
    print(f"===== {sh}: n numeric={len(rows)}; non-numeric={len(nonnum)} ejemplo {nonnum[:5]}")
    print("  bloques RMU (fila ini-fin):",[(b[0],b[1],b[2]) for b in blocks])
    # formulas
    for i in range(2,wsf.max_row+1):
        fs=[(c,wsf[f"{c}{i}"].value,wsv[f"{c}{i}"].value) for c in fcols if isinstance(wsf[f"{c}{i}"].value,str) and wsf[f"{c}{i}"].value.startswith("=")]
        if fs:
            rr=[r for (ii,r,v) in rows if ii==i]
            print("  fila",i,"RMU",rr,[(c,fo,round(vv,3) if isinstance(vv,(int,float)) else vv) for c,fo,vv in fs])
    st={}
    for r in sorted(set(x[1] for x in rows)):
        a=np.array([x[2] for x in rows if x[1]==r])
        st[r]=dict(n=int(a.size),media=float(a.mean()),P50=float(np.percentile(a,50)),P35=float(np.percentile(a,35)),P16=float(np.percentile(a,16)),P84=float(np.percentile(a,84)),sd=float(a.std(ddof=1)) if a.size>1 else None,min=float(a.min()),max=float(a.max()))
        print("  %-4s n=%5d media=%.3f P50=%.3f P35=%.3f P16=%.3f P84=%.3f sd=%.3f min=%.3f max=%.3f"%(r,st[r]['n'],st[r]['media'],st[r]['P50'],st[r]['P35'],st[r]['P16'],st[r]['P84'],st[r]['sd'] or 0,st[r]['min'],st[r]['max']))
    out[sh]=dict(stats=st,nonnum=[(a,str(b),str(c)) for a,b,c in nonnum])
json.dump(out,open(pathlib.Path(__file__).parent/"le2_stats.json","w"),indent=1)
# Resumen sheet dump (values and formulas) rows with content
wsf=wbf["Resumen"]; wsv=wbv["Resumen"]
for i in range(1,60):
    vals=[wsv.cell(i,j).value for j in range(1,12)]
    if any(v is not None for v in vals): print("R",i,vals)
for i in range(1,60):
    fo=[(wsf.cell(i,j).coordinate,wsf.cell(i,j).value) for j in range(1,12) if isinstance(wsf.cell(i,j).value,str) and wsf.cell(i,j).value.startswith("=")]
    if fo: print("F",i,fo)
