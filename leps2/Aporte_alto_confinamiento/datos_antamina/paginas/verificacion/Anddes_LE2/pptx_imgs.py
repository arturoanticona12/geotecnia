from pptx import Presentation
import pathlib
d=pathlib.Path(r"H:\Mi unidad\00_WORKSPACE\20_PROYECTOS_CONSULTORÍA\Antamina\2026_Gate3\Revisión de literatura\Leps 2.0\Data Antamina\250903 Waste Charact")
out=pathlib.Path(__file__).parent
for f,slides,tag in [(d/"Waste Sample Characterization Rev C.pptx",(8,12),"WSC"),(d/"Parametros Geotecnicos Vs RMU LE2 03092025.pptx",(4,),"PGR")]:
    p=Presentation(f)
    for i,s in enumerate(p.slides,1):
        if i not in slides: continue
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type==6: walk(sh.shapes); continue
                try:
                    im=sh.image
                except Exception: continue
                fn=out/f"{tag}_l{i}_{sh.shape_id}.{im.ext}"; fn.write_bytes(im.blob); print(fn.name, im.size if hasattr(im,'size') else '', sh.left, sh.top, sh.width, sh.height)
        walk(s.shapes)
