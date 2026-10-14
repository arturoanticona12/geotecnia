from pptx import Presentation
from pptx.util import Emu
import pathlib, re
d=pathlib.Path(r"H:\Mi unidad\00_WORKSPACE\20_PROYECTOS_CONSULTORÍA\Antamina\2026_Gate3\Revisión de literatura\Leps 2.0\Data Antamina\250903 Waste Charact")
for f in d.glob("*.pptx"):
    p=Presentation(f); print("=====",f.name)
    for i,s in enumerate(p.slides,1):
        txt=[]
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type==6: walk(sh.shapes); continue
                if sh.has_text_frame and sh.text_frame.text.strip(): txt.append(sh.text_frame.text.replace("\n"," / "))
                if getattr(sh,"has_table",False) and sh.has_table:
                    for r in sh.table.rows: txt.append(" | ".join(c.text for c in r.cells))
        walk(s.shapes)
        j=" || ".join(txt)
        if re.search(r"XDP|LMH|END|EXO|HRN|INT\b|RMU|Densidad|densidad|Mármol|Marble|BD",j): print(f"--- lámina {i}: {j[:3000]}")
