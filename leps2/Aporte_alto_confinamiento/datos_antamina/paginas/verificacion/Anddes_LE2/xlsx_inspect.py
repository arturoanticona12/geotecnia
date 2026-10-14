import openpyxl, pathlib
f=pathlib.Path(r"H:\Mi unidad\00_WORKSPACE\20_PROYECTOS_CONSULTORÍA\Antamina\2026_Gate3\Revisión de literatura\Leps 2.0\Data Antamina\250903 Waste Charact\Resultados Parametros  Geotecnicos LE2 (03092025).xlsx")
wb=openpyxl.load_workbook(f,data_only=False,read_only=False)
for ws in wb.worksheets:
    print("=== ",ws.title,ws.dimensions,ws.max_row,ws.max_column)
    for r in ws.iter_rows(min_row=1,max_row=4,values_only=True):
        print("   ",[c for c in r[:22]])
